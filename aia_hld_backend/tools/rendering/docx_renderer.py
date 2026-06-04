# from __future__ import annotations

# import logging
# import re
# from pathlib import Path
# from typing import Any, Dict, List, Optional, Tuple

# from docx import Document
# from docx.enum.text import WD_ALIGN_PARAGRAPH
# from docx.oxml import OxmlElement
# from docx.oxml.ns import qn
# from docx.shared import Inches, Pt, RGBColor
# from PIL import Image

# from .config import (
#     DOCUMENT_TITLE,
#     DOCX_FONT_NAME,
#     DOCX_FONT_SIZE_PT,
#     DOCX_TOP_MARGIN_IN,
#     DOCX_BOTTOM_MARGIN_IN,
#     DOCX_LEFT_MARGIN_IN,
#     DOCX_RIGHT_MARGIN_IN,
#     DOCX_LOGO_WIDTH_IN,
#     DOCX_MAX_IMAGE_WIDTH_FALLBACK_IN,
#     DOCX_MAX_IMAGE_HEIGHT_FALLBACK_IN,
#     DOCX_TOC_OUTLINE,
#     DOCX_UPDATE_FIELDS_ON_OPEN,
#     DOCX_KEEP_DIAGRAM_SOURCE_LINKS,
# )

# from .graphviz_renderer import (
#     normalize_to_graphviz_dot,
#     render_graphviz_to_png,
#     render_graphviz_to_assets,
# )

# from .schema_utils import (
#     METADATA_SECTIONS,
#     is_simple_metadata_dict,
#     should_include_section,
# )

# from .text_sanitizer import (
#     display_text_for_renderer,
#     renderer_safe_plain_text,
# )

# logger = logging.getLogger(__name__)


# # ---------------------------------------------------------------------
# # Internal token cleanup
# # ---------------------------------------------------------------------
# def clean_docx_internal_tokens(text: Any) -> str:
#     """
#     Remove renderer-only tokens so they never appear in DOCX output.

#     These tokens are only for PDF HTML rendering and must not be visible
#     in Word documents.
#     """
#     value = str(text or "")

#     tokens = [
#         "AIASECTIONBLOCKSTARTTOKEN",
#         "AIASECTIONBLOCKENDTOKEN",
#         "[[AIASECTIONBLOCKSTARTTOKEN]]",
#         "[[AIASECTIONBLOCKENDTOKEN]]",
#         "SECTION_BLOCK_START",
#         "SECTION_BLOCK_END",
#         "[[SECTION_BLOCK_START]]",
#         "[[SECTION_BLOCK_END]]",
#         "[[SECTIONBLOCKSTART]]",
#         "[[SECTIONBLOCKEND]]",
#     ]

#     for token in tokens:
#         value = value.replace(token, "")

#     return value.strip()


# def _is_empty_or_na(value: Any) -> bool:
#     if value is None:
#         return True

#     if isinstance(value, str):
#         return value.strip().lower() in {
#             "",
#             "n/a",
#             "na",
#             "none",
#             "null",
#             "not applicable",
#         }

#     if isinstance(value, list):
#         return len(value) == 0

#     if isinstance(value, dict):
#         return len(value) == 0

#     return False


# def _default_issues_rows() -> List[Dict[str, str]]:
#     return [
#         {
#             "id": "I001",
#             "description": "TBC - No issues have been identified at draft stage.",
#             "status": "Open",
#             "mitigation": "TBC",
#         }
#     ]


# # ---------------------------------------------------------------------
# # DOCX base helpers
# # ---------------------------------------------------------------------
# def set_docx_default_font(
#     doc: Document,
#     font_name: Optional[str] = None,
#     font_size: Optional[int] = None,
# ):
#     resolved_font_name = font_name or DOCX_FONT_NAME
#     resolved_font_size = font_size or DOCX_FONT_SIZE_PT

#     if "Normal" in doc.styles:
#         doc.styles["Normal"].font.name = resolved_font_name
#         doc.styles["Normal"].font.size = Pt(resolved_font_size)
#         doc.styles["Normal"]._element.rPr.rFonts.set(
#             qn("w:eastAsia"),
#             resolved_font_name,
#         )


# def add_docx_page_break(doc: Document):
#     """
#     Add a page break safely.
#     """
#     doc.add_page_break()


# def set_paragraph_keep_with_next(
#     paragraph,
#     keep_next: bool = True,
#     keep_lines: bool = True,
# ):
#     """
#     Configure Word paragraph pagination.

#     keep_next:
#       Keeps this paragraph with the following paragraph.

#     keep_lines:
#       Keeps lines in the paragraph together.
#     """
#     pPr = paragraph._p.get_or_add_pPr()

#     if keep_next:
#         keep_next_el = OxmlElement("w:keepNext")
#         pPr.append(keep_next_el)

#     if keep_lines:
#         keep_lines_el = OxmlElement("w:keepLines")
#         pPr.append(keep_lines_el)


# def shade_cell(
#     cell,
#     fill: str = "D9D9D9",
# ):
#     """
#     Apply background shading to a DOCX table cell.

#     fill should be hex color without '#'.
#     Example:
#       D9D9D9 = light grey
#       D9EAF7 = light blue
#     """
#     tc_pr = cell._tc.get_or_add_tcPr()

#     shd = OxmlElement("w:shd")
#     shd.set(qn("w:fill"), fill)

#     tc_pr.append(shd)


# def style_table_header_row(
#     table,
#     fill: str = "D9D9D9",
# ):
#     """
#     Make first row bold, center aligned, and apply background shading.
#     """
#     if not table.rows:
#         return

#     header_row = table.rows[0]

#     for cell in header_row.cells:
#         shade_cell(
#             cell,
#             fill=fill,
#         )

#         for paragraph in cell.paragraphs:
#             paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

#             for run in paragraph.runs:
#                 run.bold = True


# def set_table_no_row_split(table):
#     """
#     Prevent table rows from splitting across pages where Word can honour it.
#     Also repeats the first row as header row.
#     """
#     if table.rows:
#         trPr = table.rows[0]._tr.get_or_add_trPr()

#         hdr = OxmlElement("w:tblHeader")
#         hdr.set(qn("w:val"), "true")
#         trPr.append(hdr)

#     for row in table.rows:
#         trPr = row._tr.get_or_add_trPr()

#         cant = OxmlElement("w:cantSplit")
#         cant.set(qn("w:val"), "true")
#         trPr.append(cant)


# def get_docx_max_image_size(
#     doc: Document,
#     reserved_height_in: float = 0.0,
# ) -> Tuple[float, float]:
#     """
#     Return usable image size for the current DOCX page.

#     reserved_height_in is used to keep space for:
#       - section/subsection headings
#       - editable diagram links
#       - Flow Legend heading/table
#       - spacing around diagram block
#     """
#     try:
#         section = doc.sections[-1]

#         usable_width = max(
#             4.5,
#             section.page_width.inches
#             - section.left_margin.inches
#             - section.right_margin.inches,
#         )

#         usable_height = max(
#             3.8,
#             section.page_height.inches
#             - section.top_margin.inches
#             - section.bottom_margin.inches
#             - reserved_height_in,
#         )

#         return usable_width, usable_height

#     except Exception:
#         return (
#             DOCX_MAX_IMAGE_WIDTH_FALLBACK_IN,
#             max(
#                 3.8,
#                 DOCX_MAX_IMAGE_HEIGHT_FALLBACK_IN - reserved_height_in,
#             ),
#         )


# def add_picture_fit_to_page(
#     doc: Document,
#     image_path: str,
#     reserved_height_in: float = 0.0,
# ):
#     """
#     Add a PNG/JPEG image fitted to the current page.

#     The image paragraph is marked keep-with-next so Word tries to keep the
#     diagram with the following Flow Legend.
#     """
#     max_width_in, max_height_in = get_docx_max_image_size(
#         doc,
#         reserved_height_in=reserved_height_in,
#     )

#     width_in = max_width_in
#     height_in = None

#     try:
#         with Image.open(image_path) as img:
#             if img.width > 0 and img.height > 0:
#                 aspect = img.height / img.width
#                 predicted_height = width_in * aspect

#                 if predicted_height > max_height_in:
#                     height_in = max_height_in
#                     width_in = height_in / aspect

#     except Exception:
#         logger.warning(
#             "Could not inspect image size for %s; using width-only fit",
#             image_path,
#         )

#     paragraph = doc.add_paragraph()
#     paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

#     run = paragraph.add_run()

#     if height_in is not None:
#         run.add_picture(
#             str(image_path),
#             width=Inches(width_in),
#             height=Inches(height_in),
#         )
#     else:
#         run.add_picture(
#             str(image_path),
#             width=Inches(width_in),
#         )

#     set_paragraph_keep_with_next(
#         paragraph,
#         keep_next=True,
#         keep_lines=True,
#     )


# # ---------------------------------------------------------------------
# # DOCX hyperlink helpers
# # ---------------------------------------------------------------------
# def add_docx_hyperlink(
#     paragraph,
#     text: str,
#     url: str,
# ):
#     """
#     Add an external hyperlink to a DOCX paragraph.

#     Used for editable diagram asset links:
#       - DOT source
#       - SVG vector
#     """
#     part = paragraph.part

#     relationship_id = part.relate_to(
#         url,
#         "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
#         is_external=True,
#     )

#     hyperlink = OxmlElement("w:hyperlink")
#     hyperlink.set(qn("r:id"), relationship_id)

#     run = OxmlElement("w:r")
#     run_properties = OxmlElement("w:rPr")

#     color = OxmlElement("w:color")
#     color.set(qn("w:val"), "0563C1")
#     run_properties.append(color)

#     underline = OxmlElement("w:u")
#     underline.set(qn("w:val"), "single")
#     run_properties.append(underline)

#     run.append(run_properties)

#     text_element = OxmlElement("w:t")
#     text_element.text = text
#     run.append(text_element)

#     hyperlink.append(run)
#     paragraph._p.append(hyperlink)


# def add_editable_diagram_links(
#     doc: Document,
#     dot_path: Optional[str] = None,
#     svg_path: Optional[str] = None,
# ):
#     """
#     Add links above the diagram preview for editable diagram assets.

#     This paragraph is marked keep-with-next so it stays with the diagram image.
#     """
#     if not dot_path and not svg_path:
#         return

#     paragraph = doc.add_paragraph()

#     label_run = paragraph.add_run("Edit this diagram: ")
#     label_run.bold = True

#     added_any = False

#     if dot_path and Path(dot_path).exists():
#         add_docx_hyperlink(
#             paragraph,
#             "DOT source",
#             Path(dot_path).resolve().as_uri(),
#         )
#         added_any = True

#     if svg_path and Path(svg_path).exists():
#         if added_any:
#             paragraph.add_run(" | ")

#         add_docx_hyperlink(
#             paragraph,
#             "SVG vector",
#             Path(svg_path).resolve().as_uri(),
#         )

#     set_paragraph_keep_with_next(
#         paragraph,
#         keep_next=True,
#         keep_lines=True,
#     )


# # ---------------------------------------------------------------------
# # Text/table rendering
# # ---------------------------------------------------------------------
# def add_docx_paragraph_from_text(
#     doc: Document,
#     text: str,
# ):
#     clean_text = clean_docx_internal_tokens(text)

#     if not clean_text:
#         doc.add_paragraph("N/A")
#         return

#     safe_text = renderer_safe_plain_text(clean_text)

#     paragraph = doc.add_paragraph()

#     for idx, line in enumerate(safe_text.split("\n")):
#         if idx > 0:
#             paragraph.add_run().add_break()

#         paragraph.add_run(line)


# def add_docx_table(
#     doc: Document,
#     items: List[Dict[str, Any]],
# ):
#     if not items:
#         doc.add_paragraph("N/A")
#         return

#     all_keys: List[str] = []

#     for row in items:
#         for key in row.keys():
#             if key not in all_keys:
#                 all_keys.append(key)

#     table = doc.add_table(
#         rows=1,
#         cols=len(all_keys),
#     )
#     table.style = "Table Grid"

#     for idx, key in enumerate(all_keys):
#         table.rows[0].cells[idx].text = clean_docx_internal_tokens(
#             key.replace("_", " ").title()
#         )

#     style_table_header_row(
#         table,
#         fill="D9D9D9",
#     )

#     for row in items:
#         cells = table.add_row().cells

#         for idx, key in enumerate(all_keys):
#             cells[idx].text = clean_docx_internal_tokens(
#                 display_text_for_renderer(
#                     row.get(key, ""),
#                     multiline=True,
#                 )
#             )

#     set_table_no_row_split(table)
#     doc.add_paragraph("")


# def add_docx_metadata_table(
#     doc: Document,
#     data: Dict[str, Any],
# ):
#     if not data:
#         doc.add_paragraph("N/A")
#         return

#     add_docx_table(doc, [data])


# # ---------------------------------------------------------------------
# # Diagram rendering
# # ---------------------------------------------------------------------
# def add_flow_legend_docx(
#     doc: Document,
#     legend_items: List[Tuple[str, str]],
# ):
#     if not legend_items:
#         return

#     heading = doc.add_paragraph()

#     heading_run = heading.add_run("Flow Legend")
#     heading_run.bold = True
#     heading_run.font.size = Pt(11)

#     set_paragraph_keep_with_next(
#         heading,
#         keep_next=True,
#         keep_lines=True,
#     )

#     table = doc.add_table(
#         rows=1,
#         cols=2,
#     )
#     table.style = "Table Grid"

#     table.rows[0].cells[0].text = "Flow"
#     table.rows[0].cells[1].text = "Description"

#     style_table_header_row(
#         table,
#         fill="D9D9D9",
#     )

#     for color, label in legend_items:
#         row = table.add_row().cells

#         marker = row[0].paragraphs[0].add_run("■")
#         marker.font.color.rgb = RGBColor.from_string(
#             color.lstrip("#").upper()
#         )

#         row[1].text = clean_docx_internal_tokens(str(label))

#     set_table_no_row_split(table)
#     doc.add_paragraph("")


# def add_docx_diagram_with_legend(
#     doc: Document,
#     img_path: str,
#     legend: List[Tuple[str, str]],
# ):
#     """
#     Add diagram and legend as one visual block.

#     Reserve vertical space so headings, edit links, image, and legend
#     fit on one page as much as Word allows.
#     """
#     reserved_height_in = 1.35

#     if legend:
#         reserved_height_in = min(
#             3.6,
#             1.35 + (0.24 * len(legend)),
#         )

#     add_picture_fit_to_page(
#         doc,
#         img_path,
#         reserved_height_in=reserved_height_in,
#     )

#     add_flow_legend_docx(doc, legend)


# def _safe_diagram_base_name(
#     section_key: Optional[str],
#     diagram_index: int,
# ) -> str:
#     safe_section = re.sub(
#         r"[^a-zA-Z0-9_-]+",
#         "_",
#         str(section_key or "diagram"),
#     ).strip("_").lower()

#     if not safe_section:
#         safe_section = "diagram"

#     return f"{safe_section}_{diagram_index:02d}"


# def add_docx_diagram_or_placeholder(
#     doc: Document,
#     diagram_text: str,
#     section_key: Optional[str] = None,
#     diagram_assets_dir: Optional[Path] = None,
#     diagram_counter: Optional[Dict[str, int]] = None,
# ):
#     """
#     Render DOT/Mermaid-like diagram into DOCX.

#     Important:
#       Do not force page break before or after diagram here.

#     Reason:
#       build_docx() already controls top-level section page breaks.
#       Adding page breaks inside each diagram block can create blank pages,
#       especially when the next section also starts with a page break.
#     """
#     clean_diagram_text = clean_docx_internal_tokens(diagram_text)

#     if diagram_counter is not None:
#         diagram_counter["value"] = diagram_counter.get("value", 0) + 1
#         diagram_index = diagram_counter["value"]
#     else:
#         diagram_index = 1

#     base_name = _safe_diagram_base_name(
#         section_key=section_key,
#         diagram_index=diagram_index,
#     )

#     def _start_diagram_page():
#         # Intentionally no page break before diagram.
#         return

#     def _end_diagram_page():
#         # Intentionally no page break after diagram.
#         # This prevents extra/blank pages in DOCX.
#         return

#     # Preferred path: render PNG + DOT + SVG assets.
#     # Reverted to stable DOCX behaviour:
#     #   - always embed PNG preview
#     #   - optionally keep DOT/SVG source links above the diagram
#     if diagram_assets_dir is not None:
#         assets = render_graphviz_to_assets(
#             clean_diagram_text,
#             output_dir=diagram_assets_dir,
#             base_name=base_name,
#         )

#         if assets:
#             try:
#                 _start_diagram_page()

#                 if DOCX_KEEP_DIAGRAM_SOURCE_LINKS:
#                     add_editable_diagram_links(
#                         doc,
#                         dot_path=assets.get("dot_path"),
#                         svg_path=assets.get("svg_path"),
#                     )

#                 png_path = assets.get("png_path")

#                 if png_path and Path(png_path).exists():
#                     add_docx_diagram_with_legend(
#                         doc,
#                         png_path,
#                         assets.get("legend", []),
#                     )
#                 else:
#                     logger.error(
#                         "PNG diagram asset missing for DOCX. section_key=%s assets=%s",
#                         section_key,
#                         assets,
#                     )
#                     doc.add_paragraph(
#                         "[Diagram rendering failed] PNG preview image was not generated."
#                     )

#                 _end_diagram_page()

#             except Exception:
#                 logger.exception(
#                     "Failed to embed diagram assets in DOCX. section_key=%s",
#                     section_key,
#                 )
#                 doc.add_paragraph("[Diagram failed while embedding into DOCX]")

#             return

#     # Backward-compatible fallback: PNG only.
#     result = render_graphviz_to_png(clean_diagram_text)

#     if result:
#         img_path, legend = result

#         try:
#             _start_diagram_page()

#             add_docx_diagram_with_legend(
#                 doc,
#                 img_path,
#                 legend,
#             )

#             _end_diagram_page()

#         except Exception:
#             logger.exception(
#                 "Failed to embed diagram in DOCX. section_key=%s",
#                 section_key,
#             )
#             doc.add_paragraph("[Diagram failed while embedding into DOCX]")

#         return

#     logger.error(
#         "Diagram text detected but rendering failed. "
#         "DOT source will not be printed into DOCX. section_key=%s",
#         section_key,
#     )

#     doc.add_paragraph(
#         "[Diagram rendering failed] The diagram source was detected but could "
#         "not be rendered by Graphviz. Please check renderer logs for the DOT "
#         "syntax error."
#     )


# # ---------------------------------------------------------------------
# # Recursive DOCX value renderer
# # ---------------------------------------------------------------------
# def add_docx_value(
#     doc: Document,
#     val: Any,
#     section_key: Optional[str] = None,
#     diagram_assets_dir: Optional[Path] = None,
#     diagram_counter: Optional[Dict[str, int]] = None,
# ):
#     if val is None:
#         doc.add_paragraph("N/A")
#         return

#     if (
#         isinstance(val, dict)
#         and section_key in METADATA_SECTIONS
#         and is_simple_metadata_dict(val)
#     ):
#         add_docx_metadata_table(doc, val)
#         return

#     if isinstance(val, str):
#         clean_val = clean_docx_internal_tokens(val)
#         is_diagram_text = normalize_to_graphviz_dot(clean_val) is not None

#         if is_diagram_text:
#             add_docx_diagram_or_placeholder(
#                 doc,
#                 clean_val,
#                 section_key=section_key,
#                 diagram_assets_dir=diagram_assets_dir,
#                 diagram_counter=diagram_counter,
#             )
#             return

#         add_docx_paragraph_from_text(doc, clean_val)
#         return

#     if isinstance(val, list):
#         if not val:
#             doc.add_paragraph("N/A")
#             return

#         # If all items are diagram strings, render all diagrams.
#         if all(
#             isinstance(item, str)
#             and normalize_to_graphviz_dot(clean_docx_internal_tokens(item))
#             for item in val
#         ):
#             for item in val:
#                 add_docx_diagram_or_placeholder(
#                     doc,
#                     clean_docx_internal_tokens(item),
#                     section_key=section_key,
#                     diagram_assets_dir=diagram_assets_dir,
#                     diagram_counter=diagram_counter,
#                 )
#             return

#         # If all items are dicts, render as table.
#         if all(isinstance(item, dict) for item in val):
#             add_docx_table(doc, val)
#             return

#         # Mixed list: render diagram strings as diagrams, others as bullets.
#         for item in val:
#             if (
#                 isinstance(item, str)
#                 and normalize_to_graphviz_dot(clean_docx_internal_tokens(item))
#             ):
#                 add_docx_diagram_or_placeholder(
#                     doc,
#                     clean_docx_internal_tokens(item),
#                     section_key=section_key,
#                     diagram_assets_dir=diagram_assets_dir,
#                     diagram_counter=diagram_counter,
#                 )
#             else:
#                 doc.add_paragraph(
#                     clean_docx_internal_tokens(
#                         display_text_for_renderer(
#                             item,
#                             multiline=False,
#                         )
#                     ),
#                     style="List Bullet",
#                 )

#         doc.add_paragraph("")
#         return

#     if isinstance(val, dict):
#         for key, child_value in val.items():
#             key_lower = str(key).lower()
#             title = str(key).replace("_", " ").title()

#             if key_lower == "issues" and _is_empty_or_na(child_value):
#                 heading = doc.add_heading(
#                     title,
#                     level=3,
#                 )
#                 set_paragraph_keep_with_next(
#                     heading,
#                     keep_next=True,
#                     keep_lines=True,
#                 )

#                 add_docx_table(
#                     doc,
#                     _default_issues_rows(),
#                 )
#                 continue

#             heading = doc.add_heading(
#                 title,
#                 level=3,
#             )
#             set_paragraph_keep_with_next(
#                 heading,
#                 keep_next=True,
#                 keep_lines=True,
#             )

#             # If key is diagrams, keep parent section_key so generated file names
#             # relate to the parent section, not a fake "diagrams" section.
#             if key_lower == "diagrams":
#                 add_docx_value(
#                     doc,
#                     child_value,
#                     section_key=section_key,
#                     diagram_assets_dir=diagram_assets_dir,
#                     diagram_counter=diagram_counter,
#                 )
#             else:
#                 add_docx_value(
#                     doc,
#                     child_value,
#                     section_key=key,
#                     diagram_assets_dir=diagram_assets_dir,
#                     diagram_counter=diagram_counter,
#                 )

#         return

#     doc.add_paragraph(clean_docx_internal_tokens(str(val)))


# # ---------------------------------------------------------------------
# # TOC helpers
# # ---------------------------------------------------------------------
# def add_docx_toc(doc: Document):
#     paragraph = doc.add_paragraph()
#     run = paragraph.add_run()

#     begin = OxmlElement("w:fldChar")
#     begin.set(qn("w:fldCharType"), "begin")
#     run._r.append(begin)

#     instr = OxmlElement("w:instrText")
#     instr.set(qn("xml:space"), "preserve")
#     instr.text = f'TOC \\o "{DOCX_TOC_OUTLINE}" \\h \\z \\u'
#     run._r.append(instr)

#     separate = OxmlElement("w:fldChar")
#     separate.set(qn("w:fldCharType"), "separate")
#     run._r.append(separate)

#     end = OxmlElement("w:fldChar")
#     end.set(qn("w:fldCharType"), "end")
#     run._r.append(end)


# def enable_docx_update_fields(doc: Document):
#     update = OxmlElement("w:updateFields")
#     update.set(qn("w:val"), "true")
#     doc.settings._element.append(update)


# # ---------------------------------------------------------------------
# # Main DOCX builder
# # ---------------------------------------------------------------------
# def build_docx(
#     data: Dict[str, Any],
#     docx_path: str,
#     hld_model: Any,
#     selected_sections_set: Optional[set[str]] = None,
#     logo_path: Optional[str] = None,
# ):
#     doc = Document()

#     docx_output_path = Path(docx_path)

#     diagram_assets_dir = (
#         docx_output_path.parent
#         / f"{docx_output_path.stem}_diagram_assets"
#     )
#     diagram_assets_dir.mkdir(
#         parents=True,
#         exist_ok=True,
#     )

#     diagram_counter = {
#         "value": 0,
#     }

#     set_docx_default_font(
#         doc,
#         font_name=DOCX_FONT_NAME,
#         font_size=DOCX_FONT_SIZE_PT,
#     )

#     for section in doc.sections:
#         section.top_margin = Inches(DOCX_TOP_MARGIN_IN)
#         section.bottom_margin = Inches(DOCX_BOTTOM_MARGIN_IN)
#         section.left_margin = Inches(DOCX_LEFT_MARGIN_IN)
#         section.right_margin = Inches(DOCX_RIGHT_MARGIN_IN)

#     title = doc.add_heading(
#         DOCUMENT_TITLE,
#         level=1,
#     )
#     title.alignment = WD_ALIGN_PARAGRAPH.CENTER

#     set_paragraph_keep_with_next(
#         title,
#         keep_next=True,
#         keep_lines=True,
#     )

#     if logo_path:
#         try:
#             logo_paragraph = doc.add_paragraph()
#             logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
#             logo_paragraph.paragraph_format.space_before = Pt(24)

#             logo_run = logo_paragraph.add_run()
#             logo_run.add_picture(
#                 str(logo_path),
#                 width=Inches(DOCX_LOGO_WIDTH_IN),
#             )

#         except Exception:
#             logger.warning(
#                 "Logo not inserted into DOCX: %s",
#                 logo_path,
#             )

#     doc.add_page_break()
#     add_docx_toc(doc)
#     doc.add_page_break()

#     if DOCX_UPDATE_FIELDS_ON_OPEN:
#         enable_docx_update_fields(doc)

#     first = True

#     for field_name, field_info in hld_model.model_fields.items():
#         if not should_include_section(
#             field_name,
#             field_info,
#             selected_sections_set or set(),
#         ):
#             continue

#         if not first:
#             doc.add_page_break()

#         first = False

#         section_title = field_info.title or field_name.replace("_", " ").title()

#         heading = doc.add_heading(
#             section_title,
#             level=2,
#         )
#         set_paragraph_keep_with_next(
#             heading,
#             keep_next=True,
#             keep_lines=True,
#         )

#         add_docx_value(
#             doc,
#             data.get(field_name),
#             section_key=field_name,
#             diagram_assets_dir=diagram_assets_dir,
#             diagram_counter=diagram_counter,
#         )

#     doc.save(docx_path)

# from __future__ import annotations

# import logging
# import re
# from pathlib import Path
# from typing import Any, Dict, List, Optional, Tuple

# from docx import Document
# from docx.enum.text import WD_ALIGN_PARAGRAPH
# from docx.oxml import OxmlElement
# from docx.oxml.ns import qn
# from docx.shared import Inches, Pt, RGBColor
# from PIL import Image

# from .config import (
#     DOCUMENT_TITLE,
#     DOCX_FONT_NAME,
#     DOCX_FONT_SIZE_PT,
#     DOCX_TOP_MARGIN_IN,
#     DOCX_BOTTOM_MARGIN_IN,
#     DOCX_LEFT_MARGIN_IN,
#     DOCX_RIGHT_MARGIN_IN,
#     DOCX_LOGO_WIDTH_IN,
#     DOCX_MAX_IMAGE_WIDTH_FALLBACK_IN,
#     DOCX_MAX_IMAGE_HEIGHT_FALLBACK_IN,
#     DOCX_TOC_OUTLINE,
#     DOCX_UPDATE_FIELDS_ON_OPEN,
#     DOCX_KEEP_DIAGRAM_SOURCE_LINKS,
# )
# from .editable_diagram_exporter import DiagramAssets, export_editable_diagram
# from .graphviz_renderer import normalize_to_graphviz_dot, render_graphviz_to_png
# from .schema_utils import (
#     METADATA_SECTIONS,
#     is_simple_metadata_dict,
#     should_include_section,
# )
# from .text_sanitizer import display_text_for_renderer, renderer_safe_plain_text

# logger = logging.getLogger(__name__)


# # ---------------------------------------------------------------------
# # Internal token cleanup
# # ---------------------------------------------------------------------
# def clean_docx_internal_tokens(text: Any) -> str:
#     """
#     Remove renderer-only tokens so they never appear in DOCX output.

#     These tokens are only for PDF/HTML rendering and must not be visible
#     in Word documents.
#     """
#     value = str(text or "")

#     tokens = [
#         "AIASECTIONBLOCKSTARTTOKEN",
#         "AIASECTIONBLOCKENDTOKEN",
#         "[[AIASECTIONBLOCKSTARTTOKEN]]",
#         "[[AIASECTIONBLOCKENDTOKEN]]",
#         "SECTION_BLOCK_START",
#         "SECTION_BLOCK_END",
#         "[[SECTION_BLOCK_START]]",
#         "[[SECTION_BLOCK_END]]",
#         "[[SECTIONBLOCKSTART]]",
#         "[[SECTIONBLOCKEND]]",
#     ]

#     for token in tokens:
#         value = value.replace(token, "")

#     return value.strip()


# def _is_empty_or_na(value: Any) -> bool:
#     if value is None:
#         return True

#     if isinstance(value, str):
#         return value.strip().lower() in {
#             "",
#             "n/a",
#             "na",
#             "none",
#             "null",
#             "not applicable",
#         }

#     if isinstance(value, list):
#         return len(value) == 0

#     if isinstance(value, dict):
#         return len(value) == 0

#     return False


# def _default_issues_rows() -> List[Dict[str, str]]:
#     return [
#         {
#             "id": "I001",
#             "description": "TBC - No issues have been identified at draft stage.",
#             "status": "Open",
#             "mitigation": "TBC",
#         }
#     ]


# # ---------------------------------------------------------------------
# # DOCX base helpers
# # ---------------------------------------------------------------------
# def set_docx_default_font(
#     doc: Document,
#     font_name: Optional[str] = None,
#     font_size: Optional[int] = None,
# ):
#     resolved_font_name = font_name or DOCX_FONT_NAME
#     resolved_font_size = font_size or DOCX_FONT_SIZE_PT

#     if "Normal" in doc.styles:
#         doc.styles["Normal"].font.name = resolved_font_name
#         doc.styles["Normal"].font.size = Pt(resolved_font_size)
#         doc.styles["Normal"]._element.rPr.rFonts.set(
#             qn("w:eastAsia"),
#             resolved_font_name,
#         )


# def add_docx_page_break(doc: Document):
#     doc.add_page_break()


# def set_paragraph_keep_with_next(
#     paragraph,
#     keep_next: bool = True,
#     keep_lines: bool = True,
# ):
#     pPr = paragraph._p.get_or_add_pPr()

#     if keep_next:
#         keep_next_el = OxmlElement("w:keepNext")
#         pPr.append(keep_next_el)

#     if keep_lines:
#         keep_lines_el = OxmlElement("w:keepLines")
#         pPr.append(keep_lines_el)


# def shade_cell(cell, fill: str = "D9D9D9"):
#     tc_pr = cell._tc.get_or_add_tcPr()
#     shd = OxmlElement("w:shd")
#     shd.set(qn("w:fill"), fill)
#     tc_pr.append(shd)


# def style_table_header_row(table, fill: str = "D9D9D9"):
#     if not table.rows:
#         return

#     header_row = table.rows[0]

#     for cell in header_row.cells:
#         shade_cell(cell, fill=fill)

#         for paragraph in cell.paragraphs:
#             paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

#             for run in paragraph.runs:
#                 run.bold = True


# def set_table_no_row_split(table):
#     """
#     Prevent table rows from splitting across pages where Word can honour it.
#     Also repeats the first row as header row.
#     """
#     if table.rows:
#         trPr = table.rows[0]._tr.get_or_add_trPr()
#         hdr = OxmlElement("w:tblHeader")
#         hdr.set(qn("w:val"), "true")
#         trPr.append(hdr)

#     for row in table.rows:
#         trPr = row._tr.get_or_add_trPr()
#         cant = OxmlElement("w:cantSplit")
#         cant.set(qn("w:val"), "true")
#         trPr.append(cant)


# def get_docx_max_image_size(
#     doc: Document,
#     reserved_height_in: float = 0.0,
# ) -> Tuple[float, float]:
#     try:
#         section = doc.sections[-1]

#         usable_width = max(
#             4.5,
#             section.page_width.inches
#             - section.left_margin.inches
#             - section.right_margin.inches,
#         )

#         usable_height = max(
#             3.8,
#             section.page_height.inches
#             - section.top_margin.inches
#             - section.bottom_margin.inches
#             - reserved_height_in,
#         )

#         return usable_width, usable_height

#     except Exception:
#         return (
#             DOCX_MAX_IMAGE_WIDTH_FALLBACK_IN,
#             max(3.8, DOCX_MAX_IMAGE_HEIGHT_FALLBACK_IN - reserved_height_in),
#         )


# def add_picture_fit_to_page(
#     doc: Document,
#     image_path: str,
#     reserved_height_in: float = 0.0,
# ):
#     max_width_in, max_height_in = get_docx_max_image_size(
#         doc,
#         reserved_height_in=reserved_height_in,
#     )

#     width_in = max_width_in
#     height_in = None

#     try:
#         with Image.open(image_path) as img:
#             if img.width > 0 and img.height > 0:
#                 aspect = img.height / img.width
#                 predicted_height = width_in * aspect

#                 if predicted_height > max_height_in:
#                     height_in = max_height_in
#                     width_in = height_in / aspect

#     except Exception:
#         logger.warning(
#             "Could not inspect image size for %s; using width-only fit",
#             image_path,
#         )

#     paragraph = doc.add_paragraph()
#     paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

#     run = paragraph.add_run()

#     if height_in is not None:
#         run.add_picture(
#             str(image_path),
#             width=Inches(width_in),
#             height=Inches(height_in),
#         )
#     else:
#         run.add_picture(str(image_path), width=Inches(width_in))

#     set_paragraph_keep_with_next(
#         paragraph,
#         keep_next=True,
#         keep_lines=True,
#     )


# # ---------------------------------------------------------------------
# # DOCX hyperlink helpers
# # ---------------------------------------------------------------------
# def add_docx_hyperlink(paragraph, text: str, url: str):
#     part = paragraph.part

#     relationship_id = part.relate_to(
#         url,
#         "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
#         is_external=True,
#     )

#     hyperlink = OxmlElement("w:hyperlink")
#     hyperlink.set(qn("r:id"), relationship_id)

#     run = OxmlElement("w:r")
#     run_properties = OxmlElement("w:rPr")

#     color = OxmlElement("w:color")
#     color.set(qn("w:val"), "0563C1")
#     run_properties.append(color)

#     underline = OxmlElement("w:u")
#     underline.set(qn("w:val"), "single")
#     run_properties.append(underline)

#     run.append(run_properties)

#     text_element = OxmlElement("w:t")
#     text_element.text = text
#     run.append(text_element)

#     hyperlink.append(run)
#     paragraph._p.append(hyperlink)


# def add_editable_diagram_links(
#     doc: Document,
#     dot_path: Optional[str] = None,
#     svg_path: Optional[str] = None,
#     drawio_path: Optional[str] = None,
# ):
#     """
#     Add links above the diagram preview for editable diagram assets.

#     Links:
#       - DOT source
#       - SVG vector
#       - draw.io editable diagram
#     """
#     if not dot_path and not svg_path and not drawio_path:
#         return

#     paragraph = doc.add_paragraph()

#     label_run = paragraph.add_run("Edit this diagram: ")
#     label_run.bold = True

#     added_any = False

#     if dot_path and Path(dot_path).exists():
#         add_docx_hyperlink(paragraph, "DOT source", Path(dot_path).resolve().as_uri())
#         added_any = True

#     if svg_path and Path(svg_path).exists():
#         if added_any:
#             paragraph.add_run(" | ")
#         add_docx_hyperlink(paragraph, "SVG vector", Path(svg_path).resolve().as_uri())
#         added_any = True

#     if drawio_path and Path(drawio_path).exists():
#         if added_any:
#             paragraph.add_run(" | ")
#         add_docx_hyperlink(
#             paragraph,
#             "draw.io editable",
#             Path(drawio_path).resolve().as_uri(),
#         )

#     set_paragraph_keep_with_next(
#         paragraph,
#         keep_next=True,
#         keep_lines=True,
#     )


# # ---------------------------------------------------------------------
# # Text/table rendering
# # ---------------------------------------------------------------------
# def add_docx_paragraph_from_text(doc: Document, text: str):
#     clean_text = clean_docx_internal_tokens(text)

#     if not clean_text:
#         doc.add_paragraph("N/A")
#         return

#     safe_text = renderer_safe_plain_text(clean_text)
#     paragraph = doc.add_paragraph()

#     for idx, line in enumerate(safe_text.split("\n")):
#         if idx > 0:
#             paragraph.add_run().add_break()
#         paragraph.add_run(line)


# def add_docx_table(doc: Document, items: List[Dict[str, Any]]):
#     if not items:
#         doc.add_paragraph("N/A")
#         return

#     all_keys: List[str] = []

#     for row in items:
#         for key in row.keys():
#             if key not in all_keys:
#                 all_keys.append(key)

#     table = doc.add_table(rows=1, cols=len(all_keys))
#     table.style = "Table Grid"

#     for idx, key in enumerate(all_keys):
#         table.rows[0].cells[idx].text = clean_docx_internal_tokens(
#             key.replace("_", " ").title()
#         )

#     style_table_header_row(table, fill="D9D9D9")

#     for row in items:
#         cells = table.add_row().cells

#         for idx, key in enumerate(all_keys):
#             cells[idx].text = clean_docx_internal_tokens(
#                 display_text_for_renderer(row.get(key, ""), multiline=True)
#             )

#     set_table_no_row_split(table)
#     doc.add_paragraph("")


# def add_docx_metadata_table(doc: Document, data: Dict[str, Any]):
#     if not data:
#         doc.add_paragraph("N/A")
#         return

#     add_docx_table(doc, [data])


# # ---------------------------------------------------------------------
# # Diagram rendering
# # ---------------------------------------------------------------------
# def add_flow_legend_docx(doc: Document, legend_items: List[Tuple[str, str]]):
#     if not legend_items:
#         return

#     heading = doc.add_paragraph()
#     heading_run = heading.add_run("Flow Legend")
#     heading_run.bold = True
#     heading_run.font.size = Pt(11)

#     set_paragraph_keep_with_next(heading, keep_next=True, keep_lines=True)

#     table = doc.add_table(rows=1, cols=2)
#     table.style = "Table Grid"

#     table.rows[0].cells[0].text = "Flow"
#     table.rows[0].cells[1].text = "Description"

#     style_table_header_row(table, fill="D9D9D9")

#     for color, label in legend_items:
#         row = table.add_row().cells
#         marker = row[0].paragraphs[0].add_run("■")
#         marker.font.color.rgb = RGBColor.from_string(color.lstrip("#").upper())
#         row[1].text = clean_docx_internal_tokens(str(label))

#     set_table_no_row_split(table)
#     doc.add_paragraph("")


# def add_docx_diagram_with_legend(
#     doc: Document,
#     img_path: str,
#     legend: List[Tuple[str, str]],
# ):
#     reserved_height_in = 1.35

#     if legend:
#         reserved_height_in = min(3.6, 1.35 + (0.24 * len(legend)))

#     add_picture_fit_to_page(
#         doc,
#         img_path,
#         reserved_height_in=reserved_height_in,
#     )

#     add_flow_legend_docx(doc, legend)


# def _safe_diagram_base_name(section_key: Optional[str], diagram_index: int) -> str:
#     safe_section = re.sub(
#         r"[^a-zA-Z0-9_-]+",
#         "_",
#         str(section_key or "diagram"),
#     ).strip("_").lower()

#     if not safe_section:
#         safe_section = "diagram"

#     return f"{safe_section}_{diagram_index:02d}"


# def add_docx_diagram_or_placeholder(
#     doc: Document,
#     diagram_text: str,
#     section_key: Optional[str] = None,
#     diagram_assets_dir: Optional[Path] = None,
#     diagram_counter: Optional[Dict[str, int]] = None,
# ):
#     """
#     Render DOT/Mermaid-like diagram into DOCX.

#     Preferred path:
#       export_editable_diagram(...) -> .dot + _editable.svg + .png + .drawio

#     DOCX embeds the PNG preview and links to editable DOT/SVG/drawio.
#     """
#     clean_diagram_text = clean_docx_internal_tokens(diagram_text)

#     if diagram_counter is not None:
#         diagram_counter["value"] = diagram_counter.get("value", 0) + 1
#         diagram_index = diagram_counter["value"]
#     else:
#         diagram_index = 1

#     base_name = _safe_diagram_base_name(
#         section_key=section_key,
#         diagram_index=diagram_index,
#     )

#     def _start_diagram_page():
#         return

#     def _end_diagram_page():
#         return

#     if diagram_assets_dir is not None:
#         try:
#             assets: DiagramAssets = export_editable_diagram(
#                 diagram_text=clean_diagram_text,
#                 output_dir=diagram_assets_dir,
#                 base_name=base_name,
#                 generate_drawio=True,
#                 overwrite=True,
#             )

#             _start_diagram_page()

#             if DOCX_KEEP_DIAGRAM_SOURCE_LINKS:
#                 add_editable_diagram_links(
#                     doc,
#                     dot_path=str(assets.dot_path) if assets.dot_path else None,
#                     svg_path=str(assets.svg_path) if assets.svg_path else None,
#                     drawio_path=str(assets.drawio_path) if assets.drawio_path else None,
#                 )

#             if assets.png_path and Path(assets.png_path).exists():
#                 add_docx_diagram_with_legend(
#                     doc,
#                     str(assets.png_path),
#                     [],
#                 )
#             else:
#                 logger.error(
#                     "PNG diagram asset missing for DOCX. section_key=%s assets=%s",
#                     section_key,
#                     assets,
#                 )
#                 doc.add_paragraph(
#                     "[Diagram rendering failed] PNG preview image was not generated."
#                 )

#             _end_diagram_page()

#         except Exception:
#             logger.exception(
#                 "Failed to export/embed editable diagram assets in DOCX. section_key=%s",
#                 section_key,
#             )
#             doc.add_paragraph("[Diagram failed while embedding into DOCX]")

#         return

#     # Backward-compatible fallback: PNG only.
#     result = render_graphviz_to_png(clean_diagram_text)

#     if result:
#         img_path, legend = result

#         try:
#             _start_diagram_page()
#             add_docx_diagram_with_legend(doc, img_path, legend)
#             _end_diagram_page()

#         except Exception:
#             logger.exception(
#                 "Failed to embed diagram in DOCX. section_key=%s",
#                 section_key,
#             )
#             doc.add_paragraph("[Diagram failed while embedding into DOCX]")

#         return

#     logger.error(
#         "Diagram text detected but rendering failed. DOT source will not be printed into DOCX. section_key=%s",
#         section_key,
#     )

#     doc.add_paragraph(
#         "[Diagram rendering failed] The diagram source was detected but could "
#         "not be rendered by Graphviz. Please check renderer logs for the DOT "
#         "syntax error."
#     )


# # ---------------------------------------------------------------------
# # Recursive DOCX value renderer
# # ---------------------------------------------------------------------
# def add_docx_value(
#     doc: Document,
#     val: Any,
#     section_key: Optional[str] = None,
#     diagram_assets_dir: Optional[Path] = None,
#     diagram_counter: Optional[Dict[str, int]] = None,
# ):
#     if val is None:
#         doc.add_paragraph("N/A")
#         return

#     if isinstance(val, dict) and section_key in METADATA_SECTIONS and is_simple_metadata_dict(val):
#         add_docx_metadata_table(doc, val)
#         return

#     if isinstance(val, str):
#         clean_val = clean_docx_internal_tokens(val)
#         is_diagram_text = normalize_to_graphviz_dot(clean_val) is not None

#         if is_diagram_text:
#             add_docx_diagram_or_placeholder(
#                 doc,
#                 clean_val,
#                 section_key=section_key,
#                 diagram_assets_dir=diagram_assets_dir,
#                 diagram_counter=diagram_counter,
#             )
#             return

#         add_docx_paragraph_from_text(doc, clean_val)
#         return

#     if isinstance(val, list):
#         if not val:
#             doc.add_paragraph("N/A")
#             return

#         if all(
#             isinstance(item, str)
#             and normalize_to_graphviz_dot(clean_docx_internal_tokens(item))
#             for item in val
#         ):
#             for item in val:
#                 add_docx_diagram_or_placeholder(
#                     doc,
#                     clean_docx_internal_tokens(item),
#                     section_key=section_key,
#                     diagram_assets_dir=diagram_assets_dir,
#                     diagram_counter=diagram_counter,
#                 )
#             return

#         if all(isinstance(item, dict) for item in val):
#             add_docx_table(doc, val)
#             return

#         for item in val:
#             if isinstance(item, str) and normalize_to_graphviz_dot(clean_docx_internal_tokens(item)):
#                 add_docx_diagram_or_placeholder(
#                     doc,
#                     clean_docx_internal_tokens(item),
#                     section_key=section_key,
#                     diagram_assets_dir=diagram_assets_dir,
#                     diagram_counter=diagram_counter,
#                 )
#             else:
#                 doc.add_paragraph(
#                     clean_docx_internal_tokens(
#                         display_text_for_renderer(item, multiline=False)
#                     ),
#                     style="List Bullet",
#                 )

#         doc.add_paragraph("")
#         return

#     if isinstance(val, dict):
#         for key, child_value in val.items():
#             key_lower = str(key).lower()
#             title = str(key).replace("_", " ").title()

#             if key_lower == "issues" and _is_empty_or_na(child_value):
#                 heading = doc.add_heading(title, level=3)
#                 set_paragraph_keep_with_next(
#                     heading,
#                     keep_next=True,
#                     keep_lines=True,
#                 )
#                 add_docx_table(doc, _default_issues_rows())
#                 continue

#             heading = doc.add_heading(title, level=3)
#             set_paragraph_keep_with_next(
#                 heading,
#                 keep_next=True,
#                 keep_lines=True,
#             )

#             if key_lower == "diagrams":
#                 add_docx_value(
#                     doc,
#                     child_value,
#                     section_key=section_key,
#                     diagram_assets_dir=diagram_assets_dir,
#                     diagram_counter=diagram_counter,
#                 )
#             else:
#                 add_docx_value(
#                     doc,
#                     child_value,
#                     section_key=key,
#                     diagram_assets_dir=diagram_assets_dir,
#                     diagram_counter=diagram_counter,
#                 )

#         return

#     doc.add_paragraph(clean_docx_internal_tokens(str(val)))


# # ---------------------------------------------------------------------
# # TOC helpers
# # ---------------------------------------------------------------------
# def add_docx_toc(doc: Document):
#     paragraph = doc.add_paragraph()
#     run = paragraph.add_run()

#     begin = OxmlElement("w:fldChar")
#     begin.set(qn("w:fldCharType"), "begin")
#     run._r.append(begin)

#     instr = OxmlElement("w:instrText")
#     instr.set(qn("xml:space"), "preserve")
#     instr.text = f'TOC \\o "{DOCX_TOC_OUTLINE}" \\h \\z \\u'
#     run._r.append(instr)

#     separate = OxmlElement("w:fldChar")
#     separate.set(qn("w:fldCharType"), "separate")
#     run._r.append(separate)

#     end = OxmlElement("w:fldChar")
#     end.set(qn("w:fldCharType"), "end")
#     run._r.append(end)


# def enable_docx_update_fields(doc: Document):
#     update = OxmlElement("w:updateFields")
#     update.set(qn("w:val"), "true")
#     doc.settings._element.append(update)


# # ---------------------------------------------------------------------
# # Main DOCX builder
# # ---------------------------------------------------------------------
# def build_docx(
#     data: Dict[str, Any],
#     docx_path: str,
#     hld_model: Any,
#     selected_sections_set: Optional[set[str]] = None,
#     logo_path: Optional[str] = None,
# ):
#     doc = Document()
#     docx_output_path = Path(docx_path)

#     diagram_assets_dir = docx_output_path.parent / f"{docx_output_path.stem}_diagram_assets"
#     diagram_assets_dir.mkdir(parents=True, exist_ok=True)

#     diagram_counter = {"value": 0}

#     set_docx_default_font(
#         doc,
#         font_name=DOCX_FONT_NAME,
#         font_size=DOCX_FONT_SIZE_PT,
#     )

#     for section in doc.sections:
#         section.top_margin = Inches(DOCX_TOP_MARGIN_IN)
#         section.bottom_margin = Inches(DOCX_BOTTOM_MARGIN_IN)
#         section.left_margin = Inches(DOCX_LEFT_MARGIN_IN)
#         section.right_margin = Inches(DOCX_RIGHT_MARGIN_IN)

#     title = doc.add_heading(DOCUMENT_TITLE, level=1)
#     title.alignment = WD_ALIGN_PARAGRAPH.CENTER
#     set_paragraph_keep_with_next(title, keep_next=True, keep_lines=True)

#     if logo_path:
#         try:
#             logo_paragraph = doc.add_paragraph()
#             logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
#             logo_paragraph.paragraph_format.space_before = Pt(24)
#             logo_run = logo_paragraph.add_run()
#             logo_run.add_picture(str(logo_path), width=Inches(DOCX_LOGO_WIDTH_IN))
#         except Exception:
#             logger.warning("Logo not inserted into DOCX: %s", logo_path)

#     doc.add_page_break()
#     add_docx_toc(doc)
#     doc.add_page_break()

#     if DOCX_UPDATE_FIELDS_ON_OPEN:
#         enable_docx_update_fields(doc)

#     first = True

#     for field_name, field_info in hld_model.model_fields.items():
#         if not should_include_section(
#             field_name,
#             field_info,
#             selected_sections_set or set(),
#         ):
#             continue

#         if not first:
#             doc.add_page_break()

#         first = False

#         section_title = field_info.title or field_name.replace("_", " ").title()
#         heading = doc.add_heading(section_title, level=2)
#         set_paragraph_keep_with_next(
#             heading,
#             keep_next=True,
#             keep_lines=True,
#         )

#         add_docx_value(
#             doc,
#             data.get(field_name),
#             section_key=field_name,
#             diagram_assets_dir=diagram_assets_dir,
#             diagram_counter=diagram_counter,
#         )

#     doc.save(docx_path)

from __future__ import annotations

"""
docx_renderer.py

DOCX renderer for AIA HLD output.

Main responsibilities
---------------------
1. Build a Word document from a normalized HLD model dictionary.
2. Render standard text, metadata, nested dictionaries and list/table sections.
3. Render architecture diagrams as PNG previews inside the DOCX.
4. Generate and link editable diagram artefacts:
   - DOT source
   - SVG vector
   - draw.io editable file
   - Visio VSDX file, when enabled/supported by editable_diagram_exporter.py

Design notes
------------
- Diagram export is intentionally delegated to editable_diagram_exporter.py.
  This keeps DOCX rendering decoupled from editable artefact generation.
- Graphviz-only PNG fallback is preserved for backward compatibility.
- The DOCX renderer does not upload artefacts. The orchestrator collects files
  from generated_docs/<safe_name>_diagram_assets/ and uploads them.
"""

import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from .config import (
    DOCUMENT_TITLE,
    DOCX_FONT_NAME,
    DOCX_FONT_SIZE_PT,
    DOCX_TOP_MARGIN_IN,
    DOCX_BOTTOM_MARGIN_IN,
    DOCX_LEFT_MARGIN_IN,
    DOCX_RIGHT_MARGIN_IN,
    DOCX_LOGO_WIDTH_IN,
    DOCX_MAX_IMAGE_WIDTH_FALLBACK_IN,
    DOCX_MAX_IMAGE_HEIGHT_FALLBACK_IN,
    DOCX_TOC_OUTLINE,
    DOCX_UPDATE_FIELDS_ON_OPEN,
    DOCX_KEEP_DIAGRAM_SOURCE_LINKS,
)
from .editable_diagram_exporter import export_editable_diagram
from .graphviz_renderer import normalize_to_graphviz_dot, render_graphviz_to_png
from .schema_utils import (
    METADATA_SECTIONS,
    is_simple_metadata_dict,
    should_include_section,
)
from .text_sanitizer import display_text_for_renderer, renderer_safe_plain_text

logger = logging.getLogger(__name__)


# =============================================================================
# Internal token cleanup
# =============================================================================
def clean_docx_internal_tokens(text: Any) -> str:
    """
    Remove renderer-only tokens so they never appear in DOCX output.

    These tokens are only for PDF/HTML rendering and must not be visible in Word
    documents.
    """
    value = str(text or "")

    tokens = [
        "AIASECTIONBLOCKSTARTTOKEN",
        "AIASECTIONBLOCKENDTOKEN",
        "[[AIASECTIONBLOCKSTARTTOKEN]]",
        "[[AIASECTIONBLOCKENDTOKEN]]",
        "SECTION_BLOCK_START",
        "SECTION_BLOCK_END",
        "[[SECTION_BLOCK_START]]",
        "[[SECTION_BLOCK_END]]",
        "[[SECTIONBLOCKSTART]]",
        "[[SECTIONBLOCKEND]]",
    ]

    for token in tokens:
        value = value.replace(token, "")

    return value.strip()


def _is_empty_or_na(value: Any) -> bool:
    if value is None:
        return True

    if isinstance(value, str):
        return value.strip().lower() in {
            "",
            "n/a",
            "na",
            "none",
            "null",
            "not applicable",
        }

    if isinstance(value, list):
        return len(value) == 0

    if isinstance(value, dict):
        return len(value) == 0

    return False


def _default_issues_rows() -> List[Dict[str, str]]:
    return [
        {
            "id": "I001",
            "description": "TBC - No issues have been identified at draft stage.",
            "status": "Open",
            "mitigation": "TBC",
        }
    ]


# =============================================================================
# DOCX base helpers
# =============================================================================
def set_docx_default_font(
    doc: Document,
    font_name: Optional[str] = None,
    font_size: Optional[int] = None,
) -> None:
    resolved_font_name = font_name or DOCX_FONT_NAME
    resolved_font_size = font_size or DOCX_FONT_SIZE_PT

    if "Normal" in doc.styles:
        doc.styles["Normal"].font.name = resolved_font_name
        doc.styles["Normal"].font.size = Pt(resolved_font_size)
        doc.styles["Normal"]._element.rPr.rFonts.set(
            qn("w:eastAsia"),
            resolved_font_name,
        )


def add_docx_page_break(doc: Document) -> None:
    doc.add_page_break()


def set_paragraph_keep_with_next(
    paragraph,
    keep_next: bool = True,
    keep_lines: bool = True,
) -> None:
    pPr = paragraph._p.get_or_add_pPr()

    if keep_next:
        keep_next_el = OxmlElement("w:keepNext")
        pPr.append(keep_next_el)

    if keep_lines:
        keep_lines_el = OxmlElement("w:keepLines")
        pPr.append(keep_lines_el)


def shade_cell(cell, fill: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_table_header_row(table, fill: str = "D9D9D9") -> None:
    if not table.rows:
        return

    header_row = table.rows[0]

    for cell in header_row.cells:
        shade_cell(cell, fill=fill)

        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            for run in paragraph.runs:
                run.bold = True


def set_table_no_row_split(table) -> None:
    """
    Prevent table rows from splitting across pages where Word can honour it.
    Also repeats the first row as header row.
    """
    if table.rows:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trPr.append(hdr)

    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        cant.set(qn("w:val"), "true")
        trPr.append(cant)


def get_docx_max_image_size(
    doc: Document,
    reserved_height_in: float = 0.0,
) -> Tuple[float, float]:
    try:
        section = doc.sections[-1]

        usable_width = max(
            4.5,
            section.page_width.inches
            - section.left_margin.inches
            - section.right_margin.inches,
        )

        usable_height = max(
            3.8,
            section.page_height.inches
            - section.top_margin.inches
            - section.bottom_margin.inches
            - reserved_height_in,
        )

        return usable_width, usable_height

    except Exception:
        return (
            DOCX_MAX_IMAGE_WIDTH_FALLBACK_IN,
            max(3.8, DOCX_MAX_IMAGE_HEIGHT_FALLBACK_IN - reserved_height_in),
        )


def add_picture_fit_to_page(
    doc: Document,
    image_path: str,
    reserved_height_in: float = 0.0,
) -> None:
    max_width_in, max_height_in = get_docx_max_image_size(
        doc,
        reserved_height_in=reserved_height_in,
    )

    width_in = max_width_in
    height_in = None

    try:
        with Image.open(image_path) as img:
            if img.width > 0 and img.height > 0:
                aspect = img.height / img.width
                predicted_height = width_in * aspect

                if predicted_height > max_height_in:
                    height_in = max_height_in
                    width_in = height_in / aspect

    except Exception:
        logger.warning(
            "Could not inspect image size for %s; using width-only fit",
            image_path,
        )

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()

    if height_in is not None:
        run.add_picture(
            str(image_path),
            width=Inches(width_in),
            height=Inches(height_in),
        )
    else:
        run.add_picture(str(image_path), width=Inches(width_in))

    set_paragraph_keep_with_next(
        paragraph,
        keep_next=True,
        keep_lines=True,
    )


# =============================================================================
# DOCX hyperlink helpers
# =============================================================================
def add_docx_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part

    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    run.append(run_properties)

    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_link_separator_if_needed(paragraph, added_any: bool) -> None:
    if added_any:
        paragraph.add_run(" | ")


def add_editable_diagram_links(
    doc: Document,
    dot_path: Optional[str] = None,
    svg_path: Optional[str] = None,
    drawio_path: Optional[str] = None,
    vsdx_path: Optional[str] = None,
) -> None:
    """
    Add links above the diagram preview for editable diagram assets.

    Links:
      - DOT source
      - SVG vector
      - draw.io editable diagram
      - Visio VSDX, if generated
    """
    if not dot_path and not svg_path and not drawio_path and not vsdx_path:
        return

    paragraph = doc.add_paragraph()

    label_run = paragraph.add_run("Edit this diagram: ")
    label_run.bold = True

    added_any = False

    if dot_path and Path(dot_path).exists():
        add_docx_hyperlink(paragraph, "DOT source", Path(dot_path).resolve().as_uri())
        added_any = True

    if svg_path and Path(svg_path).exists():
        _add_link_separator_if_needed(paragraph, added_any)
        add_docx_hyperlink(paragraph, "SVG vector", Path(svg_path).resolve().as_uri())
        added_any = True

    if drawio_path and Path(drawio_path).exists():
        _add_link_separator_if_needed(paragraph, added_any)
        add_docx_hyperlink(
            paragraph,
            "draw.io editable",
            Path(drawio_path).resolve().as_uri(),
        )
        added_any = True

    if vsdx_path and Path(vsdx_path).exists():
        _add_link_separator_if_needed(paragraph, added_any)
        add_docx_hyperlink(
            paragraph,
            "Visio VSDX",
            Path(vsdx_path).resolve().as_uri(),
        )

    set_paragraph_keep_with_next(
        paragraph,
        keep_next=True,
        keep_lines=True,
    )


# =============================================================================
# Text/table rendering
# =============================================================================
def add_docx_paragraph_from_text(doc: Document, text: str) -> None:
    clean_text = clean_docx_internal_tokens(text)

    if not clean_text:
        doc.add_paragraph("N/A")
        return

    safe_text = renderer_safe_plain_text(clean_text)
    paragraph = doc.add_paragraph()

    for idx, line in enumerate(safe_text.split("\n")):
        if idx > 0:
            paragraph.add_run().add_break()
        paragraph.add_run(line)


def add_docx_table(doc: Document, items: List[Dict[str, Any]]) -> None:
    if not items:
        doc.add_paragraph("N/A")
        return

    all_keys: List[str] = []

    for row in items:
        for key in row.keys():
            if key not in all_keys:
                all_keys.append(key)

    table = doc.add_table(rows=1, cols=len(all_keys))
    table.style = "Table Grid"

    for idx, key in enumerate(all_keys):
        table.rows[0].cells[idx].text = clean_docx_internal_tokens(
            key.replace("_", " ").title()
        )

    style_table_header_row(table, fill="D9D9D9")

    for row in items:
        cells = table.add_row().cells

        for idx, key in enumerate(all_keys):
            cells[idx].text = clean_docx_internal_tokens(
                display_text_for_renderer(row.get(key, ""), multiline=True)
            )

    set_table_no_row_split(table)
    doc.add_paragraph("")


def add_docx_metadata_table(doc: Document, data: Dict[str, Any]) -> None:
    if not data:
        doc.add_paragraph("N/A")
        return

    add_docx_table(doc, [data])


# =============================================================================
# Diagram rendering - small decoupled helpers
# =============================================================================
def add_flow_legend_docx(doc: Document, legend_items: List[Tuple[str, str]]) -> None:
    if not legend_items:
        return

    heading = doc.add_paragraph()
    heading_run = heading.add_run("Flow Legend")
    heading_run.bold = True
    heading_run.font.size = Pt(11)

    set_paragraph_keep_with_next(heading, keep_next=True, keep_lines=True)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    table.rows[0].cells[0].text = "Flow"
    table.rows[0].cells[1].text = "Description"

    style_table_header_row(table, fill="D9D9D9")

    for color, label in legend_items:
        row = table.add_row().cells
        marker = row[0].paragraphs[0].add_run("■")
        marker.font.color.rgb = RGBColor.from_string(color.lstrip("#").upper())
        row[1].text = clean_docx_internal_tokens(str(label))

    set_table_no_row_split(table)
    doc.add_paragraph("")


def add_docx_diagram_with_legend(
    doc: Document,
    img_path: str,
    legend: List[Tuple[str, str]],
) -> None:
    reserved_height_in = 1.35

    if legend:
        reserved_height_in = min(3.6, 1.35 + (0.24 * len(legend)))

    add_picture_fit_to_page(
        doc,
        img_path,
        reserved_height_in=reserved_height_in,
    )

    add_flow_legend_docx(doc, legend)


def _safe_diagram_base_name(section_key: Optional[str], diagram_index: int) -> str:
    safe_section = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        str(section_key or "diagram"),
    ).strip("_").lower()

    if not safe_section:
        safe_section = "diagram"

    return f"{safe_section}_{diagram_index:02d}"


def _next_diagram_index(diagram_counter: Optional[Dict[str, int]]) -> int:
    if diagram_counter is None:
        return 1

    diagram_counter["value"] = diagram_counter.get("value", 0) + 1
    return diagram_counter["value"]


def _export_diagram_assets(
    diagram_text: str,
    diagram_assets_dir: Path,
    base_name: str,
):
    """
    Delegate editable artefact generation to editable_diagram_exporter.py.

    Kept as a small wrapper so VSDX/draw.io behaviour can be changed in one
    place without changing recursive DOCX rendering code.
    """
    return export_editable_diagram(
        diagram_text=diagram_text,
        output_dir=diagram_assets_dir,
        base_name=base_name,
        generate_drawio=True,
        generate_vsdx=True,
        overwrite=True,
    )


def _add_asset_links_if_enabled(doc: Document, assets: Any) -> None:
    if not DOCX_KEEP_DIAGRAM_SOURCE_LINKS:
        return

    add_editable_diagram_links(
        doc,
        dot_path=str(assets.dot_path) if getattr(assets, "dot_path", None) else None,
        svg_path=str(assets.svg_path) if getattr(assets, "svg_path", None) else None,
        drawio_path=str(assets.drawio_path) if getattr(assets, "drawio_path", None) else None,
        vsdx_path=str(assets.vsdx_path) if getattr(assets, "vsdx_path", None) else None,
    )


def _embed_asset_png_or_placeholder(
    doc: Document,
    assets: Any,
    section_key: Optional[str],
) -> None:
    png_path = getattr(assets, "png_path", None)

    if png_path and Path(png_path).exists():
        add_docx_diagram_with_legend(
            doc,
            str(png_path),
            [],
        )
        return

    logger.error(
        "PNG diagram asset missing for DOCX. section_key=%s assets=%s",
        section_key,
        assets,
    )
    doc.add_paragraph("[Diagram rendering failed] PNG preview image was not generated.")


def _render_editable_diagram_path(
    doc: Document,
    clean_diagram_text: str,
    section_key: Optional[str],
    diagram_assets_dir: Path,
    base_name: str,
) -> None:
    """Render diagram through editable asset path: DOT/SVG/PNG/draw.io/VSDX."""
    assets = _export_diagram_assets(
        diagram_text=clean_diagram_text,
        diagram_assets_dir=diagram_assets_dir,
        base_name=base_name,
    )

    _add_asset_links_if_enabled(doc, assets)
    _embed_asset_png_or_placeholder(doc, assets, section_key)


def _render_png_fallback_path(
    doc: Document,
    clean_diagram_text: str,
    section_key: Optional[str],
) -> bool:
    """
    Backward-compatible path for environments where editable asset directory is
    not provided. Returns True when a diagram was rendered.
    """
    result = render_graphviz_to_png(clean_diagram_text)

    if not result:
        return False

    img_path, legend = result

    try:
        add_docx_diagram_with_legend(doc, img_path, legend)
        return True
    except Exception:
        logger.exception(
            "Failed to embed diagram in DOCX. section_key=%s",
            section_key,
        )
        doc.add_paragraph("[Diagram failed while embedding into DOCX]")
        return True


def add_docx_diagram_or_placeholder(
    doc: Document,
    diagram_text: str,
    section_key: Optional[str] = None,
    diagram_assets_dir: Optional[Path] = None,
    diagram_counter: Optional[Dict[str, int]] = None,
) -> None:
    """
    Render DOT/Mermaid-like diagram into DOCX.

    Preferred path:
      export_editable_diagram(...) -> .dot + _editable.svg + .png + .drawio + .vsdx

    DOCX embeds the PNG preview and links to editable DOT/SVG/draw.io/VSDX.
    """
    clean_diagram_text = clean_docx_internal_tokens(diagram_text)
    diagram_index = _next_diagram_index(diagram_counter)
    base_name = _safe_diagram_base_name(section_key=section_key, diagram_index=diagram_index)

    if diagram_assets_dir is not None:
        try:
            _render_editable_diagram_path(
                doc=doc,
                clean_diagram_text=clean_diagram_text,
                section_key=section_key,
                diagram_assets_dir=diagram_assets_dir,
                base_name=base_name,
            )
        except Exception:
            logger.exception(
                "Failed to export/embed editable diagram assets in DOCX. section_key=%s",
                section_key,
            )
            doc.add_paragraph("[Diagram failed while embedding into DOCX]")
        return

    if _render_png_fallback_path(doc, clean_diagram_text, section_key):
        return

    logger.error(
        "Diagram text detected but rendering failed. DOT source will not be printed into DOCX. section_key=%s",
        section_key,
    )

    doc.add_paragraph(
        "[Diagram rendering failed] The diagram source was detected but could "
        "not be rendered by Graphviz. Please check renderer logs for the DOT "
        "syntax error."
    )


# =============================================================================
# Recursive DOCX value renderer
# =============================================================================
def _is_diagram_text(value: str) -> bool:
    return normalize_to_graphviz_dot(clean_docx_internal_tokens(value)) is not None


def add_docx_value(
    doc: Document,
    val: Any,
    section_key: Optional[str] = None,
    diagram_assets_dir: Optional[Path] = None,
    diagram_counter: Optional[Dict[str, int]] = None,
) -> None:
    if val is None:
        doc.add_paragraph("N/A")
        return

    if isinstance(val, dict) and section_key in METADATA_SECTIONS and is_simple_metadata_dict(val):
        add_docx_metadata_table(doc, val)
        return

    if isinstance(val, str):
        clean_val = clean_docx_internal_tokens(val)

        if _is_diagram_text(clean_val):
            add_docx_diagram_or_placeholder(
                doc,
                clean_val,
                section_key=section_key,
                diagram_assets_dir=diagram_assets_dir,
                diagram_counter=diagram_counter,
            )
            return

        add_docx_paragraph_from_text(doc, clean_val)
        return

    if isinstance(val, list):
        if not val:
            doc.add_paragraph("N/A")
            return

        if all(isinstance(item, str) and _is_diagram_text(item) for item in val):
            for item in val:
                add_docx_diagram_or_placeholder(
                    doc,
                    clean_docx_internal_tokens(item),
                    section_key=section_key,
                    diagram_assets_dir=diagram_assets_dir,
                    diagram_counter=diagram_counter,
                )
            return

        if all(isinstance(item, dict) for item in val):
            add_docx_table(doc, val)
            return

        for item in val:
            if isinstance(item, str) and _is_diagram_text(item):
                add_docx_diagram_or_placeholder(
                    doc,
                    clean_docx_internal_tokens(item),
                    section_key=section_key,
                    diagram_assets_dir=diagram_assets_dir,
                    diagram_counter=diagram_counter,
                )
            else:
                doc.add_paragraph(
                    clean_docx_internal_tokens(
                        display_text_for_renderer(item, multiline=False)
                    ),
                    style="List Bullet",
                )

        doc.add_paragraph("")
        return

    if isinstance(val, dict):
        for key, child_value in val.items():
            key_lower = str(key).lower()
            title = str(key).replace("_", " ").title()

            if key_lower == "issues" and _is_empty_or_na(child_value):
                heading = doc.add_heading(title, level=3)
                set_paragraph_keep_with_next(
                    heading,
                    keep_next=True,
                    keep_lines=True,
                )
                add_docx_table(doc, _default_issues_rows())
                continue

            heading = doc.add_heading(title, level=3)
            set_paragraph_keep_with_next(
                heading,
                keep_next=True,
                keep_lines=True,
            )

            if key_lower == "diagrams":
                add_docx_value(
                    doc,
                    child_value,
                    section_key=section_key,
                    diagram_assets_dir=diagram_assets_dir,
                    diagram_counter=diagram_counter,
                )
            else:
                add_docx_value(
                    doc,
                    child_value,
                    section_key=key,
                    diagram_assets_dir=diagram_assets_dir,
                    diagram_counter=diagram_counter,
                )

        return

    doc.add_paragraph(clean_docx_internal_tokens(str(val)))


# =============================================================================
# TOC helpers
# =============================================================================
def add_docx_toc(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f'TOC \\o "{DOCX_TOC_OUTLINE}" \\h \\z \\u'
    run._r.append(instr)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def enable_docx_update_fields(doc: Document) -> None:
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings._element.append(update)


# =============================================================================
# Main DOCX builder
# =============================================================================
def _prepare_docx_document() -> Document:
    doc = Document()

    set_docx_default_font(
        doc,
        font_name=DOCX_FONT_NAME,
        font_size=DOCX_FONT_SIZE_PT,
    )

    for section in doc.sections:
        section.top_margin = Inches(DOCX_TOP_MARGIN_IN)
        section.bottom_margin = Inches(DOCX_BOTTOM_MARGIN_IN)
        section.left_margin = Inches(DOCX_LEFT_MARGIN_IN)
        section.right_margin = Inches(DOCX_RIGHT_MARGIN_IN)

    return doc


def _add_title_and_logo(doc: Document, logo_path: Optional[str]) -> None:
    title = doc.add_heading(DOCUMENT_TITLE, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_keep_with_next(title, keep_next=True, keep_lines=True)

    if not logo_path:
        return

    try:
        logo_paragraph = doc.add_paragraph()
        logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_paragraph.paragraph_format.space_before = Pt(24)
        logo_run = logo_paragraph.add_run()
        logo_run.add_picture(str(logo_path), width=Inches(DOCX_LOGO_WIDTH_IN))
    except Exception:
        logger.warning("Logo not inserted into DOCX: %s", logo_path)


def _add_toc_page(doc: Document) -> None:
    doc.add_page_break()
    add_docx_toc(doc)
    doc.add_page_break()

    if DOCX_UPDATE_FIELDS_ON_OPEN:
        enable_docx_update_fields(doc)


def _render_selected_hld_sections(
    doc: Document,
    data: Dict[str, Any],
    hld_model: Any,
    selected_sections_set: Optional[set[str]],
    diagram_assets_dir: Path,
    diagram_counter: Dict[str, int],
) -> None:
    first = True

    for field_name, field_info in hld_model.model_fields.items():
        if not should_include_section(
            field_name,
            field_info,
            selected_sections_set or set(),
        ):
            continue

        if not first:
            doc.add_page_break()

        first = False

        section_title = field_info.title or field_name.replace("_", " ").title()
        heading = doc.add_heading(section_title, level=2)
        set_paragraph_keep_with_next(
            heading,
            keep_next=True,
            keep_lines=True,
        )

        add_docx_value(
            doc,
            data.get(field_name),
            section_key=field_name,
            diagram_assets_dir=diagram_assets_dir,
            diagram_counter=diagram_counter,
        )


def build_docx(
    data: Dict[str, Any],
    docx_path: str,
    hld_model: Any,
    selected_sections_set: Optional[set[str]] = None,
    logo_path: Optional[str] = None,
) -> None:
    """
    Build DOCX document.

    Diagram artefacts are written to:
        <docx_parent>/<docx_stem>_diagram_assets/

    This folder is later collected by orchestrator.py for upload/signing.
    """
    docx_output_path = Path(docx_path)
    diagram_assets_dir = docx_output_path.parent / f"{docx_output_path.stem}_diagram_assets"
    diagram_assets_dir.mkdir(parents=True, exist_ok=True)

    diagram_counter = {"value": 0}

    doc = _prepare_docx_document()
    _add_title_and_logo(doc, logo_path)
    _add_toc_page(doc)

    _render_selected_hld_sections(
        doc=doc,
        data=data,
        hld_model=hld_model,
        selected_sections_set=selected_sections_set,
        diagram_assets_dir=diagram_assets_dir,
        diagram_counter=diagram_counter,
    )

    doc.save(docx_path)
