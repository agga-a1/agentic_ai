
# from __future__ import annotations

# import html
# import logging
# import re
# import shlex
# import subprocess
# import tempfile
# from collections import defaultdict, deque
# from pathlib import Path
# from typing import Any, Dict, List, Optional, Set, Tuple

# from pptx import Presentation
# from pptx.dml.color import RGBColor
# from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
# from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
# from pptx.shapes.base import BaseShape
# from pptx.util import Inches, Pt

# from .config import ICON_DIR
# from .graphviz_renderer import normalize_to_graphviz_dot, render_graphviz_to_assets
# from .schema_utils import should_include_section
# from .text_sanitizer import renderer_safe_plain_text

# logger = logging.getLogger(__name__)

# # ---------------------------------------------------------------------
# # PPTX visual defaults
# # ---------------------------------------------------------------------
# SLIDE_WIDTH_IN = 13.333
# SLIDE_HEIGHT_IN = 7.5

# TITLE_LEFT_IN = 0.35
# TITLE_TOP_IN = 0.14
# TITLE_WIDTH_IN = 12.6
# TITLE_HEIGHT_IN = 0.42

# CANVAS_LEFT_IN = 0.35
# CANVAS_TOP_IN = 0.72
# CANVAS_WIDTH_IN = 12.3
# CANVAS_HEIGHT_IN = 5.42

# LEGEND_LEFT_IN = 0.35
# LEGEND_TOP_IN = 6.25
# LEGEND_WIDTH_IN = 12.3
# LEGEND_MAX_HEIGHT_IN = 1.05

# NODE_MIN_WIDTH_IN = 1.15
# NODE_MIN_HEIGHT_IN = 0.52
# NODE_ICON_SIZE_IN = 0.26
# NODE_INNER_PADDING_IN = 0.08
# NODE_TEXT_LEFT_WITH_ICON_IN = 0.39

# FALLBACK_NODE_WIDTH_IN = 2.15
# FALLBACK_NODE_HEIGHT_IN = 0.92
# FALLBACK_H_GAP_IN = 0.55
# FALLBACK_V_GAP_IN = 0.32
# FALLBACK_CLUSTER_PAD_IN = 0.18
# FALLBACK_CLUSTER_TITLE_H_IN = 0.26

# PREVIEW_MAX_WIDTH_IN = 12.0
# PREVIEW_MAX_HEIGHT_IN = 5.5

# DEFAULT_NODE_FILL = "F8F9FA"
# DEFAULT_NODE_BORDER = "DADCE0"
# DEFAULT_NODE_FONT = "202124"

# DEFAULT_CLUSTER_FILL = "F3F8FF"
# DEFAULT_CLUSTER_BORDER = "A8C7FA"
# DEFAULT_CLUSTER_FONT = "174EA6"

# DEFAULT_EDGE_COLOR = "5F6368"
# DEFAULT_EDGE_WIDTH_PT = 1.15

# LEGEND_BORDER = "DADCE0"
# LEGEND_FILL = "FFFFFF"

# SUPPORTED_PICTURE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif"}
# PREFERRED_ICON_EXTS = [".png", ".jpg", ".jpeg", ".bmp", ".gif"]

# COMMENT_LINE_RE = re.compile(r"//.*?$", re.MULTILINE)
# BLOCK_COMMENT_RE = re.compile(r"/\*.*?\*/", re.DOTALL)
# EDGE_RE = re.compile(r'("[^"]+"|[A-Za-z0-9_.:-]+)\s*(->|--)\s*("[^"]+"|[A-Za-z0-9_.:-]+)')

# COLOR_NAME_MAP = {
#     "red": "E60000",
#     "blue": "1D70B8",
#     "green": "28A197",
#     "orange": "F47738",
#     "purple": "4C2C92",
#     "gray": "5F6368",
#     "grey": "5F6368",
#     "black": "202124",
#     "lightgrey": "DADCE0",
#     "lightgray": "DADCE0",
# }

# ICON_KEYWORD_MAP = [
#     ("cloud storage", "cloud_storage"),
#     ("gcs", "cloud_storage"),
#     ("storage", "cloud_storage"),
#     ("cloud function", "cloud_functions"),
#     ("functions", "cloud_functions"),
#     ("function", "cloud_functions"),
#     ("vault", "secret_manager"),
#     ("hashicorp", "secret_manager"),
#     ("secret", "secret_manager"),
#     ("kms", "key_management_service"),
#     ("key management", "key_management_service"),
#     ("pub/sub", "pubsub"),
#     ("pubsub", "pubsub"),
#     ("eventarc", "eventarc"),
#     ("logging", "cloud_logging"),
#     ("monitoring", "cloud_monitoring"),
#     ("cloud build", "cloud_build"),
#     ("build", "cloud_build"),
#     ("terraform", "terraform"),
#     ("teradata", "database_migration_service"),
#     ("database", "database_migration_service"),
#     ("managed file transfer", "transfer"),
#     ("mft", "transfer"),
#     ("transfer", "transfer"),
#     ("collibra", "catalog"),
#     ("catalog", "catalog"),
# ]

# # ---------------------------------------------------------------------
# # Text / color helpers
# # ---------------------------------------------------------------------
# def _clean_internal_tokens(text: Any) -> str:
#     value = str(text or "")
#     for token in [
#         "AIASECTIONBLOCKSTARTTOKEN", "AIASECTIONBLOCKENDTOKEN",
#         "[[AIASECTIONBLOCKSTARTTOKEN]]", "[[AIASECTIONBLOCKENDTOKEN]]",
#         "SECTION_BLOCK_START", "SECTION_BLOCK_END",
#         "[[SECTION_BLOCK_START]]", "[[SECTION_BLOCK_END]]",
#         "[[SECTIONBLOCKSTART]]", "[[SECTIONBLOCKEND]]",
#     ]:
#         value = value.replace(token, "")
#     return html.unescape(value.strip())


# def _safe_title(text: Any) -> str:
#     return renderer_safe_plain_text(_clean_internal_tokens(text)).strip() or "Diagram"


# def _safe_name(text: str) -> str:
#     return re.sub(r"[^a-zA-Z0-9_-]+", "_", str(text or "")).strip("_").lower() or "diagram"


# def _clean_label(value: Any) -> str:
#     text = html.unescape(str(value or "").strip()).strip('"')
#     text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
#     text = re.sub(r"<[^>]+>", " ", text)
#     text = text.replace("\\N", "\n").replace("\\n", "\n").replace("&nbsp;", " ")
#     text = re.sub(r"\s+\n", "\n", text)
#     text = re.sub(r"\n\s+", "\n", text)
#     text = re.sub(r"[ \t]+", " ", text)
#     return text.strip()


# def _label(raw: Optional[str], fallback: str) -> str:
#     label = _clean_label(raw or "")
#     if not label:
#         label = str(fallback or "").replace("_", " ").replace("-", " ").title()
#     return label[:180]


# def _safe_color(value: Optional[str], default: str = DEFAULT_EDGE_COLOR) -> str:
#     raw = str(value or "").strip().strip('"').strip().lstrip("#")
#     if not raw:
#         return default
#     lowered = raw.lower()
#     if lowered in COLOR_NAME_MAP:
#         return COLOR_NAME_MAP[lowered]
#     raw = raw.upper()
#     if len(raw) == 6 and all(ch in "0123456789ABCDEF" for ch in raw):
#         return raw
#     return default


# def _rgb(hex_color: str, default: str = "000000") -> RGBColor:
#     return RGBColor.from_string(_safe_color(hex_color, default))


# def _is_pseudo_node(node_id: str, label: Optional[str] = None) -> bool:
#     text = _clean_label(label if label is not None else node_id).lower().strip()
#     raw = str(node_id or "").lower().strip()
#     if text in {"", "n/a", "na", "none", "null"}:
#         return True
#     if raw in {"node", "edge", "graph", "rank", "label", "style", "color", "fillcolor", "rankdir"}:
#         return True
#     if text.startswith(("label=", "style=", "color=", "fillcolor=", "rankdir=", "fontsize=", "fontname=", "margin=", "pad=")):
#         return True
#     if raw.startswith(("label=", "style=", "color=", "fillcolor=", "rankdir=")):
#         return True
#     if text in {"filled", "lightgrey", "lightgray", "blue", "green", "purple", "orange", "red", "gray", "grey"}:
#         return True
#     return False

# # ---------------------------------------------------------------------
# # Icon resolution
# # ---------------------------------------------------------------------
# def _try_existing_icon_resolver(icon_hint: str) -> Optional[Path]:
#     try:
#         from . import icon_resolver as ir
#     except Exception:
#         return None

#     candidate_fns = ["resolve_icon_path", "resolve_icon", "get_icon_path", "find_icon_path"]
#     for fn_name in candidate_fns:
#         fn = getattr(ir, fn_name, None)
#         if not callable(fn):
#             continue
#         for args in [(icon_hint,), (icon_hint, True), (icon_hint, None)]:
#             try:
#                 result = fn(*args)
#                 if result:
#                     path = Path(str(result))
#                     if path.exists():
#                         return path
#             except TypeError:
#                 continue
#             except Exception:
#                 logger.debug("Existing icon resolver failed: %s", fn_name, exc_info=True)
#     return None


# def _search_icon_dir(icon_hint: str) -> Optional[Path]:
#     if not icon_hint:
#         return None
#     hint = _safe_name(icon_hint)
#     raw = str(icon_hint).strip()
#     bases = {
#         hint,
#         raw,
#         raw.lower(),
#         raw.replace(" ", "_").lower(),
#         raw.replace("-", "_").lower(),
#         raw.replace("/", "_").lower(),
#     }
#     for base in bases:
#         if not base:
#             continue
#         for ext in [".png", ".jpg", ".jpeg", ".bmp", ".gif", ".svg"]:
#             candidate = ICON_DIR / f"{base}{ext}"
#             if candidate.exists():
#                 return candidate

#     try:
#         for path in ICON_DIR.rglob("*"):
#             if path.is_file() and hint in path.stem.lower():
#                 return path
#     except Exception:
#         logger.debug("Icon directory search failed.", exc_info=True)
#     return None


# def _keyword_icon_hint(label_or_id: str) -> Optional[str]:
#     text = _clean_label(label_or_id).lower()
#     for keyword, icon_hint in ICON_KEYWORD_MAP:
#         if keyword in text:
#             return icon_hint
#     return None


# def _preferred_picture_icon_path(icon_hint: str) -> Optional[Path]:
#     candidates = []
#     if icon_hint:
#         candidates.append(icon_hint)
#         keyword_hint = _keyword_icon_hint(icon_hint)
#         if keyword_hint:
#             candidates.insert(0, keyword_hint)

#     for candidate in candidates:
#         path = _try_existing_icon_resolver(candidate) or _search_icon_dir(candidate)
#         if not path:
#             continue
#         if path.suffix.lower() in SUPPORTED_PICTURE_EXTS:
#             return path
#         if path.suffix.lower() == ".svg":
#             for ext in PREFERRED_ICON_EXTS:
#                 sibling = path.with_suffix(ext)
#                 if sibling.exists():
#                     return sibling
#     return None


# def _guess_icon_hint(node_id: str, node_label: str, attrs: Dict[str, str]) -> str:
#     image_attr = attrs.get("image") or attrs.get("icon") or attrs.get("imagepath")
#     if image_attr:
#         stem = Path(str(image_attr)).stem
#         if stem:
#             return stem
#     keyword = _keyword_icon_hint(node_label) or _keyword_icon_hint(node_id)
#     if keyword:
#         return keyword
#     for candidate in [node_id, node_label.split("\n", 1)[0], node_label.replace("\n", " ")]:
#         cleaned = _safe_name(candidate)
#         if cleaned:
#             return cleaned
#     return _safe_name(node_id or node_label or "node")

# # ---------------------------------------------------------------------
# # DOT parsing helpers
# # ---------------------------------------------------------------------
# def _strip_dot_comments(dot_text: str) -> str:
#     text = BLOCK_COMMENT_RE.sub("", dot_text or "")
#     text = COMMENT_LINE_RE.sub("", text)
#     return text


# def _split_statements(dot_text: str) -> List[str]:
#     statements: List[str] = []
#     buf = []
#     bracket_depth = 0
#     brace_depth = 0
#     angle_depth = 0
#     for ch in dot_text or "":
#         if ch == "[":
#             bracket_depth += 1
#         elif ch == "]":
#             bracket_depth = max(0, bracket_depth - 1)
#         elif ch == "{":
#             brace_depth += 1
#         elif ch == "}":
#             brace_depth = max(0, brace_depth - 1)
#         elif ch == "<":
#             angle_depth += 1
#         elif ch == ">":
#             angle_depth = max(0, angle_depth - 1)
#         if ch == ";" and bracket_depth == 0 and brace_depth == 0 and angle_depth == 0:
#             statement = "".join(buf).strip()
#             if statement:
#                 statements.append(statement)
#             buf = []
#         else:
#             buf.append(ch)
#     tail = "".join(buf).strip()
#     if tail:
#         statements.append(tail)
#     return statements


# def _extract_attr_pairs(attr_text: str) -> Dict[str, str]:
#     attrs: Dict[str, str] = {}
#     if not attr_text:
#         return attrs
#     text = html.unescape(str(attr_text))
#     # permissive parser for key="value", key=<value>, key=value
#     for match in re.finditer(r'([A-Za-z0-9_:-]+)\s*=\s*("[^"]*"|<[^>]*>|[^,\n\r;]+)', text):
#         key = match.group(1).lower().strip()
#         value = match.group(2).strip().strip('"')
#         if value.startswith("<") and value.endswith(">"):
#             value = value[1:-1]
#         attrs[key] = value.strip()
#     return attrs


# def _extract_bracket_attr(statement: str) -> Tuple[str, Dict[str, str]]:
#     statement = html.unescape((statement or "").strip().rstrip(";").strip())
#     if "[" not in statement or "]" not in statement:
#         return statement, {}
#     start = statement.find("[")
#     end = statement.rfind("]")
#     if end <= start:
#         return statement, {}
#     return statement[:start].strip(), _extract_attr_pairs(statement[start + 1 : end])


# def _unquote_id(value: str) -> str:
#     return str(value or "").strip().strip('"').strip()


# def _find_subgraph_blocks(dot_text: str) -> List[Tuple[str, str]]:
#     text = dot_text or ""
#     results: List[Tuple[str, str]] = []
#     idx = 0
#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             break
#         start = idx + match.start()
#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break
#         header = text[start:brace_start].strip()
#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1
#         if end >= len(text):
#             break
#         results.append((header, text[brace_start + 1 : end]))
#         idx = end + 1
#     return results


# def _remove_subgraph_blocks(dot_text: str) -> str:
#     text = dot_text or ""
#     pieces = []
#     idx = 0
#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             pieces.append(text[idx:])
#             break
#         start = idx + match.start()
#         pieces.append(text[idx:start])
#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break
#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1
#         if end >= len(text):
#             break
#         idx = end + 1
#     return "".join(pieces)


# def _extract_cluster_label(body: str, fallback_name: str) -> str:
#     for stmt in _split_statements(body):
#         stmt = stmt.strip()
#         if stmt.lower().startswith("label"):
#             if "=" in stmt:
#                 try:
#                     val = stmt.split("=", 1)[1].strip().strip('"').strip()
#                     if val:
#                         return _clean_label(val)
#                 except Exception:
#                     pass
#             _, attrs = _extract_bracket_attr(stmt)
#             if attrs.get("label"):
#                 return _clean_label(attrs["label"])
#     return fallback_name


# def _extract_cluster_nodes(body: str) -> Set[str]:
#     node_ids: Set[str] = set()
#     for stmt in _split_statements(body):
#         s = stmt.strip()
#         if not s or "->" in s or "--" in s:
#             continue
#         if s.lower().startswith(("node ", "edge ", "graph ", "rank ", "label", "style", "color", "fillcolor")):
#             continue
#         prefix, attrs = _extract_bracket_attr(s)
#         node_id = _unquote_id(prefix)
#         label = attrs.get("label", node_id)
#         if node_id and not _is_pseudo_node(node_id, label) and re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
#             node_ids.add(node_id)
#     for match in EDGE_RE.finditer(body or ""):
#         for group_idx in [1, 3]:
#             node_id = _unquote_id(match.group(group_idx))
#             if node_id and not _is_pseudo_node(node_id):
#                 node_ids.add(node_id)
#     return node_ids


# def _parse_clusters_from_dot(dot: str) -> List[Dict[str, Any]]:
#     clusters: List[Dict[str, Any]] = []
#     for header, body in _find_subgraph_blocks(dot):
#         fallback_name = _unquote_id(header.replace("subgraph", "", 1).strip()) or "Cluster"
#         cluster_label = _extract_cluster_label(body, fallback_name)
#         cluster_nodes = sorted(_extract_cluster_nodes(body))
#         if cluster_nodes and not _is_pseudo_node(fallback_name, cluster_label):
#             clusters.append({"id": _safe_name(fallback_name), "label": cluster_label, "nodes": cluster_nodes})
#     return clusters


# def _parse_node_attrs_from_dot(dot: str) -> Dict[str, Dict[str, str]]:
#     node_attrs: Dict[str, Dict[str, str]] = {}
#     for stmt in _split_statements(dot):
#         s = stmt.strip()
#         if not s or "->" in s or "--" in s:
#             continue
#         if s.lower().startswith(("node ", "edge ", "graph ", "rank", "label", "style", "color", "fillcolor")):
#             continue
#         prefix, attrs = _extract_bracket_attr(s)
#         node_id = _unquote_id(prefix)
#         if not node_id:
#             continue
#         label = attrs.get("label", node_id)
#         if _is_pseudo_node(node_id, label):
#             continue
#         node_attrs[node_id] = attrs
#     return node_attrs


# def _parse_edge_attrs_from_dot(dot: str) -> Dict[Tuple[str, str], Dict[str, str]]:
#     edge_attrs: Dict[Tuple[str, str], Dict[str, str]] = {}
#     for stmt in _split_statements(dot):
#         s = stmt.strip()
#         if "->" not in s and "--" not in s:
#             continue
#         prefix, attrs = _extract_bracket_attr(s)
#         tokens = re.findall(r'"[^"]+"|[A-Za-z0-9_.:-]+', prefix)
#         if len(tokens) >= 2:
#             for src_raw, dst_raw in zip(tokens[:-1], tokens[1:]):
#                 src = _unquote_id(src_raw)
#                 dst = _unquote_id(dst_raw)
#                 if src and dst:
#                     edge_attrs[(src, dst)] = attrs
#     return edge_attrs


# def _sanitize_dot_for_plain(dot: str) -> str:
#     cleaned_statements: List[str] = []
#     for stmt in _split_statements(dot):
#         s = html.unescape(stmt.strip())
#         if not s:
#             continue
#         if "->" in s or "--" in s or s.lower().startswith(("digraph", "graph", "subgraph")):
#             cleaned_statements.append(stmt)
#             continue
#         if "[" not in s and "=" in s:
#             lhs = s.split("=", 1)[0].strip().lower()
#             if lhs in {"label", "color", "style", "fillcolor", "fontcolor", "fontsize", "fontname", "rankdir", "splines", "bgcolor", "margin", "pad"}:
#                 continue
#         prefix, attrs = _extract_bracket_attr(s)
#         node_id = _unquote_id(prefix)
#         label = attrs.get("label", node_id)
#         if node_id and _is_pseudo_node(node_id, label):
#             continue
#         cleaned_statements.append(stmt)
#     body = ";\n".join(cleaned_statements)
#     if "digraph" not in body.lower() and "graph" not in body.lower():
#         body = "digraph G {\n" + body + "\n}"
#     return body

# # ---------------------------------------------------------------------
# # Graphviz coordinate rendering path
# # ---------------------------------------------------------------------
# def _run_graphviz_plain(dot: str) -> str:
#     with tempfile.TemporaryDirectory() as tmp:
#         dot_path = Path(tmp) / "diagram.dot"
#         dot_path.write_text(dot, encoding="utf-8")
#         proc = subprocess.run(["dot", "-Tplain", str(dot_path)], check=False, capture_output=True, text=True)
#         if proc.returncode != 0:
#             raise RuntimeError(f"Graphviz plain layout failed: {proc.stderr.strip()}")
#         return proc.stdout


# def _parse_graphviz_plain(plain_text: str) -> Dict[str, Any]:
#     graph_width = 1.0
#     graph_height = 1.0
#     nodes: Dict[str, Dict[str, Any]] = {}
#     edges: List[Dict[str, Any]] = []
#     for raw_line in (plain_text or "").splitlines():
#         line = raw_line.strip()
#         if not line or line == "stop":
#             continue
#         try:
#             parts = shlex.split(line)
#         except Exception:
#             logger.debug("Could not parse Graphviz plain line: %s", line, exc_info=True)
#             continue
#         if not parts:
#             continue
#         if parts[0] == "graph" and len(parts) >= 4:
#             graph_width = max(0.1, float(parts[2]))
#             graph_height = max(0.1, float(parts[3]))
#         elif parts[0] == "node" and len(parts) >= 6:
#             node_id = parts[1]
#             label = parts[6] if len(parts) >= 7 else node_id
#             if _is_pseudo_node(node_id, label):
#                 continue
#             nodes[node_id] = {
#                 "id": node_id,
#                 "x": float(parts[2]),
#                 "y": float(parts[3]),
#                 "width": max(NODE_MIN_WIDTH_IN, float(parts[4])),
#                 "height": max(NODE_MIN_HEIGHT_IN, float(parts[5])),
#                 "label": _label(label, node_id),
#             }
#         elif parts[0] == "edge" and len(parts) >= 5:
#             src = parts[1]
#             dst = parts[2]
#             try:
#                 point_count = int(parts[3])
#             except Exception:
#                 continue
#             coords = parts[4 : 4 + point_count * 2]
#             points = []
#             for idx in range(0, len(coords), 2):
#                 try:
#                     points.append((float(coords[idx]), float(coords[idx + 1])))
#                 except Exception:
#                     pass
#             if points:
#                 edges.append({"source": src, "target": dst, "points": points})
#     return {"width": graph_width, "height": graph_height, "nodes": nodes, "edges": edges}


# def _build_graphviz_layout_model(diagram_text: str) -> Dict[str, Any]:
#     normalized = normalize_to_graphviz_dot(diagram_text)
#     if not normalized:
#         raise ValueError("Diagram text could not be normalized to DOT.")
#     normalized = html.unescape(normalized)
#     dot = _strip_dot_comments(normalized)
#     clusters = _parse_clusters_from_dot(dot)
#     node_attrs = _parse_node_attrs_from_dot(dot)
#     edge_attrs = _parse_edge_attrs_from_dot(dot)
#     plain = _run_graphviz_plain(_sanitize_dot_for_plain(dot))
#     layout = _parse_graphviz_plain(plain)
#     if not layout.get("nodes"):
#         logger.warning("Graphviz sanitized plain layout returned no nodes; retrying original DOT.")
#         layout = _parse_graphviz_plain(_run_graphviz_plain(dot))
#     nodes: Dict[str, Dict[str, Any]] = {}
#     for node_id, plain_node in layout["nodes"].items():
#         attrs = node_attrs.get(node_id, {})
#         label = _label(attrs.get("label") or plain_node.get("label"), node_id)
#         if _is_pseudo_node(node_id, label):
#             continue
#         nodes[node_id] = {**plain_node, "label": label, "attrs": attrs, "icon_hint": _guess_icon_hint(node_id, label, attrs)}
#     edges: List[Dict[str, Any]] = []
#     legend: List[Tuple[str, str]] = []
#     seen_legend = set()
#     for edge in layout["edges"]:
#         src = edge["source"]
#         dst = edge["target"]
#         if src not in nodes or dst not in nodes:
#             continue
#         attrs = edge_attrs.get((src, dst), {})
#         color = _safe_color(attrs.get("color") or attrs.get("fontcolor"), DEFAULT_EDGE_COLOR)
#         label = _clean_label(attrs.get("label", ""))
#         edges.append({**edge, "color": color, "label": label})
#         if label:
#             key = (color, label)
#             if key not in seen_legend:
#                 seen_legend.add(key)
#                 legend.append(key)
#     for cluster in clusters:
#         cluster["nodes"] = [nid for nid in cluster["nodes"] if nid in nodes]
#     clusters = [cluster for cluster in clusters if cluster["nodes"]]
#     cluster_by_node = {nid: cluster["id"] for cluster in clusters for nid in cluster["nodes"]}
#     logger.info("Editable PPTX Graphviz model parsed. nodes=%s edges=%s clusters=%s legend=%s", len(nodes), len(edges), len(clusters), len(legend))
#     return {"mode": "graphviz", "graph_width": layout["width"], "graph_height": layout["height"], "nodes": nodes, "edges": edges, "clusters": clusters, "cluster_by_node": cluster_by_node, "legend": legend}


# def _compute_coordinate_transform(model: Dict[str, Any]) -> Dict[str, float]:
#     graph_width = max(0.1, float(model.get("graph_width") or 1.0))
#     graph_height = max(0.1, float(model.get("graph_height") or 1.0))
#     scale = min(CANVAS_WIDTH_IN / graph_width, CANVAS_HEIGHT_IN / graph_height)
#     used_w = graph_width * scale
#     used_h = graph_height * scale
#     offset_x = CANVAS_LEFT_IN + max(0.0, (CANVAS_WIDTH_IN - used_w) / 2.0)
#     offset_y = CANVAS_TOP_IN + max(0.0, (CANVAS_HEIGHT_IN - used_h) / 2.0)
#     return {"scale": scale, "offset_x": offset_x, "offset_y": offset_y, "graph_height": graph_height}


# def _point_to_ppt(x: float, y: float, transform: Dict[str, float]) -> Tuple[float, float]:
#     scale = transform["scale"]
#     px = transform["offset_x"] + x * scale
#     py = transform["offset_y"] + (transform["graph_height"] - y) * scale
#     return px, py


# def _node_box_to_ppt(node: Dict[str, Any], transform: Dict[str, float]) -> Dict[str, float]:
#     cx, cy = _point_to_ppt(float(node["x"]), float(node["y"]), transform)
#     width = max(0.82, float(node["width"]) * transform["scale"])
#     height = max(0.42, float(node["height"]) * transform["scale"])
#     return {"left": cx - width / 2.0, "top": cy - height / 2.0, "width": width, "height": height}


# def _cluster_boxes_from_nodes(model: Dict[str, Any], placements: Dict[str, Dict[str, float]]) -> List[Dict[str, Any]]:
#     boxes: List[Dict[str, Any]] = []
#     for cluster in model.get("clusters", []):
#         node_boxes = [placements[nid] for nid in cluster["nodes"] if nid in placements]
#         if not node_boxes:
#             continue
#         left = min(box["left"] for box in node_boxes) - 0.14
#         top = min(box["top"] for box in node_boxes) - 0.32
#         right = max(box["left"] + box["width"] for box in node_boxes) + 0.14
#         bottom = max(box["top"] + box["height"] for box in node_boxes) + 0.14
#         boxes.append({"id": cluster["id"], "label": cluster["label"], "left": max(CANVAS_LEFT_IN, left), "top": max(CANVAS_TOP_IN, top), "width": max(0.75, right - left), "height": max(0.55, bottom - top)})
#     return boxes

# # ---------------------------------------------------------------------
# # Fallback editable model if Graphviz-coordinate mode fails
# # ---------------------------------------------------------------------
# def _parse_fallback_editable_model(diagram_text: str) -> Dict[str, Any]:
#     normalized = normalize_to_graphviz_dot(diagram_text)
#     if not normalized:
#         raise ValueError("Diagram text could not be normalized to DOT.")
#     dot = _strip_dot_comments(html.unescape(normalized))
#     clusters = _parse_clusters_from_dot(dot)
#     node_attrs = _parse_node_attrs_from_dot(dot)
#     edge_attrs = _parse_edge_attrs_from_dot(dot)
#     nodes: Dict[str, Dict[str, Any]] = {}
#     edges: List[Dict[str, Any]] = []
#     # node declarations
#     for node_id, attrs in node_attrs.items():
#         label = _label(attrs.get("label"), node_id)
#         if _is_pseudo_node(node_id, label):
#             continue
#         nodes[node_id] = {"id": node_id, "label": label, "attrs": attrs, "icon_hint": _guess_icon_hint(node_id, label, attrs)}
#     # edge endpoints + labels
#     for (src, dst), attrs in edge_attrs.items():
#         if _is_pseudo_node(src) or _is_pseudo_node(dst):
#             continue
#         for nid in [src, dst]:
#             if nid not in nodes:
#                 label = _label(None, nid)
#                 nodes[nid] = {"id": nid, "label": label, "attrs": {}, "icon_hint": _guess_icon_hint(nid, label, {})}
#         edges.append({"source": src, "target": dst, "color": _safe_color(attrs.get("color") or attrs.get("fontcolor"), DEFAULT_EDGE_COLOR), "label": _clean_label(attrs.get("label", ""))})
#     legend: List[Tuple[str, str]] = []
#     seen = set()
#     for edge in edges:
#         if edge.get("label"):
#             key = (edge["color"], edge["label"])
#             if key not in seen:
#                 seen.add(key)
#                 legend.append(key)
#     cluster_by_node = {nid: cluster["id"] for cluster in clusters for nid in cluster["nodes"] if nid in nodes}
#     return {"mode": "fallback", "nodes": list(nodes.values()), "edges": edges, "clusters": clusters, "cluster_by_node": cluster_by_node, "legend": legend}


# def _compute_layers(node_ids: List[str], edges: List[Dict[str, Any]]) -> Dict[str, int]:
#     indeg = {nid: 0 for nid in node_ids}
#     outgoing: Dict[str, List[str]] = defaultdict(list)
#     for edge in edges:
#         src = edge["source"]
#         dst = edge["target"]
#         indeg.setdefault(src, 0)
#         indeg.setdefault(dst, 0)
#         outgoing[src].append(dst)
#         indeg[dst] += 1
#     q = deque([nid for nid, deg in indeg.items() if deg == 0])
#     layer = {nid: 0 for nid in q}
#     visited = 0
#     while q:
#         cur = q.popleft()
#         visited += 1
#         for nxt in outgoing.get(cur, []):
#             layer[nxt] = max(layer.get(nxt, 0), layer.get(cur, 0) + 1)
#             indeg[nxt] -= 1
#             if indeg[nxt] == 0:
#                 q.append(nxt)
#     if visited < len(indeg):
#         fallback_layer = max(layer.values(), default=0)
#         for nid in indeg:
#             layer.setdefault(nid, fallback_layer)
#     return layer


# def _layout_fallback_model(model: Dict[str, Any]) -> Dict[str, Any]:
#     nodes = model["nodes"]
#     edges = model["edges"]
#     clusters = model["clusters"]
#     cluster_by_node = model["cluster_by_node"]
#     node_ids = [node["id"] for node in nodes]
#     layer_by_node = _compute_layers(node_ids, edges)
#     cluster_nodes: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
#     cluster_order: List[str] = []
#     for node in nodes:
#         cid = cluster_by_node.get(node["id"])
#         if cid:
#             cluster_nodes[cid].append(node)
#             if cid not in cluster_order:
#                 cluster_order.append(cid)
#     unclustered = [n for n in nodes if n["id"] not in cluster_by_node]
#     cluster_map = {cluster["id"]: cluster for cluster in clusters}
#     x = CANVAS_LEFT_IN
#     placements: Dict[str, Dict[str, float]] = {}
#     cluster_boxes: List[Dict[str, Any]] = []
#     def place_group(group_nodes: List[Dict[str, Any]], label: Optional[str], group_id: Optional[str]):
#         nonlocal x
#         if not group_nodes:
#             return
#         layers: Dict[int, List[Dict[str, Any]]] = defaultdict(list)
#         for node in group_nodes:
#             layers[layer_by_node.get(node["id"], 0)].append(node)
#         layer_keys = sorted(layers.keys())
#         max_rows = max(len(layers[k]) for k in layer_keys)
#         inner_w = len(layer_keys) * FALLBACK_NODE_WIDTH_IN + max(0, len(layer_keys) - 1) * FALLBACK_H_GAP_IN
#         inner_h = max_rows * FALLBACK_NODE_HEIGHT_IN + max(0, max_rows - 1) * FALLBACK_V_GAP_IN
#         group_w = inner_w + 2 * FALLBACK_CLUSTER_PAD_IN
#         group_h = inner_h + FALLBACK_CLUSTER_TITLE_H_IN + 2 * FALLBACK_CLUSTER_PAD_IN
#         group_left = x
#         group_top = CANVAS_TOP_IN
#         for col_idx, layer_key in enumerate(layer_keys):
#             col_nodes = sorted(layers[layer_key], key=lambda n: n["label"].lower())
#             col_x = group_left + FALLBACK_CLUSTER_PAD_IN + col_idx * (FALLBACK_NODE_WIDTH_IN + FALLBACK_H_GAP_IN)
#             total_h = len(col_nodes) * FALLBACK_NODE_HEIGHT_IN + max(0, len(col_nodes) - 1) * FALLBACK_V_GAP_IN
#             start_y = group_top + FALLBACK_CLUSTER_TITLE_H_IN + FALLBACK_CLUSTER_PAD_IN + max(0, (inner_h - total_h) / 2)
#             for row_idx, node in enumerate(col_nodes):
#                 placements[node["id"]] = {"left": col_x, "top": start_y + row_idx * (FALLBACK_NODE_HEIGHT_IN + FALLBACK_V_GAP_IN), "width": FALLBACK_NODE_WIDTH_IN, "height": FALLBACK_NODE_HEIGHT_IN}
#         if group_id:
#             cluster_boxes.append({"id": group_id, "label": label or group_id, "left": group_left, "top": group_top, "width": group_w, "height": group_h})
#         x += group_w + FALLBACK_H_GAP_IN
#     for cid in cluster_order:
#         place_group(cluster_nodes[cid], cluster_map.get(cid, {}).get("label", cid), cid)
#     if unclustered:
#         place_group(unclustered, None, None)
#     # scale down to fit canvas
#     if placements:
#         right = max(v["left"] + v["width"] for v in placements.values())
#         bottom = max(v["top"] + v["height"] for v in placements.values())
#         for box in cluster_boxes:
#             right = max(right, box["left"] + box["width"])
#             bottom = max(bottom, box["top"] + box["height"])
#         scale = min(1.0, CANVAS_WIDTH_IN / max(0.1, right - CANVAS_LEFT_IN), CANVAS_HEIGHT_IN / max(0.1, bottom - CANVAS_TOP_IN))
#         if scale < 1:
#             def sx(v): return CANVAS_LEFT_IN + (v - CANVAS_LEFT_IN) * scale
#             def sy(v): return CANVAS_TOP_IN + (v - CANVAS_TOP_IN) * scale
#             for v in placements.values():
#                 v["left"] = sx(v["left"]); v["top"] = sy(v["top"]); v["width"] *= scale; v["height"] *= scale
#             for box in cluster_boxes:
#                 box["left"] = sx(box["left"]); box["top"] = sy(box["top"]); box["width"] *= scale; box["height"] *= scale
#     return {"placements": placements, "cluster_boxes": cluster_boxes}

# # ---------------------------------------------------------------------
# # Drawing helpers
# # ---------------------------------------------------------------------
# def _set_text_frame(shape: BaseShape, text: str, font_size_pt: float = 10, bold: bool = False, color_hex: str = DEFAULT_NODE_FONT, align=PP_ALIGN.LEFT):
#     tf = shape.text_frame
#     tf.clear()
#     tf.word_wrap = True
#     tf.margin_left = Inches(0.02)
#     tf.margin_right = Inches(0.02)
#     tf.margin_top = Inches(0.01)
#     tf.margin_bottom = Inches(0.01)
#     tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
#     p = tf.paragraphs[0]
#     p.alignment = align
#     run = p.add_run()
#     run.text = _safe_title(text)
#     font = run.font
#     font.name = "Calibri"
#     font.size = Pt(font_size_pt)
#     font.bold = bold
#     font.color.rgb = _rgb(color_hex, DEFAULT_NODE_FONT)


# def _font_for_label(label: str, width: float, height: float) -> float:
#     clean = _safe_title(label)
#     lines = max(1, clean.count("\n") + 1)
#     length = len(clean.replace("\n", " "))
#     if width < 1.2 or height < 0.5 or length > 55 or lines >= 3:
#         return 7.3
#     if length > 38 or lines == 2:
#         return 8.2
#     return 9.2


# def _add_title(slide, title_text: str):
#     shape = slide.shapes.add_textbox(Inches(TITLE_LEFT_IN), Inches(TITLE_TOP_IN), Inches(TITLE_WIDTH_IN), Inches(TITLE_HEIGHT_IN))
#     _set_text_frame(shape, title_text, font_size_pt=18, bold=True, color_hex="202124")


# def _add_cluster_box(slide, box: Dict[str, Any]):
#     shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
#     shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(DEFAULT_CLUSTER_FILL)
#     shape.line.color.rgb = _rgb(DEFAULT_CLUSTER_BORDER); shape.line.width = Pt(1.1)
#     try: shape.adjustments[0] = 0.06
#     except Exception: pass
#     label_box = slide.shapes.add_textbox(Inches(box["left"] + 0.06), Inches(box["top"] + 0.02), Inches(max(0.5, box["width"] - 0.12)), Inches(0.20))
#     _set_text_frame(label_box, box.get("label", ""), font_size_pt=8.5, bold=True, color_hex=DEFAULT_CLUSTER_FONT)


# def _add_node_box(slide, box: Dict[str, float], label: str, icon_path: Optional[Path] = None):
#     shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
#     shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(DEFAULT_NODE_FILL)
#     shape.line.color.rgb = _rgb(DEFAULT_NODE_BORDER); shape.line.width = Pt(1.0)
#     try: shape.adjustments[0] = 0.08
#     except Exception: pass
#     text_left = box["left"] + NODE_INNER_PADDING_IN
#     icon_size = min(NODE_ICON_SIZE_IN, max(0.16, box["height"] * 0.40))
#     if icon_path and icon_path.exists() and box["width"] > 0.72:
#         try:
#             slide.shapes.add_picture(str(icon_path), Inches(box["left"] + NODE_INNER_PADDING_IN), Inches(box["top"] + (box["height"] - icon_size) / 2.0), width=Inches(icon_size), height=Inches(icon_size))
#             text_left = box["left"] + NODE_TEXT_LEFT_WITH_ICON_IN
#         except Exception:
#             logger.debug("Could not add icon to PPTX node: %s", icon_path, exc_info=True)
#     text_box = slide.shapes.add_textbox(Inches(text_left), Inches(box["top"] + 0.03), Inches(max(0.35, box["width"] - (text_left - box["left"]) - 0.05)), Inches(max(0.22, box["height"] - 0.06)))
#     _set_text_frame(text_box, label, font_size_pt=_font_for_label(label, box["width"], box["height"]), bold=True, color_hex=DEFAULT_NODE_FONT)


# def _add_connector(slide, p1: Tuple[float, float], p2: Tuple[float, float], color_hex: str = DEFAULT_EDGE_COLOR):
#     conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p1[0]), Inches(p1[1]), Inches(p2[0]), Inches(p2[1]))
#     conn.line.color.rgb = _rgb(color_hex, DEFAULT_EDGE_COLOR)
#     conn.line.width = Pt(DEFAULT_EDGE_WIDTH_PT)
#     return conn


# def _add_edge_polyline(slide, points: List[Tuple[float, float]], color_hex: str = DEFAULT_EDGE_COLOR):
#     if len(points) < 2:
#         return
#     for idx in range(len(points) - 1):
#         segment = _add_connector(slide, points[idx], points[idx + 1], color_hex)
#         if idx == len(points) - 2:
#             try: segment.line.end_arrowhead = True
#             except Exception: pass


# def _add_straight_box_edge(slide, src_box: Dict[str, float], dst_box: Dict[str, float], color_hex: str = DEFAULT_EDGE_COLOR):
#     p1 = (src_box["left"] + src_box["width"], src_box["top"] + src_box["height"] / 2)
#     p2 = (dst_box["left"], dst_box["top"] + dst_box["height"] / 2)
#     segment = _add_connector(slide, p1, p2, color_hex)
#     try: segment.line.end_arrowhead = True
#     except Exception: pass


# def _add_legend(slide, legend_items: List[Tuple[str, str]]):
#     if not legend_items:
#         return
#     title_box = slide.shapes.add_textbox(Inches(LEGEND_LEFT_IN), Inches(LEGEND_TOP_IN), Inches(1.2), Inches(0.18))
#     _set_text_frame(title_box, "Flow Legend", font_size_pt=9.2, bold=True, color_hex="202124")
#     current_x = LEGEND_LEFT_IN
#     current_y = LEGEND_TOP_IN + 0.24
#     max_x = LEGEND_LEFT_IN + LEGEND_WIDTH_IN
#     for color, label in legend_items[:16]:
#         label = _safe_title(label)
#         item_width = max(1.15, min(3.0, 0.42 + len(label) * 0.052))
#         if current_x + item_width > max_x:
#             current_x = LEGEND_LEFT_IN
#             current_y += 0.22
#         if current_y > LEGEND_TOP_IN + 1.05:
#             break
#         chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(current_x), Inches(current_y), Inches(item_width), Inches(0.18))
#         chip.fill.solid(); chip.fill.fore_color.rgb = _rgb(LEGEND_FILL)
#         chip.line.color.rgb = _rgb(LEGEND_BORDER); chip.line.width = Pt(0.6)
#         dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(current_x + 0.04), Inches(current_y + 0.045), Inches(0.09), Inches(0.09))
#         dot.fill.solid(); dot.fill.fore_color.rgb = _rgb(color, DEFAULT_EDGE_COLOR); dot.line.color.rgb = _rgb(color, DEFAULT_EDGE_COLOR)
#         label_box = slide.shapes.add_textbox(Inches(current_x + 0.16), Inches(current_y + 0.01), Inches(max(0.3, item_width - 0.18)), Inches(0.15))
#         _set_text_frame(label_box, label, font_size_pt=7.4, bold=False, color_hex="202124")
#         current_x += item_width + 0.08


# def _add_fallback_preview_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str):
#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, f"{slide_title} (Preview)")
#     try:
#         assets = render_graphviz_to_assets(diagram_text, output_dir=output_dir, base_name=base_name)
#     except Exception:
#         logger.exception("Fallback preview asset generation failed.")
#         assets = None
#     if assets and assets.get("png_path") and Path(assets["png_path"]).exists():
#         slide.shapes.add_picture(str(assets["png_path"]), Inches(CANVAS_LEFT_IN), Inches(CANVAS_TOP_IN), width=Inches(PREVIEW_MAX_WIDTH_IN))
#         _add_legend(slide, assets.get("legend", []))

# # ---------------------------------------------------------------------
# # Diagram extraction
# # ---------------------------------------------------------------------
# def _is_diagram_text(value: Any) -> bool:
#     return isinstance(value, str) and normalize_to_graphviz_dot(_clean_internal_tokens(value)) is not None


# def _walk_for_diagrams(value: Any, path_parts: List[str], results: List[Dict[str, Any]]):
#     if value is None:
#         return
#     if isinstance(value, str):
#         if _is_diagram_text(value):
#             results.append({"title": " / ".join(path_parts) if path_parts else "Diagram", "diagram_text": _clean_internal_tokens(value), "section_key": _safe_name(path_parts[-1] if path_parts else "diagram")})
#         return
#     if isinstance(value, list):
#         for idx, item in enumerate(value, start=1):
#             if _is_diagram_text(item):
#                 results.append({"title": " / ".join(path_parts) if path_parts else f"Diagram {idx}", "diagram_text": _clean_internal_tokens(item), "section_key": _safe_name(path_parts[-1] if path_parts else f"diagram_{idx}")})
#             elif isinstance(item, (dict, list)):
#                 _walk_for_diagrams(item, path_parts, results)
#         return
#     if isinstance(value, dict):
#         for key, child in value.items():
#             key_text = str(key).replace("_", " ").title()
#             next_path = path_parts if str(key).lower() == "diagrams" else path_parts + [key_text]
#             _walk_for_diagrams(child, next_path, results)


# def extract_diagram_specs(data: Dict[str, Any], hld_model: Any, selected_sections_set: Optional[Set[str]] = None) -> List[Dict[str, Any]]:
#     results: List[Dict[str, Any]] = []
#     for field_name, field_info in hld_model.model_fields.items():
#         if not should_include_section(field_name, field_info, selected_sections_set or set()):
#             continue
#         section_title = field_info.title or field_name.replace("_", " ").title()
#         _walk_for_diagrams(data.get(field_name), [section_title], results)
#     title_counts: Dict[str, int] = defaultdict(int)
#     for item in results:
#         title_counts[item["title"]] += 1
#     dedupe_counter: Dict[str, int] = defaultdict(int)
#     for item in results:
#         if title_counts[item["title"]] > 1:
#             dedupe_counter[item["title"]] += 1
#             item["title"] = f'{item["title"]} ({dedupe_counter[item["title"]]})'
#     return results

# # ---------------------------------------------------------------------
# # Slide rendering
# # ---------------------------------------------------------------------
# def add_diagram_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str):
#     # First try Graphviz coordinate layout for PDF/Word-like placement.
#     try:
#         model = _build_graphviz_layout_model(diagram_text)
#         if not model.get("nodes"):
#             raise ValueError("No nodes parsed from Graphviz coordinate model")
#         transform = _compute_coordinate_transform(model)
#         placements = {node_id: _node_box_to_ppt(node, transform) for node_id, node in model["nodes"].items()}
#         cluster_boxes = _cluster_boxes_from_nodes(model, placements)
#         mode = "graphviz"
#     except Exception:
#         logger.warning("Graphviz-coordinate editable rendering failed for slide '%s'. Trying editable fallback layout.", slide_title, exc_info=True)
#         try:
#             model = _parse_fallback_editable_model(diagram_text)
#             layout = _layout_fallback_model(model)
#             placements = layout["placements"]
#             cluster_boxes = layout["cluster_boxes"]
#             mode = "fallback"
#             if not placements:
#                 raise ValueError("Fallback editable layout produced no placements")
#         except Exception:
#             logger.exception("Editable PPTX rendering failed for slide '%s'. Falling back to preview image.", slide_title)
#             _add_fallback_preview_slide(prs, slide_title, diagram_text, output_dir, base_name)
#             return
#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, slide_title)
#     # cluster backgrounds
#     for box in cluster_boxes:
#         _add_cluster_box(slide, box)
#     # edges behind nodes
#     if mode == "graphviz":
#         for edge in model.get("edges", []):
#             points = [_point_to_ppt(x, y, transform) for x, y in edge.get("points", [])]
#             _add_edge_polyline(slide, points, color_hex=edge.get("color") or DEFAULT_EDGE_COLOR)
#         node_items = model["nodes"].items()
#     else:
#         for edge in model.get("edges", []):
#             src = placements.get(edge["source"])
#             dst = placements.get(edge["target"])
#             if src and dst:
#                 _add_straight_box_edge(slide, src, dst, edge.get("color") or DEFAULT_EDGE_COLOR)
#         node_items = [(node["id"], node) for node in model["nodes"]]
#     # nodes/icons/text on top
#     for node_id, node in node_items:
#         box = placements.get(node_id)
#         if not box:
#             continue
#         icon_path = _preferred_picture_icon_path(node.get("icon_hint") or node.get("label") or node_id)
#         _add_node_box(slide, box, label=node.get("label") or node_id, icon_path=icon_path)
#     _add_legend(slide, model.get("legend", []))

# # ---------------------------------------------------------------------
# # Main public builder
# # ---------------------------------------------------------------------
# def build_editable_pptx(data: Dict[str, Any], pptx_path: str, hld_model: Any, selected_sections_set: Optional[Set[str]] = None):
#     output_path = Path(pptx_path)
#     output_path.parent.mkdir(parents=True, exist_ok=True)
#     diagram_assets_dir = output_path.parent / f"{output_path.stem}_diagram_assets"
#     diagram_assets_dir.mkdir(parents=True, exist_ok=True)
#     diagram_specs = extract_diagram_specs(data=data, hld_model=hld_model, selected_sections_set=selected_sections_set)
#     prs = Presentation()
#     prs.slide_width = Inches(SLIDE_WIDTH_IN)
#     prs.slide_height = Inches(SLIDE_HEIGHT_IN)
#     if not diagram_specs:
#         slide = prs.slides.add_slide(prs.slide_layouts[6])
#         _add_title(slide, "Editable Diagrams")
#         msg = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11.5), Inches(1.0))
#         _set_text_frame(msg, "No diagrams were detected in the selected sections.", font_size_pt=16, bold=True, color_hex="5F6368", align=PP_ALIGN.CENTER)
#         prs.save(str(output_path))
#         return
#     for counter, spec in enumerate(diagram_specs, start=1):
#         base_name = f"{_safe_name(spec['section_key'])}_{counter:02d}"
#         add_diagram_slide(prs=prs, slide_title=spec["title"], diagram_text=spec["diagram_text"], output_dir=diagram_assets_dir, base_name=base_name)
#     prs.save(str(output_path))
#     logger.info("Editable PPTX generated: %s", output_path)




#2 running code
# from __future__ import annotations

# import html
# import logging
# import math
# import re
# from collections import defaultdict, deque
# from pathlib import Path
# from typing import Any, Dict, List, Optional, Set, Tuple

# from pptx import Presentation
# from pptx.dml.color import RGBColor
# from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
# from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
# from pptx.shapes.base import BaseShape
# from pptx.util import Inches, Pt

# from .config import ICON_DIR
# from .graphviz_renderer import normalize_to_graphviz_dot, render_graphviz_to_assets
# from .schema_utils import should_include_section
# from .text_sanitizer import renderer_safe_plain_text

# logger = logging.getLogger(__name__)

# # =============================================================================
# # Editable PPTX Diagram Renderer
# # =============================================================================
# # This renderer intentionally uses a stable editable grid/flow layout instead of
# # trying to reproduce every Graphviz coordinate. The goal is a PPTX diagram that
# # is complete, readable and editable: boxes, text, icons, arrows and legend.
# # =============================================================================

# SLIDE_WIDTH_IN = 13.333
# SLIDE_HEIGHT_IN = 7.5

# TITLE_LEFT_IN = 0.35
# TITLE_TOP_IN = 0.12
# TITLE_WIDTH_IN = 12.6
# TITLE_HEIGHT_IN = 0.42

# CANVAS_LEFT_IN = 0.35
# CANVAS_TOP_IN = 0.68
# CANVAS_WIDTH_IN = 12.3
# CANVAS_HEIGHT_IN = 5.30

# LEGEND_LEFT_IN = 0.35
# LEGEND_TOP_IN = 6.14
# LEGEND_WIDTH_IN = 12.3
# LEGEND_MAX_HEIGHT_IN = 1.20

# NODE_WIDTH_IN = 2.05
# NODE_HEIGHT_IN = 0.82
# NODE_ICON_SIZE_IN = 0.23
# NODE_PADDING_IN = 0.07
# NODE_TEXT_LEFT_WITH_ICON_IN = 0.34

# H_GAP_IN = 0.34
# V_GAP_IN = 0.42
# CLUSTER_PAD_IN = 0.16
# CLUSTER_TITLE_HEIGHT_IN = 0.24

# PREVIEW_MAX_WIDTH_IN = 12.0
# PREVIEW_MAX_HEIGHT_IN = 5.5

# DEFAULT_NODE_FILL = "F8F9FA"
# DEFAULT_NODE_BORDER = "DADCE0"
# DEFAULT_NODE_FONT = "202124"
# DEFAULT_CLUSTER_FILL = "F3F8FF"
# DEFAULT_CLUSTER_BORDER = "A8C7FA"
# DEFAULT_CLUSTER_FONT = "174EA6"
# DEFAULT_EDGE_COLOR = "5F6368"
# DEFAULT_EDGE_WIDTH_PT = 1.25
# LEGEND_FILL = "FFFFFF"
# LEGEND_BORDER = "DADCE0"

# SUPPORTED_PICTURE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif"}
# PREFERRED_ICON_EXTS = [".png", ".jpg", ".jpeg", ".bmp", ".gif"]

# COMMENT_LINE_RE = re.compile(r"//.*?$", re.MULTILINE)
# BLOCK_COMMENT_RE = re.compile(r"/\*.*?\*/", re.DOTALL)
# EDGE_RE = re.compile(r'("[^"]+"|[A-Za-z0-9_.:-]+)\s*(->|--)\s*("[^"]+"|[A-Za-z0-9_.:-]+)')
# NODE_DECL_RE = re.compile(r'(?m)^\s*("[^"]+"|[A-Za-z0-9_.:-]+)\s*\[(.*?)\]\s*;?\s*$', re.DOTALL)

# FLOW_PALETTE = ["E60000", "1D70B8", "28A197", "F47738", "6F35A5", "007C89", "D53880", "85994B", "B58840", "5F6368"]
# COLOR_NAME_MAP = {
#     "red": "E60000", "blue": "1D70B8", "green": "28A197", "orange": "F47738",
#     "purple": "6F35A5", "teal": "007C89", "gray": "5F6368", "grey": "5F6368",
#     "black": "202124", "lightgrey": "DADCE0", "lightgray": "DADCE0",
# }
# ICON_KEYWORD_MAP = [
#     ("cloud storage", "cloud_storage"), ("gcs", "cloud_storage"), ("storage", "cloud_storage"),
#     ("cloud function", "cloud_functions"), ("cloud functions", "cloud_functions"), ("function", "cloud_functions"),
#     ("pub/sub", "pubsub"), ("pubsub", "pubsub"), ("notification", "pubsub"), ("eventarc", "eventarc"),
#     ("vault", "secret_manager"), ("hashicorp", "secret_manager"), ("secret", "secret_manager"),
#     ("kms", "key_management_service"), ("key management", "key_management_service"),
#     ("logging", "cloud_logging"), ("monitoring", "cloud_monitoring"),
#     ("cloud build", "cloud_build"), ("build", "cloud_build"), ("terraform", "terraform"),
#     ("teradata", "database_migration_service"), ("database", "database_migration_service"),
#     ("mft", "transfer"), ("managed file transfer", "transfer"), ("transfer", "transfer"),
#     ("interconnect", "cloud_interconnect"), ("vpn", "cloud_vpn"),
#     ("collibra", "catalog"), ("catalog", "catalog"),
# ]

# # -----------------------------------------------------------------------------
# # Text, color and filtering helpers
# # -----------------------------------------------------------------------------
# def _clean_internal_tokens(text: Any) -> str:
#     value = str(text or "")
#     for token in [
#         "AIASECTIONBLOCKSTARTTOKEN", "AIASECTIONBLOCKENDTOKEN",
#         "[[AIASECTIONBLOCKSTARTTOKEN]]", "[[AIASECTIONBLOCKENDTOKEN]]",
#         "SECTION_BLOCK_START", "SECTION_BLOCK_END",
#         "[[SECTION_BLOCK_START]]", "[[SECTION_BLOCK_END]]",
#         "[[SECTIONBLOCKSTART]]", "[[SECTIONBLOCKEND]]",
#     ]:
#         value = value.replace(token, "")
#     return html.unescape(value.strip())


# def _safe_title(text: Any) -> str:
#     return renderer_safe_plain_text(_clean_internal_tokens(text)).strip() or "Diagram"


# def _safe_name(text: str) -> str:
#     return re.sub(r"[^a-zA-Z0-9_-]+", "_", str(text or "")).strip("_").lower() or "diagram"


# def _clean_label(value: Any) -> str:
#     text = html.unescape(str(value or "").strip()).strip('"')
#     text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
#     text = re.sub(r"<[^>]+>", " ", text)
#     text = text.replace("\\N", "\n").replace("\\n", "\n").replace("&nbsp;", " ")
#     text = re.sub(r"\s+\n", "\n", text)
#     text = re.sub(r"\n\s+", "\n", text)
#     text = re.sub(r"[ \t]+", " ", text)
#     return text.strip()


# def _label(raw: Optional[str], fallback: str) -> str:
#     label = _clean_label(raw or "")
#     if not label:
#         label = str(fallback or "").replace("_", " ").replace("-", " ").title()
#     return label[:180]


# def _safe_color(value: Optional[str], default: str = DEFAULT_EDGE_COLOR) -> str:
#     raw = str(value or "").strip().strip('"').lstrip("#")
#     if not raw:
#         return default
#     if raw.lower() in COLOR_NAME_MAP:
#         return COLOR_NAME_MAP[raw.lower()]
#     raw = raw.upper()
#     return raw if len(raw) == 6 and all(ch in "0123456789ABCDEF" for ch in raw) else default


# def _rgb(hex_color: str, default: str = "000000") -> RGBColor:
#     return RGBColor.from_string(_safe_color(hex_color, default))


# def _is_pseudo_node(node_id: str, label: Optional[str] = None) -> bool:
#     text = _clean_label(label if label is not None else node_id).lower().strip()
#     raw = str(node_id or "").lower().strip()
#     if text in {"", "n/a", "na", "none", "null"}:
#         return True
#     if raw in {"node", "edge", "graph", "digraph", "subgraph", "rank", "label", "style", "color", "fillcolor", "fontcolor", "rankdir"}:
#         return True
#     if raw.startswith(("digraph", "graph", "subgraph", "cluster_", "cluster ")):
#         return True
#     if raw.startswith(("label=", "style=", "color=", "fillcolor=", "fontcolor=", "rankdir=")):
#         return True
#     if text.startswith(("label=", "style=", "color=", "fillcolor=", "fontcolor=", "rankdir=", "fontsize=", "fontname=", "margin=", "pad=")):
#         return True
#     if text in {"filled", "lightgrey", "lightgray", "blue", "green", "purple", "orange", "red", "gray", "grey"}:
#         return True
#     return False


# def _is_graphviz_control_statement(statement: str) -> bool:
#     s = html.unescape(str(statement or "")).strip().lower()
#     if not s or s in {"{", "}"}:
#         return True
#     return s.startswith((
#         "digraph ", "graph ", "subgraph ", "node ", "edge ", "rank ",
#         "rankdir", "label", "style", "color", "fillcolor", "fontcolor",
#         "fontsize", "fontname", "margin", "pad", "splines", "bgcolor",
#     ))

# # -----------------------------------------------------------------------------
# # Icon helpers
# # -----------------------------------------------------------------------------
# def _try_existing_icon_resolver(icon_hint: str) -> Optional[Path]:
#     try:
#         from . import icon_resolver as ir
#     except Exception:
#         return None
#     for fn_name in ["resolve_icon_path", "resolve_icon", "get_icon_path", "find_icon_path"]:
#         fn = getattr(ir, fn_name, None)
#         if not callable(fn):
#             continue
#         for args in [(icon_hint,), (icon_hint, True), (icon_hint, None)]:
#             try:
#                 result = fn(*args)
#                 if result:
#                     path = Path(str(result))
#                     if path.exists():
#                         return path
#             except TypeError:
#                 continue
#             except Exception:
#                 logger.debug("Icon resolver failed: %s", fn_name, exc_info=True)
#     return None


# def _search_icon_dir(icon_hint: str) -> Optional[Path]:
#     if not icon_hint:
#         return None
#     hint = _safe_name(icon_hint)
#     bases = {
#         hint, icon_hint, icon_hint.lower(), icon_hint.replace(" ", "_").lower(),
#         icon_hint.replace("-", "_").lower(), icon_hint.replace("/", "_").lower(),
#     }
#     for base in bases:
#         if not base:
#             continue
#         for ext in [".png", ".jpg", ".jpeg", ".bmp", ".gif", ".svg"]:
#             candidate = ICON_DIR / f"{base}{ext}"
#             if candidate.exists():
#                 return candidate
#     try:
#         for path in ICON_DIR.rglob("*"):
#             if path.is_file() and hint in path.stem.lower():
#                 return path
#     except Exception:
#         logger.debug("Icon directory search failed", exc_info=True)
#     return None


# def _keyword_icon_hint(label_or_id: str) -> Optional[str]:
#     text = _clean_label(label_or_id).lower()
#     for keyword, icon_hint in ICON_KEYWORD_MAP:
#         if keyword in text:
#             return icon_hint
#     return None


# def _preferred_picture_icon_path(icon_hint: str) -> Optional[Path]:
#     candidates: List[str] = []
#     keyword = _keyword_icon_hint(icon_hint)
#     if keyword:
#         candidates.append(keyword)
#     if icon_hint:
#         candidates.append(icon_hint)
#     for candidate in candidates:
#         path = _try_existing_icon_resolver(candidate) or _search_icon_dir(candidate)
#         if not path:
#             continue
#         if path.suffix.lower() in SUPPORTED_PICTURE_EXTS:
#             return path
#         if path.suffix.lower() == ".svg":
#             for ext in PREFERRED_ICON_EXTS:
#                 sibling = path.with_suffix(ext)
#                 if sibling.exists():
#                     return sibling
#     return None


# def _guess_icon_hint(node_id: str, node_label: str, attrs: Dict[str, str]) -> str:
#     image_attr = attrs.get("image") or attrs.get("icon") or attrs.get("imagepath")
#     if image_attr:
#         stem = Path(str(image_attr)).stem
#         if stem:
#             return stem
#     keyword = _keyword_icon_hint(node_label) or _keyword_icon_hint(node_id)
#     if keyword:
#         return keyword
#     for candidate in [node_id, node_label.split("\n", 1)[0], node_label.replace("\n", " ")]:
#         cleaned = _safe_name(candidate)
#         if cleaned:
#             return cleaned
#     return _safe_name(node_id or node_label or "node")

# # -----------------------------------------------------------------------------
# # DOT parser helpers
# # -----------------------------------------------------------------------------
# def _strip_dot_comments(dot_text: str) -> str:
#     text = BLOCK_COMMENT_RE.sub("", dot_text or "")
#     return COMMENT_LINE_RE.sub("", text)


# def _outer_graph_body(dot: str) -> str:
#     text = dot or ""
#     start = text.find("{")
#     end = text.rfind("}")
#     if start != -1 and end != -1 and end > start:
#         return text[start + 1:end]
#     return text


# def _split_statements(dot_text: str) -> List[str]:
#     statements: List[str] = []
#     buf: List[str] = []
#     bracket_depth = brace_depth = angle_depth = 0
#     for ch in dot_text or "":
#         if ch == "[":
#             bracket_depth += 1
#         elif ch == "]":
#             bracket_depth = max(0, bracket_depth - 1)
#         elif ch == "{":
#             brace_depth += 1
#         elif ch == "}":
#             brace_depth = max(0, brace_depth - 1)
#         elif ch == "<":
#             angle_depth += 1
#         elif ch == ">":
#             angle_depth = max(0, angle_depth - 1)
#         if ch == ";" and bracket_depth == 0 and brace_depth == 0 and angle_depth == 0:
#             stmt = "".join(buf).strip()
#             if stmt:
#                 statements.append(stmt)
#             buf = []
#         else:
#             buf.append(ch)
#     tail = "".join(buf).strip()
#     if tail:
#         statements.append(tail)
#     return statements


# def _extract_attr_pairs(attr_text: str) -> Dict[str, str]:
#     attrs: Dict[str, str] = {}
#     if not attr_text:
#         return attrs
#     text = html.unescape(str(attr_text))
#     for match in re.finditer(r'([A-Za-z0-9_:-]+)\s*=\s*("[^"]*"|<[^>]*>|[^,\n\r;]+)', text):
#         key = match.group(1).lower().strip()
#         value = match.group(2).strip().strip('"')
#         if value.startswith("<") and value.endswith(">"):
#             value = value[1:-1]
#         attrs[key] = value.strip()
#     return attrs


# def _extract_bracket_attr(statement: str) -> Tuple[str, Dict[str, str]]:
#     statement = html.unescape((statement or "").strip().rstrip(";").strip())
#     if "[" not in statement or "]" not in statement:
#         return statement, {}
#     start = statement.find("[")
#     end = statement.rfind("]")
#     if end <= start:
#         return statement, {}
#     return statement[:start].strip(), _extract_attr_pairs(statement[start + 1:end])


# def _unquote_id(value: str) -> str:
#     return str(value or "").strip().strip('"').strip()


# def _find_subgraph_blocks(dot_text: str) -> List[Tuple[str, str]]:
#     text = dot_text or ""
#     results: List[Tuple[str, str]] = []
#     idx = 0
#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             break
#         start = idx + match.start()
#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break
#         header = text[start:brace_start].strip()
#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1
#         if end >= len(text):
#             break
#         results.append((header, text[brace_start + 1:end]))
#         idx = end + 1
#     return results


# def _remove_subgraph_blocks(dot_text: str) -> str:
#     text = dot_text or ""
#     pieces: List[str] = []
#     idx = 0
#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             pieces.append(text[idx:])
#             break
#         start = idx + match.start()
#         pieces.append(text[idx:start])
#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break
#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1
#         if end >= len(text):
#             break
#         idx = end + 1
#     return "".join(pieces)


# def _extract_cluster_label(body: str, fallback_name: str) -> str:
#     for stmt in _split_statements(body):
#         s = html.unescape(stmt or "").strip()
#         if s.lower().startswith("label") and "=" in s:
#             try:
#                 value = s.split("=", 1)[1].strip().strip('"').strip()
#                 if value:
#                     return _clean_label(value)
#             except Exception:
#                 pass
#     return fallback_name


# def _extract_cluster_nodes(body: str) -> Set[str]:
#     node_ids: Set[str] = set()
#     for stmt in _split_statements(body):
#         s = html.unescape(stmt or "").strip()
#         if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s):
#             continue
#         prefix, attrs = _extract_bracket_attr(s)
#         node_id = _unquote_id(prefix)
#         label = attrs.get("label", node_id)
#         if not node_id or _is_pseudo_node(node_id, label):
#             continue
#         if re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
#             node_ids.add(node_id)
#     for match in EDGE_RE.finditer(body or ""):
#         for idx in [1, 3]:
#             node_id = _unquote_id(match.group(idx))
#             if node_id and not _is_pseudo_node(node_id):
#                 node_ids.add(node_id)
#     return node_ids


# def _parse_clusters_from_dot(dot: str) -> List[Dict[str, Any]]:
#     clusters: List[Dict[str, Any]] = []
#     for header, body in _find_subgraph_blocks(dot):
#         fallback_name = _unquote_id(header.replace("subgraph", "", 1).strip()) or "Cluster"
#         cluster_label = _extract_cluster_label(body, fallback_name)
#         cluster_nodes = sorted(_extract_cluster_nodes(body))
#         if cluster_nodes and not _is_pseudo_node(fallback_name, cluster_label):
#             clusters.append({"id": _safe_name(fallback_name), "label": cluster_label, "nodes": cluster_nodes})
#     return clusters


# def _parse_node_attrs_from_dot(dot: str) -> Tuple[Dict[str, Dict[str, str]], List[str]]:
#     attrs_by_node: Dict[str, Dict[str, str]] = {}
#     node_order: List[str] = []

#     def parse_node_statements(text: str):
#         for stmt in _split_statements(text):
#             s = html.unescape(stmt or "").strip()
#             if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s):
#                 continue
#             prefix, attrs = _extract_bracket_attr(s)
#             node_id = _unquote_id(prefix)
#             if not node_id:
#                 continue
#             label = attrs.get("label", node_id)
#             if _is_pseudo_node(node_id, label):
#                 continue
#             if not re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
#                 continue
#             if node_id not in attrs_by_node:
#                 node_order.append(node_id)
#             attrs_by_node[node_id] = attrs

#     body = _outer_graph_body(dot)
#     parse_node_statements(_remove_subgraph_blocks(body))
#     for _header, subgraph_body in _find_subgraph_blocks(dot):
#         parse_node_statements(subgraph_body)
#     return attrs_by_node, node_order


# def _parse_edge_attrs_from_dot(dot: str) -> Dict[Tuple[str, str], Dict[str, str]]:
#     attrs_by_edge: Dict[Tuple[str, str], Dict[str, str]] = {}

#     def parse_edge_statements(text: str):
#         for stmt in _split_statements(text):
#             s = html.unescape(stmt or "").strip()
#             if "->" not in s and "--" not in s:
#                 continue
#             prefix, attrs = _extract_bracket_attr(s)
#             tokens = re.findall(r'"[^"]+"|[A-Za-z0-9_.:-]+', prefix)
#             if len(tokens) < 2:
#                 continue
#             for src_raw, dst_raw in zip(tokens[:-1], tokens[1:]):
#                 src = _unquote_id(src_raw)
#                 dst = _unquote_id(dst_raw)
#                 if not src or not dst or _is_pseudo_node(src) or _is_pseudo_node(dst):
#                     continue
#                 attrs_by_edge[(src, dst)] = attrs

#     body = _outer_graph_body(dot)
#     parse_edge_statements(_remove_subgraph_blocks(body))
#     for _header, subgraph_body in _find_subgraph_blocks(dot):
#         parse_edge_statements(subgraph_body)

#     # Last-resort regex over full DOT. This avoids "half diagrams" if statement
#     # splitting fails because the model emitted unusual formatting.
#     for match in EDGE_RE.finditer(dot or ""):
#         src = _unquote_id(match.group(1))
#         dst = _unquote_id(match.group(3))
#         if src and dst and not _is_pseudo_node(src) and not _is_pseudo_node(dst):
#             attrs_by_edge.setdefault((src, dst), {})

#     return attrs_by_edge

# # -----------------------------------------------------------------------------
# # Model and layout
# # -----------------------------------------------------------------------------
# def _build_editable_model(diagram_text: str) -> Dict[str, Any]:
#     normalized = normalize_to_graphviz_dot(diagram_text)
#     if not normalized:
#         raise ValueError("Diagram text could not be normalized to DOT.")
#     dot = _strip_dot_comments(html.unescape(normalized))

#     clusters = _parse_clusters_from_dot(dot)
#     node_attrs, node_order = _parse_node_attrs_from_dot(dot)
#     edge_attrs = _parse_edge_attrs_from_dot(dot)

#     nodes: Dict[str, Dict[str, Any]] = {}
#     for node_id in node_order:
#         attrs = node_attrs.get(node_id, {})
#         label = _label(attrs.get("label"), node_id)
#         if _is_pseudo_node(node_id, label):
#             continue
#         nodes[node_id] = {"id": node_id, "label": label, "attrs": attrs, "icon_hint": _guess_icon_hint(node_id, label, attrs)}

#     edges: List[Dict[str, Any]] = []
#     for (src, dst), attrs in edge_attrs.items():
#         for node_id in [src, dst]:
#             if node_id not in nodes:
#                 label = _label(None, node_id)
#                 if not _is_pseudo_node(node_id, label):
#                     nodes[node_id] = {"id": node_id, "label": label, "attrs": {}, "icon_hint": _guess_icon_hint(node_id, label, {})}
#                     node_order.append(node_id)
#         if src in nodes and dst in nodes:
#             edges.append({"source": src, "target": dst, "raw_color": attrs.get("color") or attrs.get("fontcolor"), "label": _clean_label(attrs.get("label", ""))})

#     # If the diagram model contains nodes but edges were not parsed, infer a
#     # sequential left-to-right flow. This prevents vertical, tiny, half-rendered slides.
#     if not edges and len(node_order) > 1:
#         for src, dst in zip(node_order[:-1], node_order[1:]):
#             if src in nodes and dst in nodes:
#                 edges.append({"source": src, "target": dst, "raw_color": None, "label": ""})

#     cluster_by_node: Dict[str, str] = {}
#     clean_clusters: List[Dict[str, Any]] = []
#     for cluster in clusters:
#         cluster_nodes = [node_id for node_id in cluster["nodes"] if node_id in nodes]
#         if not cluster_nodes:
#             continue
#         clean_clusters.append({**cluster, "nodes": cluster_nodes})
#         for node_id in cluster_nodes:
#             cluster_by_node[node_id] = cluster["id"]

#     _assign_flow_colors(edges, nodes)
#     legend = _legend_from_edges(edges)
#     logger.info("Editable PPTX model parsed. nodes=%s edges=%s clusters=%s legend=%s", len(nodes), len(edges), len(clean_clusters), len(legend))
#     return {"nodes": [nodes[nid] for nid in node_order if nid in nodes], "node_map": nodes, "node_order": node_order, "edges": edges, "clusters": clean_clusters, "cluster_by_node": cluster_by_node, "legend": legend}


# def _assign_flow_colors(edges: List[Dict[str, Any]], nodes: Dict[str, Dict[str, Any]]):
#     used_by_label: Dict[str, str] = {}
#     palette_idx = 0
#     for edge in edges:
#         raw_color = _safe_color(edge.get("raw_color"), "") if edge.get("raw_color") else ""
#         flow_label = edge.get("label") or _edge_default_label(edge, nodes)
#         edge["display_label"] = flow_label
#         if raw_color:
#             color = raw_color
#         elif flow_label in used_by_label:
#             color = used_by_label[flow_label]
#         else:
#             color = FLOW_PALETTE[palette_idx % len(FLOW_PALETTE)]
#             used_by_label[flow_label] = color
#             palette_idx += 1
#         edge["color"] = color


# def _edge_default_label(edge: Dict[str, Any], nodes: Dict[str, Dict[str, Any]]) -> str:
#     src_label = _clean_label(nodes.get(edge["source"], {}).get("label", edge["source"])).replace("\n", " ")[:24]
#     dst_label = _clean_label(nodes.get(edge["target"], {}).get("label", edge["target"])).replace("\n", " ")[:24]
#     return f"{src_label} → {dst_label}"


# def _legend_from_edges(edges: List[Dict[str, Any]]) -> List[Tuple[str, str]]:
#     legend: List[Tuple[str, str]] = []
#     seen = set()
#     for edge in edges:
#         label = edge.get("display_label") or edge.get("label")
#         color = edge.get("color") or DEFAULT_EDGE_COLOR
#         if not label:
#             continue
#         key = (color, label)
#         if key not in seen:
#             seen.add(key)
#             legend.append(key)
#     return legend


# def _topological_order(node_order: List[str], edges: List[Dict[str, Any]]) -> List[str]:
#     ordered_known = [nid for nid in node_order]
#     indeg = {nid: 0 for nid in ordered_known}
#     outgoing: Dict[str, List[str]] = defaultdict(list)
#     for edge in edges:
#         src = edge["source"]
#         dst = edge["target"]
#         if src not in indeg or dst not in indeg:
#             continue
#         outgoing[src].append(dst)
#         indeg[dst] += 1
#     q = deque([nid for nid in ordered_known if indeg.get(nid, 0) == 0])
#     result: List[str] = []
#     while q:
#         cur = q.popleft()
#         result.append(cur)
#         for nxt in outgoing.get(cur, []):
#             indeg[nxt] -= 1
#             if indeg[nxt] == 0:
#                 q.append(nxt)
#     # Preserve any cyclic/missed nodes in original order.
#     for nid in ordered_known:
#         if nid not in result:
#             result.append(nid)
#     return result


# def _layout_model(model: Dict[str, Any]) -> Dict[str, Any]:
#     nodes_by_id = {node["id"]: node for node in model["nodes"]}
#     node_order = [nid for nid in model.get("node_order", []) if nid in nodes_by_id]
#     if not node_order:
#         node_order = [node["id"] for node in model["nodes"]]
#     ordered_ids = _topological_order(node_order, model["edges"])

#     n = max(1, len(ordered_ids))
#     cols = min(5, max(2, math.ceil(math.sqrt(n * 1.65))))
#     rows = math.ceil(n / cols)

#     node_w = min(NODE_WIDTH_IN, (CANVAS_WIDTH_IN - (cols - 1) * H_GAP_IN) / cols)
#     node_h = min(NODE_HEIGHT_IN, (CANVAS_HEIGHT_IN - (rows - 1) * V_GAP_IN - 0.25) / rows)
#     node_w = max(1.25, node_w)
#     node_h = max(0.58, node_h)

#     total_w = cols * node_w + (cols - 1) * H_GAP_IN
#     total_h = rows * node_h + (rows - 1) * V_GAP_IN
#     start_x = CANVAS_LEFT_IN + max(0.0, (CANVAS_WIDTH_IN - total_w) / 2.0)
#     start_y = CANVAS_TOP_IN + 0.20 + max(0.0, (CANVAS_HEIGHT_IN - total_h - 0.35) / 2.0)

#     placements: Dict[str, Dict[str, float]] = {}
#     for idx, node_id in enumerate(ordered_ids):
#         row = idx // cols
#         col = idx % cols
#         placements[node_id] = {"left": start_x + col * (node_w + H_GAP_IN), "top": start_y + row * (node_h + V_GAP_IN), "width": node_w, "height": node_h}

#     cluster_boxes = _cluster_boxes_from_placements(model["clusters"], placements)
#     return {"placements": placements, "cluster_boxes": cluster_boxes}


# def _cluster_boxes_from_placements(clusters: List[Dict[str, Any]], placements: Dict[str, Dict[str, float]]) -> List[Dict[str, Any]]:
#     boxes: List[Dict[str, Any]] = []
#     for cluster in clusters:
#         member_boxes = [placements[nid] for nid in cluster.get("nodes", []) if nid in placements]
#         if not member_boxes:
#             continue
#         left = min(b["left"] for b in member_boxes) - CLUSTER_PAD_IN
#         top = min(b["top"] for b in member_boxes) - CLUSTER_TITLE_HEIGHT_IN - 0.08
#         right = max(b["left"] + b["width"] for b in member_boxes) + CLUSTER_PAD_IN
#         bottom = max(b["top"] + b["height"] for b in member_boxes) + CLUSTER_PAD_IN
#         boxes.append({"id": cluster["id"], "label": cluster.get("label", cluster["id"]), "left": max(CANVAS_LEFT_IN, left), "top": max(CANVAS_TOP_IN, top), "width": max(0.8, right - left), "height": max(0.55, bottom - top)})
#     return boxes

# # -----------------------------------------------------------------------------
# # PPTX drawing helpers
# # -----------------------------------------------------------------------------
# def _set_text_frame(shape: BaseShape, text: str, font_size_pt: float = 10, bold: bool = False, color_hex: str = DEFAULT_NODE_FONT, align=PP_ALIGN.LEFT):
#     tf = shape.text_frame
#     tf.clear()
#     tf.word_wrap = True
#     tf.margin_left = Inches(0.03)
#     tf.margin_right = Inches(0.03)
#     tf.margin_top = Inches(0.01)
#     tf.margin_bottom = Inches(0.01)
#     tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
#     paragraph = tf.paragraphs[0]
#     paragraph.alignment = align
#     run = paragraph.add_run()
#     run.text = _safe_title(text)
#     font = run.font
#     font.name = "Calibri"
#     font.size = Pt(font_size_pt)
#     font.bold = bold
#     font.color.rgb = _rgb(color_hex, DEFAULT_NODE_FONT)


# def _font_for_label(label: str, width: float, height: float) -> float:
#     clean = _safe_title(label)
#     line_count = max(1, clean.count("\n") + 1)
#     length = len(clean.replace("\n", " "))
#     if width < 1.20 or height < 0.58 or length > 62 or line_count >= 3:
#         return 7.0
#     if width < 1.55 or length > 44 or line_count == 2:
#         return 8.0
#     return 9.0


# def _add_title(slide, title_text: str):
#     shape = slide.shapes.add_textbox(Inches(TITLE_LEFT_IN), Inches(TITLE_TOP_IN), Inches(TITLE_WIDTH_IN), Inches(TITLE_HEIGHT_IN))
#     _set_text_frame(shape, title_text, font_size_pt=18, bold=True, color_hex="202124")


# def _add_cluster_box(slide, box: Dict[str, Any]):
#     shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
#     shape.fill.solid()
#     shape.fill.fore_color.rgb = _rgb(DEFAULT_CLUSTER_FILL)
#     shape.line.color.rgb = _rgb(DEFAULT_CLUSTER_BORDER)
#     shape.line.width = Pt(1.0)
#     try:
#         shape.adjustments[0] = 0.06
#     except Exception:
#         pass
#     label_box = slide.shapes.add_textbox(Inches(box["left"] + 0.06), Inches(box["top"] + 0.02), Inches(max(0.5, box["width"] - 0.12)), Inches(0.20))
#     _set_text_frame(label_box, box.get("label", ""), font_size_pt=8.4, bold=True, color_hex=DEFAULT_CLUSTER_FONT)


# def _add_node_box(slide, box: Dict[str, float], label: str, icon_path: Optional[Path] = None):
#     shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
#     shape.fill.solid()
#     shape.fill.fore_color.rgb = _rgb(DEFAULT_NODE_FILL)
#     shape.line.color.rgb = _rgb(DEFAULT_NODE_BORDER)
#     shape.line.width = Pt(1.0)
#     try:
#         shape.adjustments[0] = 0.08
#     except Exception:
#         pass

#     icon_size = min(NODE_ICON_SIZE_IN, max(0.15, box["height"] * 0.38))
#     text_left = box["left"] + NODE_PADDING_IN
#     if icon_path and icon_path.exists() and box["width"] > 0.75:
#         try:
#             slide.shapes.add_picture(str(icon_path), Inches(box["left"] + NODE_PADDING_IN), Inches(box["top"] + (box["height"] - icon_size) / 2.0), width=Inches(icon_size), height=Inches(icon_size))
#             text_left = box["left"] + NODE_TEXT_LEFT_WITH_ICON_IN
#         except Exception:
#             logger.debug("Could not add icon to PPTX node: %s", icon_path, exc_info=True)

#     text_width = max(0.35, box["width"] - (text_left - box["left"]) - 0.06)
#     text_box = slide.shapes.add_textbox(Inches(text_left), Inches(box["top"] + 0.03), Inches(text_width), Inches(max(0.22, box["height"] - 0.06)))
#     _set_text_frame(text_box, label, font_size_pt=_font_for_label(label, box["width"], box["height"]), bold=True, color_hex=DEFAULT_NODE_FONT)


# def _edge_points_between_boxes(src_box: Dict[str, float], dst_box: Dict[str, float]) -> List[Tuple[float, float]]:
#     src_right = src_box["left"] + src_box["width"]
#     src_left = src_box["left"]
#     src_y = src_box["top"] + src_box["height"] / 2.0
#     dst_left = dst_box["left"]
#     dst_right = dst_box["left"] + dst_box["width"]
#     dst_y = dst_box["top"] + dst_box["height"] / 2.0
#     if dst_left >= src_right:
#         mid_x = (src_right + dst_left) / 2.0
#         return [(src_right, src_y), (mid_x, src_y), (mid_x, dst_y), (dst_left, dst_y)]
#     if src_left >= dst_right:
#         mid_x = (dst_right + src_left) / 2.0
#         return [(src_left, src_y), (mid_x, src_y), (mid_x, dst_y), (dst_right, dst_y)]
#     if dst_y >= src_y:
#         return [(src_box["left"] + src_box["width"] / 2.0, src_box["top"] + src_box["height"]), (dst_box["left"] + dst_box["width"] / 2.0, dst_box["top"])]
#     return [(src_box["left"] + src_box["width"] / 2.0, src_box["top"]), (dst_box["left"] + dst_box["width"] / 2.0, dst_box["top"] + dst_box["height"])]


# def _add_line_segment(slide, p1: Tuple[float, float], p2: Tuple[float, float], color_hex: str):
#     connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p1[0]), Inches(p1[1]), Inches(p2[0]), Inches(p2[1]))
#     connector.line.color.rgb = _rgb(color_hex, DEFAULT_EDGE_COLOR)
#     connector.line.width = Pt(DEFAULT_EDGE_WIDTH_PT)
#     return connector


# def _add_polyline_arrow(slide, points: List[Tuple[float, float]], color_hex: str):
#     if len(points) < 2:
#         return
#     clean_points = [points[0]]
#     for point in points[1:]:
#         if abs(point[0] - clean_points[-1][0]) > 0.01 or abs(point[1] - clean_points[-1][1]) > 0.01:
#             clean_points.append(point)
#     if len(clean_points) < 2:
#         return
#     for idx in range(len(clean_points) - 1):
#         segment = _add_line_segment(slide, clean_points[idx], clean_points[idx + 1], color_hex)
#         if idx == len(clean_points) - 2:
#             try:
#                 segment.line.end_arrowhead = True
#             except Exception:
#                 pass


# def _add_legend(slide, legend_items: List[Tuple[str, str]]):
#     if not legend_items:
#         return
#     title_box = slide.shapes.add_textbox(Inches(LEGEND_LEFT_IN), Inches(LEGEND_TOP_IN), Inches(1.15), Inches(0.18))
#     _set_text_frame(title_box, "Flow Legend", font_size_pt=9.2, bold=True, color_hex="202124")
#     x = LEGEND_LEFT_IN
#     y = LEGEND_TOP_IN + 0.24
#     max_x = LEGEND_LEFT_IN + LEGEND_WIDTH_IN
#     for color, label in legend_items[:18]:
#         label = _safe_title(label)
#         item_width = max(1.25, min(3.15, 0.42 + len(label) * 0.050))
#         if x + item_width > max_x:
#             x = LEGEND_LEFT_IN
#             y += 0.22
#         if y > LEGEND_TOP_IN + LEGEND_MAX_HEIGHT_IN:
#             break
#         chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(item_width), Inches(0.18))
#         chip.fill.solid()
#         chip.fill.fore_color.rgb = _rgb(LEGEND_FILL)
#         chip.line.color.rgb = _rgb(LEGEND_BORDER)
#         chip.line.width = Pt(0.6)
#         dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.04), Inches(y + 0.045), Inches(0.09), Inches(0.09))
#         dot.fill.solid()
#         dot.fill.fore_color.rgb = _rgb(color, DEFAULT_EDGE_COLOR)
#         dot.line.color.rgb = _rgb(color, DEFAULT_EDGE_COLOR)
#         label_box = slide.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.01), Inches(max(0.3, item_width - 0.18)), Inches(0.15))
#         _set_text_frame(label_box, label, font_size_pt=7.3, bold=False, color_hex="202124")
#         x += item_width + 0.08


# def _add_fallback_preview_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str):
#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, f"{slide_title} (Preview)")
#     try:
#         assets = render_graphviz_to_assets(diagram_text, output_dir=output_dir, base_name=base_name)
#     except Exception:
#         logger.exception("Fallback preview asset generation failed.")
#         assets = None
#     if assets and assets.get("png_path") and Path(assets["png_path"]).exists():
#         slide.shapes.add_picture(str(assets["png_path"]), Inches(CANVAS_LEFT_IN), Inches(CANVAS_TOP_IN), width=Inches(PREVIEW_MAX_WIDTH_IN))
#         _add_legend(slide, assets.get("legend", []))

# # -----------------------------------------------------------------------------
# # Diagram extraction
# # -----------------------------------------------------------------------------
# def _is_diagram_text(value: Any) -> bool:
#     return isinstance(value, str) and normalize_to_graphviz_dot(_clean_internal_tokens(value)) is not None


# def _walk_for_diagrams(value: Any, path_parts: List[str], results: List[Dict[str, Any]]):
#     if value is None:
#         return
#     if isinstance(value, str):
#         if _is_diagram_text(value):
#             results.append({"title": " / ".join(path_parts) if path_parts else "Diagram", "diagram_text": _clean_internal_tokens(value), "section_key": _safe_name(path_parts[-1] if path_parts else "diagram")})
#         return
#     if isinstance(value, list):
#         for idx, item in enumerate(value, start=1):
#             if _is_diagram_text(item):
#                 results.append({"title": " / ".join(path_parts) if path_parts else f"Diagram {idx}", "diagram_text": _clean_internal_tokens(item), "section_key": _safe_name(path_parts[-1] if path_parts else f"diagram_{idx}")})
#             elif isinstance(item, (dict, list)):
#                 _walk_for_diagrams(item, path_parts, results)
#         return
#     if isinstance(value, dict):
#         for key, child in value.items():
#             key_text = str(key).replace("_", " ").title()
#             next_path = path_parts if str(key).lower() == "diagrams" else path_parts + [key_text]
#             _walk_for_diagrams(child, next_path, results)


# def extract_diagram_specs(data: Dict[str, Any], hld_model: Any, selected_sections_set: Optional[Set[str]] = None) -> List[Dict[str, Any]]:
#     results: List[Dict[str, Any]] = []
#     for field_name, field_info in hld_model.model_fields.items():
#         if not should_include_section(field_name, field_info, selected_sections_set or set()):
#             continue
#         section_title = field_info.title or field_name.replace("_", " ").title()
#         _walk_for_diagrams(data.get(field_name), [section_title], results)
#     title_counts: Dict[str, int] = defaultdict(int)
#     for item in results:
#         title_counts[item["title"]] += 1
#     dedupe_counter: Dict[str, int] = defaultdict(int)
#     for item in results:
#         if title_counts[item["title"]] > 1:
#             dedupe_counter[item["title"]] += 1
#             item["title"] = f'{item["title"]} ({dedupe_counter[item["title"]]})'
#     return results

# # -----------------------------------------------------------------------------
# # Slide rendering
# # -----------------------------------------------------------------------------
# def add_diagram_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str):
#     try:
#         model = _build_editable_model(diagram_text)
#         if not model["nodes"]:
#             raise ValueError("No editable nodes parsed")
#         layout = _layout_model(model)
#         placements = layout["placements"]
#         cluster_boxes = layout["cluster_boxes"]
#         if not placements:
#             raise ValueError("No placements generated")
#     except Exception:
#         logger.exception("Editable PPTX rendering failed for slide '%s'. Falling back to preview image.", slide_title)
#         _add_fallback_preview_slide(prs, slide_title, diagram_text, output_dir, base_name)
#         return

#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, slide_title)
#     for box in cluster_boxes:
#         _add_cluster_box(slide, box)
#     for edge in model["edges"]:
#         src_box = placements.get(edge["source"])
#         dst_box = placements.get(edge["target"])
#         if not src_box or not dst_box:
#             continue
#         _add_polyline_arrow(slide, _edge_points_between_boxes(src_box, dst_box), edge.get("color") or DEFAULT_EDGE_COLOR)
#     for node in model["nodes"]:
#         box = placements.get(node["id"])
#         if not box:
#             continue
#         icon_path = _preferred_picture_icon_path(node.get("icon_hint") or node.get("label") or node["id"])
#         _add_node_box(slide, box, node.get("label") or node["id"], icon_path)
#     _add_legend(slide, model.get("legend", []))

# # -----------------------------------------------------------------------------
# # Public builder
# # -----------------------------------------------------------------------------
# def build_editable_pptx(data: Dict[str, Any], pptx_path: str, hld_model: Any, selected_sections_set: Optional[Set[str]] = None):
#     output_path = Path(pptx_path)
#     output_path.parent.mkdir(parents=True, exist_ok=True)
#     diagram_assets_dir = output_path.parent / f"{output_path.stem}_diagram_assets"
#     diagram_assets_dir.mkdir(parents=True, exist_ok=True)
#     diagram_specs = extract_diagram_specs(data=data, hld_model=hld_model, selected_sections_set=selected_sections_set)
#     prs = Presentation()
#     prs.slide_width = Inches(SLIDE_WIDTH_IN)
#     prs.slide_height = Inches(SLIDE_HEIGHT_IN)
#     if not diagram_specs:
#         slide = prs.slides.add_slide(prs.slide_layouts[6])
#         _add_title(slide, "Editable Diagrams")
#         msg = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11.5), Inches(1.0))
#         _set_text_frame(msg, "No diagrams were detected in the selected sections.", font_size_pt=16, bold=True, color_hex="5F6368", align=PP_ALIGN.CENTER)
#         prs.save(str(output_path))
#         return
#     for counter, spec in enumerate(diagram_specs, start=1):
#         base_name = f"{_safe_name(spec['section_key'])}_{counter:02d}"
#         add_diagram_slide(prs=prs, slide_title=spec["title"], diagram_text=spec["diagram_text"], output_dir=diagram_assets_dir, base_name=base_name)
#     prs.save(str(output_path))
#     logger.info("Editable PPTX generated: %s", output_path)

#3 rd-party imports


# from __future__ import annotations

# import html
# import logging
# import math
# import re
# from collections import defaultdict, deque
# from pathlib import Path
# from typing import Any, Dict, List, Optional, Set, Tuple

# from pptx import Presentation
# from pptx.dml.color import RGBColor
# from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
# from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
# from pptx.shapes.base import BaseShape
# from pptx.util import Inches, Pt
# from pptx.oxml.xmlchemy import OxmlElement

# from .config import ICON_DIR
# from .graphviz_renderer import normalize_to_graphviz_dot, render_graphviz_to_assets
# from .schema_utils import should_include_section
# from .text_sanitizer import renderer_safe_plain_text

# logger = logging.getLogger(__name__)

# # =============================================================================
# # Editable PPTX Diagram Renderer - swimlane version
# # =============================================================================
# # Renders editable PowerPoint diagrams with grouped zone boxes similar to PDF:
# #   - On-Premise box
# #   - GCP box
# #   - External / CI-CD / Observability boxes when applicable
# #   - editable node cards with icons
# #   - colored routed arrows with arrowheads
# #   - matching Flow Legend
# # =============================================================================

# SLIDE_WIDTH_IN = 13.333
# SLIDE_HEIGHT_IN = 7.5
# TITLE_LEFT_IN = 0.35
# TITLE_TOP_IN = 0.12
# TITLE_WIDTH_IN = 12.6
# TITLE_HEIGHT_IN = 0.42
# CANVAS_LEFT_IN = 0.35
# CANVAS_TOP_IN = 0.68
# CANVAS_WIDTH_IN = 12.3
# CANVAS_HEIGHT_IN = 5.25
# LEGEND_LEFT_IN = 0.35
# LEGEND_TOP_IN = 6.12
# LEGEND_WIDTH_IN = 12.3
# LEGEND_MAX_HEIGHT_IN = 1.22
# NODE_W = 1.78
# NODE_H = 0.78
# NODE_ICON_SIZE_IN = 0.22
# NODE_PADDING_IN = 0.07
# NODE_TEXT_LEFT_WITH_ICON_IN = 0.33
# LANE_PAD = 0.18
# LANE_TITLE_H = 0.26
# H_GAP = 0.28
# V_GAP = 0.28
# DEFAULT_NODE_FILL = "F8F9FA"
# DEFAULT_NODE_BORDER = "DADCE0"
# DEFAULT_NODE_FONT = "202124"
# LANE_FILL = "F3F8FF"
# LANE_BORDER = "A8C7FA"
# LANE_TITLE_COLOR = "174EA6"
# EXTERNAL_FILL = "FFF8E1"
# EXTERNAL_BORDER = "F9AB00"
# CICD_FILL = "FCE8E6"
# CICD_BORDER = "F28B82"
# OBS_FILL = "F3E8FD"
# OBS_BORDER = "D7AEFB"
# DEFAULT_EDGE_COLOR = "5F6368"
# DEFAULT_EDGE_WIDTH_PT = 1.25
# LEGEND_FILL = "FFFFFF"
# LEGEND_BORDER = "DADCE0"
# PREVIEW_MAX_WIDTH_IN = 12.0
# SUPPORTED_PICTURE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif"}
# PREFERRED_ICON_EXTS = [".png", ".jpg", ".jpeg", ".bmp", ".gif"]
# COMMENT_LINE_RE = re.compile(r"//.*?$", re.MULTILINE)
# BLOCK_COMMENT_RE = re.compile(r"/\*.*?\*/", re.DOTALL)
# EDGE_RE = re.compile(r'("[^"]+"|[A-Za-z0-9_.:-]+)\s*(->|--)\s*("[^"]+"|[A-Za-z0-9_.:-]+)')
# FLOW_PALETTE = ["E60000", "1D70B8", "28A197", "F47738", "6F35A5", "007C89", "D53880", "85994B", "B58840", "5F6368"]
# COLOR_NAME_MAP = {"red":"E60000","blue":"1D70B8","green":"28A197","orange":"F47738","purple":"6F35A5","teal":"007C89","gray":"5F6368","grey":"5F6368","black":"202124","lightgrey":"DADCE0","lightgray":"DADCE0"}
# ICON_KEYWORD_MAP = [
#     ("cloud storage", "cloud_storage"), ("gcs", "cloud_storage"), ("storage", "cloud_storage"),
#     ("cloud function", "cloud_functions"), ("cloud functions", "cloud_functions"), ("function", "cloud_functions"),
#     ("pub/sub", "pubsub"), ("pubsub", "pubsub"), ("notification", "pubsub"), ("eventarc", "eventarc"),
#     ("vault", "secret_manager"), ("hashicorp", "secret_manager"), ("secret", "secret_manager"),
#     ("kms", "key_management_service"), ("key management", "key_management_service"),
#     ("logging", "cloud_logging"), ("monitoring", "cloud_monitoring"),
#     ("cloud build", "cloud_build"), ("build", "cloud_build"), ("terraform", "terraform"),
#     ("teradata", "database_migration_service"), ("database", "database_migration_service"),
#     ("mft", "transfer"), ("managed file transfer", "transfer"), ("transfer", "transfer"),
#     ("interconnect", "cloud_interconnect"), ("vpn", "cloud_vpn"), ("collibra", "catalog"), ("catalog", "catalog"),
# ]


# def _clean_internal_tokens(text: Any) -> str:
#     value = str(text or "")
#     for token in ["AIASECTIONBLOCKSTARTTOKEN","AIASECTIONBLOCKENDTOKEN","[[AIASECTIONBLOCKSTARTTOKEN]]","[[AIASECTIONBLOCKENDTOKEN]]","SECTION_BLOCK_START","SECTION_BLOCK_END","[[SECTION_BLOCK_START]]","[[SECTION_BLOCK_END]]","[[SECTIONBLOCKSTART]]","[[SECTIONBLOCKEND]]"]:
#         value = value.replace(token, "")
#     return html.unescape(value.strip())


# def _safe_title(text: Any) -> str:
#     return renderer_safe_plain_text(_clean_internal_tokens(text)).strip() or "Diagram"


# def _safe_name(text: str) -> str:
#     return re.sub(r"[^a-zA-Z0-9_-]+", "_", str(text or "")).strip("_").lower() or "diagram"


# def _clean_label(value: Any) -> str:
#     text = html.unescape(str(value or "").strip()).strip('"')
#     text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
#     text = re.sub(r"<[^>]+>", " ", text)
#     text = text.replace("\\N", "\n").replace("\\n", "\n").replace("&nbsp;", " ")
#     text = re.sub(r"\s+\n", "\n", text)
#     text = re.sub(r"\n\s+", "\n", text)
#     text = re.sub(r"[ \t]+", " ", text)
#     return text.strip()


# def _label(raw: Optional[str], fallback: str) -> str:
#     label = _clean_label(raw or "")
#     if not label:
#         label = str(fallback or "").replace("_", " ").replace("-", " ").title()
#     return label[:180]


# def _safe_color(value: Optional[str], default: str = DEFAULT_EDGE_COLOR) -> str:
#     raw = str(value or "").strip().strip('"').lstrip("#")
#     if not raw:
#         return default
#     if raw.lower() in COLOR_NAME_MAP:
#         return COLOR_NAME_MAP[raw.lower()]
#     raw = raw.upper()
#     return raw if len(raw) == 6 and all(ch in "0123456789ABCDEF" for ch in raw) else default


# def _rgb(hex_color: str, default: str = "000000") -> RGBColor:
#     return RGBColor.from_string(_safe_color(hex_color, default))


# def _is_pseudo_node(node_id: str, label: Optional[str] = None) -> bool:
#     text = _clean_label(label if label is not None else node_id).lower().strip()
#     raw = str(node_id or "").lower().strip()
#     if text in {"", "n/a", "na", "none", "null"}:
#         return True
#     if raw in {"node", "edge", "graph", "digraph", "subgraph", "rank", "label", "style", "color", "fillcolor", "fontcolor", "rankdir"}:
#         return True
#     if raw.startswith(("digraph", "graph", "subgraph", "cluster_", "cluster ", "label=", "style=", "color=", "fillcolor=", "fontcolor=", "rankdir=")):
#         return True
#     if text.startswith(("label=", "style=", "color=", "fillcolor=", "fontcolor=", "rankdir=", "fontsize=", "fontname=", "margin=", "pad=")):
#         return True
#     if text in {"filled", "lightgrey", "lightgray", "blue", "green", "purple", "orange", "red", "gray", "grey"}:
#         return True
#     return False


# def _is_graphviz_control_statement(statement: str) -> bool:
#     s = html.unescape(str(statement or "")).strip().lower()
#     if not s or s in {"{", "}"}:
#         return True
#     return s.startswith(("digraph ", "graph ", "subgraph ", "node ", "edge ", "rank ", "rankdir", "label", "style", "color", "fillcolor", "fontcolor", "fontsize", "fontname", "margin", "pad", "splines", "bgcolor"))


# def _try_existing_icon_resolver(icon_hint: str) -> Optional[Path]:
#     try:
#         from . import icon_resolver as ir
#     except Exception:
#         return None
#     for fn_name in ["resolve_icon_path", "resolve_icon", "get_icon_path", "find_icon_path"]:
#         fn = getattr(ir, fn_name, None)
#         if not callable(fn):
#             continue
#         for args in [(icon_hint,), (icon_hint, True), (icon_hint, None)]:
#             try:
#                 result = fn(*args)
#                 if result:
#                     path = Path(str(result))
#                     if path.exists():
#                         return path
#             except TypeError:
#                 continue
#             except Exception:
#                 logger.debug("Icon resolver failed: %s", fn_name, exc_info=True)
#     return None


# def _search_icon_dir(icon_hint: str) -> Optional[Path]:
#     if not icon_hint:
#         return None
#     hint = _safe_name(icon_hint)
#     bases = {hint, icon_hint, icon_hint.lower(), icon_hint.replace(" ", "_").lower(), icon_hint.replace("-", "_").lower(), icon_hint.replace("/", "_").lower()}
#     for base in bases:
#         if not base:
#             continue
#         for ext in [".png", ".jpg", ".jpeg", ".bmp", ".gif", ".svg"]:
#             candidate = ICON_DIR / f"{base}{ext}"
#             if candidate.exists():
#                 return candidate
#     try:
#         for path in ICON_DIR.rglob("*"):
#             if path.is_file() and hint in path.stem.lower():
#                 return path
#     except Exception:
#         logger.debug("Icon directory search failed", exc_info=True)
#     return None


# def _keyword_icon_hint(label_or_id: str) -> Optional[str]:
#     text = _clean_label(label_or_id).lower()
#     for keyword, icon_hint in ICON_KEYWORD_MAP:
#         if keyword in text:
#             return icon_hint
#     return None


# def _preferred_picture_icon_path(icon_hint: str) -> Optional[Path]:
#     candidates: List[str] = []
#     keyword = _keyword_icon_hint(icon_hint)
#     if keyword:
#         candidates.append(keyword)
#     if icon_hint:
#         candidates.append(icon_hint)
#     for candidate in candidates:
#         path = _try_existing_icon_resolver(candidate) or _search_icon_dir(candidate)
#         if not path:
#             continue
#         if path.suffix.lower() in SUPPORTED_PICTURE_EXTS:
#             return path
#         if path.suffix.lower() == ".svg":
#             for ext in PREFERRED_ICON_EXTS:
#                 sibling = path.with_suffix(ext)
#                 if sibling.exists():
#                     return sibling
#     return None


# def _guess_icon_hint(node_id: str, node_label: str, attrs: Dict[str, str]) -> str:
#     image_attr = attrs.get("image") or attrs.get("icon") or attrs.get("imagepath")
#     if image_attr:
#         stem = Path(str(image_attr)).stem
#         if stem:
#             return stem
#     keyword = _keyword_icon_hint(node_label) or _keyword_icon_hint(node_id)
#     if keyword:
#         return keyword
#     for candidate in [node_id, node_label.split("\n", 1)[0], node_label.replace("\n", " ")]:
#         cleaned = _safe_name(candidate)
#         if cleaned:
#             return cleaned
#     return _safe_name(node_id or node_label or "node")


# def _strip_dot_comments(dot_text: str) -> str:
#     text = BLOCK_COMMENT_RE.sub("", dot_text or "")
#     return COMMENT_LINE_RE.sub("", text)


# def _outer_graph_body(dot: str) -> str:
#     text = dot or ""
#     start = text.find("{")
#     end = text.rfind("}")
#     if start != -1 and end != -1 and end > start:
#         return text[start + 1:end]
#     return text


# def _split_statements(dot_text: str) -> List[str]:
#     statements: List[str] = []
#     buf: List[str] = []
#     bracket_depth = brace_depth = angle_depth = 0
#     for ch in dot_text or "":
#         if ch == "[":
#             bracket_depth += 1
#         elif ch == "]":
#             bracket_depth = max(0, bracket_depth - 1)
#         elif ch == "{":
#             brace_depth += 1
#         elif ch == "}":
#             brace_depth = max(0, brace_depth - 1)
#         elif ch == "<":
#             angle_depth += 1
#         elif ch == ">":
#             angle_depth = max(0, angle_depth - 1)
#         if ch == ";" and bracket_depth == 0 and brace_depth == 0 and angle_depth == 0:
#             stmt = "".join(buf).strip()
#             if stmt:
#                 statements.append(stmt)
#             buf = []
#         else:
#             buf.append(ch)
#     tail = "".join(buf).strip()
#     if tail:
#         statements.append(tail)
#     return statements


# def _extract_attr_pairs(attr_text: str) -> Dict[str, str]:
#     attrs: Dict[str, str] = {}
#     if not attr_text:
#         return attrs
#     text = html.unescape(str(attr_text))
#     for match in re.finditer(r'([A-Za-z0-9_:-]+)\s*=\s*("[^"]*"|<[^>]*>|[^,\n\r;]+)', text):
#         key = match.group(1).lower().strip()
#         value = match.group(2).strip().strip('"')
#         if value.startswith("<") and value.endswith(">"):
#             value = value[1:-1]
#         attrs[key] = value.strip()
#     return attrs


# def _extract_bracket_attr(statement: str) -> Tuple[str, Dict[str, str]]:
#     statement = html.unescape((statement or "").strip().rstrip(";").strip())
#     if "[" not in statement or "]" not in statement:
#         return statement, {}
#     start = statement.find("[")
#     end = statement.rfind("]")
#     if end <= start:
#         return statement, {}
#     return statement[:start].strip(), _extract_attr_pairs(statement[start + 1:end])


# def _unquote_id(value: str) -> str:
#     return str(value or "").strip().strip('"').strip()


# def _find_subgraph_blocks(dot_text: str) -> List[Tuple[str, str]]:
#     text = dot_text or ""
#     results: List[Tuple[str, str]] = []
#     idx = 0
#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             break
#         start = idx + match.start()
#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break
#         header = text[start:brace_start].strip()
#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1
#         if end >= len(text):
#             break
#         results.append((header, text[brace_start + 1:end]))
#         idx = end + 1
#     return results


# def _remove_subgraph_blocks(dot_text: str) -> str:
#     text = dot_text or ""
#     pieces: List[str] = []
#     idx = 0
#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             pieces.append(text[idx:])
#             break
#         start = idx + match.start()
#         pieces.append(text[idx:start])
#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break
#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1
#         if end >= len(text):
#             break
#         idx = end + 1
#     return "".join(pieces)


# def _extract_cluster_label(body: str, fallback_name: str) -> str:
#     for stmt in _split_statements(body):
#         s = html.unescape(stmt or "").strip()
#         if s.lower().startswith("label") and "=" in s:
#             try:
#                 value = s.split("=", 1)[1].strip().strip('"').strip()
#                 if value:
#                     return _clean_label(value)
#             except Exception:
#                 pass
#     return fallback_name


# def _extract_cluster_nodes(body: str) -> Set[str]:
#     node_ids: Set[str] = set()
#     for stmt in _split_statements(body):
#         s = html.unescape(stmt or "").strip()
#         if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s):
#             continue
#         prefix, attrs = _extract_bracket_attr(s)
#         node_id = _unquote_id(prefix)
#         label = attrs.get("label", node_id)
#         if not node_id or _is_pseudo_node(node_id, label):
#             continue
#         if re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
#             node_ids.add(node_id)
#     for match in EDGE_RE.finditer(body or ""):
#         for idx in [1, 3]:
#             node_id = _unquote_id(match.group(idx))
#             if node_id and not _is_pseudo_node(node_id):
#                 node_ids.add(node_id)
#     return node_ids


# def _parse_clusters_from_dot(dot: str) -> List[Dict[str, Any]]:
#     clusters: List[Dict[str, Any]] = []
#     for header, body in _find_subgraph_blocks(dot):
#         fallback_name = _unquote_id(header.replace("subgraph", "", 1).strip()) or "Cluster"
#         cluster_label = _extract_cluster_label(body, fallback_name)
#         cluster_nodes = sorted(_extract_cluster_nodes(body))
#         if cluster_nodes and not _is_pseudo_node(fallback_name, cluster_label):
#             clusters.append({"id": _safe_name(fallback_name), "label": cluster_label, "nodes": cluster_nodes})
#     return clusters


# def _parse_node_attrs_from_dot(dot: str) -> Tuple[Dict[str, Dict[str, str]], List[str]]:
#     attrs_by_node: Dict[str, Dict[str, str]] = {}
#     node_order: List[str] = []

#     def parse_node_statements(text: str):
#         for stmt in _split_statements(text):
#             s = html.unescape(stmt or "").strip()
#             if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s):
#                 continue
#             prefix, attrs = _extract_bracket_attr(s)
#             node_id = _unquote_id(prefix)
#             if not node_id:
#                 continue
#             label = attrs.get("label", node_id)
#             if _is_pseudo_node(node_id, label):
#                 continue
#             if not re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
#                 continue
#             if node_id not in attrs_by_node:
#                 node_order.append(node_id)
#             attrs_by_node[node_id] = attrs

#     body = _outer_graph_body(dot)
#     parse_node_statements(_remove_subgraph_blocks(body))
#     for _header, subgraph_body in _find_subgraph_blocks(dot):
#         parse_node_statements(subgraph_body)
#     return attrs_by_node, node_order


# def _parse_edge_attrs_from_dot(dot: str) -> Dict[Tuple[str, str], Dict[str, str]]:
#     attrs_by_edge: Dict[Tuple[str, str], Dict[str, str]] = {}

#     def parse_edge_statements(text: str):
#         for stmt in _split_statements(text):
#             s = html.unescape(stmt or "").strip()
#             if "->" not in s and "--" not in s:
#                 continue
#             prefix, attrs = _extract_bracket_attr(s)
#             tokens = re.findall(r'"[^"]+"|[A-Za-z0-9_.:-]+', prefix)
#             if len(tokens) < 2:
#                 continue
#             for src_raw, dst_raw in zip(tokens[:-1], tokens[1:]):
#                 src = _unquote_id(src_raw)
#                 dst = _unquote_id(dst_raw)
#                 if not src or not dst or _is_pseudo_node(src) or _is_pseudo_node(dst):
#                     continue
#                 attrs_by_edge[(src, dst)] = attrs

#     body = _outer_graph_body(dot)
#     parse_edge_statements(_remove_subgraph_blocks(body))
#     for _header, subgraph_body in _find_subgraph_blocks(dot):
#         parse_edge_statements(subgraph_body)
#     for match in EDGE_RE.finditer(dot or ""):
#         src = _unquote_id(match.group(1))
#         dst = _unquote_id(match.group(3))
#         if src and dst and not _is_pseudo_node(src) and not _is_pseudo_node(dst):
#             attrs_by_edge.setdefault((src, dst), {})
#     return attrs_by_edge

# # -----------------------------------------------------------------------------
# # Model
# # -----------------------------------------------------------------------------
# def _build_editable_model(diagram_text: str) -> Dict[str, Any]:
#     normalized = normalize_to_graphviz_dot(diagram_text)
#     if not normalized:
#         raise ValueError("Diagram text could not be normalized to DOT.")
#     dot = _strip_dot_comments(html.unescape(normalized))
#     clusters = _parse_clusters_from_dot(dot)
#     node_attrs, node_order = _parse_node_attrs_from_dot(dot)
#     edge_attrs = _parse_edge_attrs_from_dot(dot)

#     nodes: Dict[str, Dict[str, Any]] = {}
#     for node_id in node_order:
#         attrs = node_attrs.get(node_id, {})
#         label = _label(attrs.get("label"), node_id)
#         if _is_pseudo_node(node_id, label):
#             continue
#         nodes[node_id] = {"id": node_id, "label": label, "attrs": attrs, "icon_hint": _guess_icon_hint(node_id, label, attrs), "zone": _infer_zone(node_id, label)}

#     edges: List[Dict[str, Any]] = []
#     for (src, dst), attrs in edge_attrs.items():
#         for node_id in [src, dst]:
#             if node_id not in nodes:
#                 label = _label(None, node_id)
#                 if not _is_pseudo_node(node_id, label):
#                     nodes[node_id] = {"id": node_id, "label": label, "attrs": {}, "icon_hint": _guess_icon_hint(node_id, label, {}), "zone": _infer_zone(node_id, label)}
#                     node_order.append(node_id)
#         if src in nodes and dst in nodes:
#             edges.append({"source": src, "target": dst, "raw_color": attrs.get("color") or attrs.get("fontcolor"), "label": _clean_label(attrs.get("label", ""))})

#     if not edges and len(node_order) > 1:
#         for src, dst in zip(node_order[:-1], node_order[1:]):
#             if src in nodes and dst in nodes:
#                 edges.append({"source": src, "target": dst, "raw_color": None, "label": ""})

#     for cluster in clusters:
#         for nid in cluster.get("nodes", []):
#             if nid in nodes:
#                 lname = _clean_label(cluster.get("label", "")).lower()
#                 if "prem" in lname or "source" in lname:
#                     nodes[nid]["zone"] = "onprem"
#                 elif "external" in lname or "vault" in lname or "secret" in lname:
#                     nodes[nid]["zone"] = "external"
#                 elif "build" in lname or "deploy" in lname or "ci" in lname or "cd" in lname:
#                     nodes[nid]["zone"] = "cicd"
#                 elif "observ" in lname or "logging" in lname or "monitor" in lname:
#                     nodes[nid]["zone"] = "observability"
#                 elif "gcp" in lname or "cloud" in lname or "target" in lname:
#                     nodes[nid]["zone"] = "gcp"

#     _assign_flow_colors(edges, nodes)
#     legend = _legend_from_edges(edges)
#     ordered_nodes = [nodes[nid] for nid in node_order if nid in nodes]
#     return {"nodes": ordered_nodes, "node_map": nodes, "node_order": [n["id"] for n in ordered_nodes], "edges": edges, "legend": legend}


# def _infer_zone(node_id: str, label: str) -> str:
#     text = f"{node_id} {label}".lower()
#     if any(k in text for k in ["teradata", "on-prem", "on premise", "mft", "source db", "source system", "interconnect", "vpn"]):
#         return "onprem"
#     if any(k in text for k in ["vault", "hashicorp", "external", "secret manager"]):
#         return "external"
#     if any(k in text for k in ["terraform", "cloud build", "build", "iac", "ci/cd", "cicd"]):
#         return "cicd"
#     if any(k in text for k in ["logging", "monitoring", "observability"]):
#         return "observability"
#     return "gcp"


# def _assign_flow_colors(edges: List[Dict[str, Any]], nodes: Dict[str, Dict[str, Any]]):
#     used_by_label: Dict[str, str] = {}
#     palette_idx = 0
#     for edge in edges:
#         raw_color = _safe_color(edge.get("raw_color"), "") if edge.get("raw_color") else ""
#         flow_label = edge.get("label") or _edge_default_label(edge, nodes)
#         edge["display_label"] = flow_label
#         if raw_color:
#             color = raw_color
#         elif flow_label in used_by_label:
#             color = used_by_label[flow_label]
#         else:
#             color = FLOW_PALETTE[palette_idx % len(FLOW_PALETTE)]
#             used_by_label[flow_label] = color
#             palette_idx += 1
#         edge["color"] = color


# def _edge_default_label(edge: Dict[str, Any], nodes: Dict[str, Dict[str, Any]]) -> str:
#     src_label = _clean_label(nodes.get(edge["source"], {}).get("label", edge["source"])).replace("\n", " ")[:24]
#     dst_label = _clean_label(nodes.get(edge["target"], {}).get("label", edge["target"])).replace("\n", " ")[:24]
#     return f"{src_label} → {dst_label}"


# def _legend_from_edges(edges: List[Dict[str, Any]]) -> List[Tuple[str, str]]:
#     legend: List[Tuple[str, str]] = []
#     seen = set()
#     for edge in edges:
#         label = edge.get("display_label") or edge.get("label")
#         color = edge.get("color") or DEFAULT_EDGE_COLOR
#         if label and (color, label) not in seen:
#             seen.add((color, label))
#             legend.append((color, label))
#     return legend

# # -----------------------------------------------------------------------------
# # Swimlane layout
# # -----------------------------------------------------------------------------
# def _layout_model(model: Dict[str, Any]) -> Dict[str, Any]:
#     nodes = model["nodes"]
#     edges = model["edges"]
#     order = _topological_order(model.get("node_order", [n["id"] for n in nodes]), edges)
#     by_id = {n["id"]: n for n in nodes}

#     zone_order = _zone_order_for_nodes(nodes)
#     lane_boxes = _compute_lane_boxes(zone_order)
#     zone_nodes: Dict[str, List[str]] = defaultdict(list)
#     for nid in order:
#         if nid in by_id:
#             zone_nodes[by_id[nid].get("zone", "gcp")].append(nid)

#     placements: Dict[str, Dict[str, float]] = {}
#     for zone in zone_order:
#         ids = zone_nodes.get(zone, [])
#         if not ids:
#             continue
#         lane = lane_boxes[zone]
#         inner_left = lane["left"] + LANE_PAD
#         inner_top = lane["top"] + LANE_TITLE_H + LANE_PAD
#         inner_w = max(0.5, lane["width"] - 2 * LANE_PAD)
#         inner_h = max(0.5, lane["height"] - LANE_TITLE_H - 2 * LANE_PAD)

#         if zone in {"onprem", "external", "cicd", "observability"}:
#             cols = 1 if len(ids) <= 3 else 2
#         else:
#             cols = min(4, max(2, math.ceil(math.sqrt(len(ids) * 1.4))))
#         rows = math.ceil(len(ids) / cols)
#         node_w = min(NODE_W, (inner_w - (cols - 1) * H_GAP) / cols)
#         node_h = min(NODE_H, (inner_h - (rows - 1) * V_GAP) / rows)
#         node_w = max(1.15, node_w)
#         node_h = max(0.56, node_h)
#         total_w = cols * node_w + (cols - 1) * H_GAP
#         total_h = rows * node_h + (rows - 1) * V_GAP
#         start_x = inner_left + max(0.0, (inner_w - total_w) / 2)
#         start_y = inner_top + max(0.0, (inner_h - total_h) / 2)
#         for idx, nid in enumerate(ids):
#             row = idx // cols
#             col = idx % cols
#             placements[nid] = {"left": start_x + col * (node_w + H_GAP), "top": start_y + row * (node_h + V_GAP), "width": node_w, "height": node_h}
#     return {"placements": placements, "lane_boxes": list(lane_boxes.values())}


# def _zone_order_for_nodes(nodes: List[Dict[str, Any]]) -> List[str]:
#     present = {n.get("zone", "gcp") for n in nodes}
#     order = []
#     for zone in ["onprem", "gcp", "external", "cicd", "observability"]:
#         if zone in present:
#             order.append(zone)
#     return order or ["gcp"]


# def _compute_lane_boxes(zone_order: List[str]) -> Dict[str, Dict[str, Any]]:
#     boxes: Dict[str, Dict[str, Any]] = {}
#     has_onprem = "onprem" in zone_order
#     has_gcp = "gcp" in zone_order
#     right_zones = [z for z in zone_order if z not in {"onprem", "gcp"}]
#     top = CANVAS_TOP_IN
#     height = CANVAS_HEIGHT_IN
#     if has_onprem and has_gcp:
#         onprem_w = 2.45
#         right_w = 2.05 if right_zones else 0
#         gcp_w = CANVAS_WIDTH_IN - onprem_w - right_w - (0.18 if right_w else 0)
#         boxes["onprem"] = _lane_box("onprem", "On-Premise", CANVAS_LEFT_IN, top, onprem_w, height)
#         boxes["gcp"] = _lane_box("gcp", "Google Cloud Platform", CANVAS_LEFT_IN + onprem_w + 0.18, top, gcp_w, height)
#         if right_zones:
#             rz_left = CANVAS_LEFT_IN + onprem_w + 0.18 + gcp_w + 0.18
#             rz_h = height / len(right_zones)
#             for idx, zone in enumerate(right_zones):
#                 boxes[zone] = _lane_box(zone, _zone_label(zone), rz_left, top + idx * rz_h, right_w, rz_h - 0.08)
#     else:
#         n = max(1, len(zone_order))
#         lane_w = (CANVAS_WIDTH_IN - (n - 1) * 0.18) / n
#         for idx, zone in enumerate(zone_order):
#             boxes[zone] = _lane_box(zone, _zone_label(zone), CANVAS_LEFT_IN + idx * (lane_w + 0.18), top, lane_w, height)
#     return boxes


# def _lane_box(zone: str, label: str, left: float, top: float, width: float, height: float) -> Dict[str, Any]:
#     style = {
#         "onprem": ("Source / On-Premise", "EEF4FF", "8AB4F8"),
#         "gcp": ("Google Cloud Platform", "F3F8FF", "A8C7FA"),
#         "external": ("External Secret Management", EXTERNAL_FILL, EXTERNAL_BORDER),
#         "cicd": ("CI/CD & Deployment", CICD_FILL, CICD_BORDER),
#         "observability": ("Observability", OBS_FILL, OBS_BORDER),
#     }.get(zone, (label, LANE_FILL, LANE_BORDER))
#     return {"id": zone, "label": style[0], "left": left, "top": top, "width": width, "height": height, "fill": style[1], "border": style[2]}


# def _zone_label(zone: str) -> str:
#     return {"onprem": "Source / On-Premise", "gcp": "Google Cloud Platform", "external": "External Secret Management", "cicd": "CI/CD & Deployment", "observability": "Observability"}.get(zone, zone.title())


# def _topological_order(node_order: List[str], edges: List[Dict[str, Any]]) -> List[str]:
#     indeg = {nid: 0 for nid in node_order}
#     outgoing: Dict[str, List[str]] = defaultdict(list)
#     for e in edges:
#         s, d = e["source"], e["target"]
#         if s in indeg and d in indeg:
#             outgoing[s].append(d)
#             indeg[d] += 1
#     q = deque([nid for nid in node_order if indeg.get(nid, 0) == 0])
#     result = []
#     while q:
#         cur = q.popleft()
#         result.append(cur)
#         for nxt in outgoing.get(cur, []):
#             indeg[nxt] -= 1
#             if indeg[nxt] == 0:
#                 q.append(nxt)
#     for nid in node_order:
#         if nid not in result:
#             result.append(nid)
#     return result

# # -----------------------------------------------------------------------------
# # Drawing helpers
# # -----------------------------------------------------------------------------
# def _set_text_frame(shape: BaseShape, text: str, font_size_pt: float = 10, bold: bool = False, color_hex: str = DEFAULT_NODE_FONT, align=PP_ALIGN.LEFT):
#     tf = shape.text_frame
#     tf.clear()
#     tf.word_wrap = True
#     tf.margin_left = Inches(0.03)
#     tf.margin_right = Inches(0.03)
#     tf.margin_top = Inches(0.01)
#     tf.margin_bottom = Inches(0.01)
#     tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
#     p = tf.paragraphs[0]
#     p.alignment = align
#     run = p.add_run()
#     run.text = _safe_title(text)
#     run.font.name = "Calibri"
#     run.font.size = Pt(font_size_pt)
#     run.font.bold = bold
#     run.font.color.rgb = _rgb(color_hex, DEFAULT_NODE_FONT)


# def _font_for_label(label: str, width: float, height: float) -> float:
#     clean = _safe_title(label)
#     line_count = max(1, clean.count("\n") + 1)
#     length = len(clean.replace("\n", " "))
#     if width < 1.15 or height < 0.58 or length > 58 or line_count >= 3:
#         return 6.8
#     if width < 1.5 or length > 42 or line_count == 2:
#         return 7.8
#     return 8.8


# def _add_title(slide, title_text: str):
#     shape = slide.shapes.add_textbox(Inches(TITLE_LEFT_IN), Inches(TITLE_TOP_IN), Inches(TITLE_WIDTH_IN), Inches(TITLE_HEIGHT_IN))
#     _set_text_frame(shape, title_text, font_size_pt=18, bold=True, color_hex="202124")


# def _add_lane_box(slide, box: Dict[str, Any]):
#     shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
#     shape.fill.solid()
#     shape.fill.fore_color.rgb = _rgb(box.get("fill", LANE_FILL))
#     shape.line.color.rgb = _rgb(box.get("border", LANE_BORDER))
#     shape.line.width = Pt(1.2)
#     try:
#         shape.adjustments[0] = 0.04
#     except Exception:
#         pass
#     label_box = slide.shapes.add_textbox(Inches(box["left"] + 0.08), Inches(box["top"] + 0.03), Inches(max(0.5, box["width"] - 0.16)), Inches(0.20))
#     _set_text_frame(label_box, box.get("label", ""), font_size_pt=8.7, bold=True, color_hex=LANE_TITLE_COLOR)


# def _add_node_box(slide, box: Dict[str, float], label: str, icon_path: Optional[Path] = None):
#     shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
#     shape.fill.solid()
#     shape.fill.fore_color.rgb = _rgb(DEFAULT_NODE_FILL)
#     shape.line.color.rgb = _rgb(DEFAULT_NODE_BORDER)
#     shape.line.width = Pt(1.0)
#     try:
#         shape.adjustments[0] = 0.08
#     except Exception:
#         pass
#     icon_size = min(NODE_ICON_SIZE_IN, max(0.15, box["height"] * 0.36))
#     text_left = box["left"] + NODE_PADDING_IN
#     if icon_path and icon_path.exists() and box["width"] > 0.75:
#         try:
#             slide.shapes.add_picture(str(icon_path), Inches(box["left"] + NODE_PADDING_IN), Inches(box["top"] + (box["height"] - icon_size) / 2.0), width=Inches(icon_size), height=Inches(icon_size))
#             text_left = box["left"] + NODE_TEXT_LEFT_WITH_ICON_IN
#         except Exception:
#             logger.debug("Could not add icon to PPTX node: %s", icon_path, exc_info=True)
#     text_width = max(0.35, box["width"] - (text_left - box["left"]) - 0.05)
#     text_box = slide.shapes.add_textbox(Inches(text_left), Inches(box["top"] + 0.03), Inches(text_width), Inches(max(0.22, box["height"] - 0.06)))
#     _set_text_frame(text_box, label, font_size_pt=_font_for_label(label, box["width"], box["height"]), bold=True, color_hex=DEFAULT_NODE_FONT)


# def _edge_points_between_boxes(src: Dict[str, float], dst: Dict[str, float]) -> List[Tuple[float, float]]:
#     sr, sl = src["left"] + src["width"], src["left"]
#     sy = src["top"] + src["height"] / 2
#     dl, dr = dst["left"], dst["left"] + dst["width"]
#     dy = dst["top"] + dst["height"] / 2
#     if dl >= sr:
#         mx = (sr + dl) / 2
#         return [(sr, sy), (mx, sy), (mx, dy), (dl, dy)]
#     if sl >= dr:
#         mx = (dr + sl) / 2
#         return [(sl, sy), (mx, sy), (mx, dy), (dr, dy)]
#     if dy >= sy:
#         return [(src["left"] + src["width"] / 2, src["top"] + src["height"]), (dst["left"] + dst["width"] / 2, dst["top"])]
#     return [(src["left"] + src["width"] / 2, src["top"]), (dst["left"] + dst["width"] / 2, dst["top"] + dst["height"])]


# def _add_line_segment(slide, p1: Tuple[float, float], p2: Tuple[float, float], color_hex: str):
#     connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p1[0]), Inches(p1[1]), Inches(p2[0]), Inches(p2[1]))
#     connector.line.color.rgb = _rgb(color_hex, DEFAULT_EDGE_COLOR)
#     connector.line.width = Pt(DEFAULT_EDGE_WIDTH_PT)
#     return connector


# def _add_polyline_arrow(slide, points: List[Tuple[float, float]], color_hex: str):
#     if len(points) < 2:
#         return
#     clean = [points[0]]
#     for point in points[1:]:
#         if abs(point[0] - clean[-1][0]) > 0.01 or abs(point[1] - clean[-1][1]) > 0.01:
#             clean.append(point)
#     for idx in range(len(clean) - 1):
#         seg = _add_line_segment(slide, clean[idx], clean[idx + 1], color_hex)
#         if idx == len(clean) - 2:
#             try:
#                 seg.line.end_arrowhead = True
#             except Exception:
#                 pass


# def _add_legend(slide, legend_items: List[Tuple[str, str]]):
#     if not legend_items:
#         return
#     title_box = slide.shapes.add_textbox(Inches(LEGEND_LEFT_IN), Inches(LEGEND_TOP_IN), Inches(1.15), Inches(0.18))
#     _set_text_frame(title_box, "Flow Legend", font_size_pt=9.2, bold=True, color_hex="202124")
#     x, y, max_x = LEGEND_LEFT_IN, LEGEND_TOP_IN + 0.24, LEGEND_LEFT_IN + LEGEND_WIDTH_IN
#     for color, label in legend_items[:18]:
#         label = _safe_title(label)
#         item_w = max(1.25, min(3.10, 0.42 + len(label) * 0.050))
#         if x + item_w > max_x:
#             x, y = LEGEND_LEFT_IN, y + 0.22
#         if y > LEGEND_TOP_IN + LEGEND_MAX_HEIGHT_IN:
#             break
#         chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(item_w), Inches(0.18))
#         chip.fill.solid(); chip.fill.fore_color.rgb = _rgb(LEGEND_FILL); chip.line.color.rgb = _rgb(LEGEND_BORDER); chip.line.width = Pt(0.6)
#         dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.04), Inches(y + 0.045), Inches(0.09), Inches(0.09))
#         dot.fill.solid(); dot.fill.fore_color.rgb = _rgb(color, DEFAULT_EDGE_COLOR); dot.line.color.rgb = _rgb(color, DEFAULT_EDGE_COLOR)
#         label_box = slide.shapes.add_textbox(Inches(x + 0.16), Inches(y + 0.01), Inches(max(0.3, item_w - 0.18)), Inches(0.15))
#         _set_text_frame(label_box, label, font_size_pt=7.2, bold=False, color_hex="202124")
#         x += item_w + 0.08


# def _add_fallback_preview_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str):
#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, f"{slide_title} (Preview)")
#     try:
#         assets = render_graphviz_to_assets(diagram_text, output_dir=output_dir, base_name=base_name)
#     except Exception:
#         logger.exception("Fallback preview asset generation failed.")
#         assets = None
#     if assets and assets.get("png_path") and Path(assets["png_path"]).exists():
#         slide.shapes.add_picture(str(assets["png_path"]), Inches(CANVAS_LEFT_IN), Inches(CANVAS_TOP_IN), width=Inches(PREVIEW_MAX_WIDTH_IN))
#         _add_legend(slide, assets.get("legend", []))

# # -----------------------------------------------------------------------------
# # Diagram extraction
# # -----------------------------------------------------------------------------
# def _is_diagram_text(value: Any) -> bool:
#     return isinstance(value, str) and normalize_to_graphviz_dot(_clean_internal_tokens(value)) is not None


# def _walk_for_diagrams(value: Any, path_parts: List[str], results: List[Dict[str, Any]]):
#     if value is None:
#         return
#     if isinstance(value, str):
#         if _is_diagram_text(value):
#             results.append({"title": " / ".join(path_parts) if path_parts else "Diagram", "diagram_text": _clean_internal_tokens(value), "section_key": _safe_name(path_parts[-1] if path_parts else "diagram")})
#         return
#     if isinstance(value, list):
#         for idx, item in enumerate(value, start=1):
#             if _is_diagram_text(item):
#                 results.append({"title": " / ".join(path_parts) if path_parts else f"Diagram {idx}", "diagram_text": _clean_internal_tokens(item), "section_key": _safe_name(path_parts[-1] if path_parts else f"diagram_{idx}")})
#             elif isinstance(item, (dict, list)):
#                 _walk_for_diagrams(item, path_parts, results)
#         return
#     if isinstance(value, dict):
#         for key, child in value.items():
#             key_text = str(key).replace("_", " ").title()
#             next_path = path_parts if str(key).lower() == "diagrams" else path_parts + [key_text]
#             _walk_for_diagrams(child, next_path, results)


# def extract_diagram_specs(data: Dict[str, Any], hld_model: Any, selected_sections_set: Optional[Set[str]] = None) -> List[Dict[str, Any]]:
#     results: List[Dict[str, Any]] = []
#     for field_name, field_info in hld_model.model_fields.items():
#         if not should_include_section(field_name, field_info, selected_sections_set or set()):
#             continue
#         section_title = field_info.title or field_name.replace("_", " ").title()
#         _walk_for_diagrams(data.get(field_name), [section_title], results)
#     counts: Dict[str, int] = defaultdict(int)
#     for item in results:
#         counts[item["title"]] += 1
#     seen: Dict[str, int] = defaultdict(int)
#     for item in results:
#         if counts[item["title"]] > 1:
#             seen[item["title"]] += 1
#             item["title"] = f'{item["title"]} ({seen[item["title"]]})'
#     return results

# # -----------------------------------------------------------------------------
# # Slide rendering
# # -----------------------------------------------------------------------------
# def add_diagram_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str):
#     try:
#         model = _build_editable_model(diagram_text)
#         if not model["nodes"]:
#             raise ValueError("No editable nodes parsed")
#         layout = _layout_model(model)
#         placements = layout["placements"]
#         lane_boxes = layout["lane_boxes"]
#     except Exception:
#         logger.exception("Editable PPTX rendering failed for slide '%s'. Falling back to preview image.", slide_title)
#         _add_fallback_preview_slide(prs, slide_title, diagram_text, output_dir, base_name)
#         return

#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, slide_title)
#     for box in lane_boxes:
#         _add_lane_box(slide, box)
#     for edge in model["edges"]:
#         src_box = placements.get(edge["source"])
#         dst_box = placements.get(edge["target"])
#         if src_box and dst_box:
#             _add_polyline_arrow(slide, _edge_points_between_boxes(src_box, dst_box), edge.get("color") or DEFAULT_EDGE_COLOR)
#     for node in model["nodes"]:
#         box = placements.get(node["id"])
#         if not box:
#             continue
#         icon_path = _preferred_picture_icon_path(node.get("icon_hint") or node.get("label") or node["id"])
#         _add_node_box(slide, box, node.get("label") or node["id"], icon_path)
#     _add_legend(slide, model.get("legend", []))

# # -----------------------------------------------------------------------------
# # Public builder
# # -----------------------------------------------------------------------------
# def build_editable_pptx(data: Dict[str, Any], pptx_path: str, hld_model: Any, selected_sections_set: Optional[Set[str]] = None):
#     output_path = Path(pptx_path)
#     output_path.parent.mkdir(parents=True, exist_ok=True)
#     diagram_assets_dir = output_path.parent / f"{output_path.stem}_diagram_assets"
#     diagram_assets_dir.mkdir(parents=True, exist_ok=True)
#     diagram_specs = extract_diagram_specs(data=data, hld_model=hld_model, selected_sections_set=selected_sections_set)
#     prs = Presentation()
#     prs.slide_width = Inches(SLIDE_WIDTH_IN)
#     prs.slide_height = Inches(SLIDE_HEIGHT_IN)
#     if not diagram_specs:
#         slide = prs.slides.add_slide(prs.slide_layouts[6])
#         _add_title(slide, "Editable Diagrams")
#         msg = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11.5), Inches(1.0))
#         _set_text_frame(msg, "No diagrams were detected in the selected sections.", font_size_pt=16, bold=True, color_hex="5F6368", align=PP_ALIGN.CENTER)
#         prs.save(str(output_path))
#         return
#     for counter, spec in enumerate(diagram_specs, start=1):
#         add_diagram_slide(prs, spec["title"], spec["diagram_text"], diagram_assets_dir, f"{_safe_name(spec['section_key'])}_{counter:02d}")
#     prs.save(str(output_path))
#     logger.info("Editable PPTX generated: %s", output_path)

#$4th Import
# from __future__ import annotations

# import html
# import logging
# import math
# import re
# from collections import defaultdict, deque
# from pathlib import Path
# from typing import Any, Dict, List, Optional, Set, Tuple

# from pptx import Presentation
# from pptx.dml.color import RGBColor
# from pptx.enum.dml import MSO_LINE_DASH_STYLE
# from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
# from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
# from pptx.oxml.xmlchemy import OxmlElement
# from pptx.shapes.base import BaseShape
# from pptx.util import Inches, Pt

# from .config import ICON_DIR
# from .graphviz_renderer import normalize_to_graphviz_dot, render_graphviz_to_assets
# from .schema_utils import should_include_section
# from .text_sanitizer import renderer_safe_plain_text

# logger = logging.getLogger(__name__)

# # =============================================================================
# # Editable PPTX Diagram Renderer - swimlane version
# # =============================================================================
# # Renders editable PowerPoint diagrams with grouped zone boxes similar to PDF:
# #   - On-Premise box
# #   - GCP box
# #   - External / CI-CD / Observability boxes when applicable
# #   - editable node cards with icons
# #   - colored routed arrows with arrowheads
# #   - matching Flow Legend
# # =============================================================================

# SLIDE_WIDTH_IN = 13.333
# SLIDE_HEIGHT_IN = 7.5
# TITLE_LEFT_IN = 0.35
# TITLE_TOP_IN = 0.12
# TITLE_WIDTH_IN = 12.6
# TITLE_HEIGHT_IN = 0.42
# CANVAS_LEFT_IN = 0.35
# CANVAS_TOP_IN = 0.68
# CANVAS_WIDTH_IN = 12.3
# CANVAS_HEIGHT_IN = 5.25
# LEGEND_LEFT_IN = 0.35
# LEGEND_TOP_IN = 6.12
# LEGEND_WIDTH_IN = 12.3
# LEGEND_MAX_HEIGHT_IN = 1.22

# NODE_W = 1.78
# NODE_H = 0.78
# NODE_ICON_SIZE_IN = 0.22
# NODE_PADDING_IN = 0.07
# NODE_TEXT_LEFT_WITH_ICON_IN = 0.33

# LANE_PAD = 0.18
# LANE_TITLE_H = 0.26
# H_GAP = 0.28
# V_GAP = 0.28

# DEFAULT_NODE_FILL = "F8F9FA"
# DEFAULT_NODE_BORDER = "DADCE0"
# DEFAULT_NODE_FONT = "202124"

# LANE_FILL = "F3F8FF"
# LANE_BORDER = "A8C7FA"
# LANE_TITLE_COLOR = "174EA6"

# EXTERNAL_FILL = "FFF8E1"
# EXTERNAL_BORDER = "F9AB00"
# CICD_FILL = "FCE8E6"
# CICD_BORDER = "F28B82"
# OBS_FILL = "F3E8FD"
# OBS_BORDER = "D7AEFB"

# DEFAULT_EDGE_COLOR = "5F6368"
# DEFAULT_EDGE_WIDTH_PT = 1.6
# LEGEND_FILL = "FFFFFF"
# LEGEND_BORDER = "DADCE0"

# PREVIEW_MAX_WIDTH_IN = 12.0
# SUPPORTED_PICTURE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif"}
# PREFERRED_ICON_EXTS = [".png", ".jpg", ".jpeg", ".bmp", ".gif"]

# COMMENT_LINE_RE = re.compile(r"//.*?$", re.MULTILINE)
# BLOCK_COMMENT_RE = re.compile(r"/\*.*?\*/", re.DOTALL)
# EDGE_RE = re.compile(r'("[^"]+"|[A-Za-z0-9_.:-]+)\s*(->|--)\s*("[^"]+"|[A-Za-z0-9_.:-]+)')

# FLOW_PALETTE = [
#     "E60000", "1D70B8", "28A197", "F47738", "6F35A5",
#     "007C89", "D53880", "85994B", "B58840", "5F6368"
# ]

# COLOR_NAME_MAP = {
#     "red": "E60000",
#     "blue": "1D70B8",
#     "green": "28A197",
#     "orange": "F47738",
#     "purple": "6F35A5",
#     "teal": "007C89",
#     "gray": "5F6368",
#     "grey": "5F6368",
#     "black": "202124",
#     "lightgrey": "DADCE0",
#     "lightgray": "DADCE0",
# }

# ICON_KEYWORD_MAP = [
#     ("cloud storage", "cloud_storage"),
#     ("gcs", "cloud_storage"),
#     ("storage", "cloud_storage"),
#     ("cloud function", "cloud_functions"),
#     ("cloud functions", "cloud_functions"),
#     ("function", "cloud_functions"),
#     ("pub/sub", "pubsub"),
#     ("pubsub", "pubsub"),
#     ("notification", "pubsub"),
#     ("eventarc", "eventarc"),
#     ("vault", "secret_manager"),
#     ("hashicorp", "secret_manager"),
#     ("secret", "secret_manager"),
#     ("kms", "key_management_service"),
#     ("key management", "key_management_service"),
#     ("logging", "cloud_logging"),
#     ("monitoring", "cloud_monitoring"),
#     ("cloud build", "cloud_build"),
#     ("build", "cloud_build"),
#     ("terraform", "terraform"),
#     ("teradata", "database_migration_service"),
#     ("database", "database_migration_service"),
#     ("mft", "transfer"),
#     ("managed file transfer", "transfer"),
#     ("transfer", "transfer"),
#     ("interconnect", "cloud_interconnect"),
#     ("vpn", "cloud_vpn"),
#     ("collibra", "catalog"),
#     ("catalog", "catalog"),
# ]

# ARROW_HEAD_TYPE = "triangle"
# ARROW_HEAD_WIDTH = "med"
# ARROW_HEAD_LENGTH = "med"
# ARROW_TIP_GAP_IN = 0.07


# def _clean_internal_tokens(text: Any) -> str:
#     value = str(text or "")
#     for token in [
#         "AIASECTIONBLOCKSTARTTOKEN",
#         "AIASECTIONBLOCKENDTOKEN",
#         "[[AIASECTIONBLOCKSTARTTOKEN]]",
#         "[[AIASECTIONBLOCKENDTOKEN]]",
#         "SECTION_BLOCK_START",
#         "SECTION_BLOCK_END",
#         "[[SECTION_BLOCK_START]]",
#         "[[SECTION_BLOCK_END]]",
#         "[[SECTIONBLOCKSTART]]",
#         "[[SECTIONBLOCKEND]]",
#     ]:
#         value = value.replace(token, "")
#     return html.unescape(value.strip())


# def _safe_title(text: Any) -> str:
#     return renderer_safe_plain_text(_clean_internal_tokens(text)).strip() or "Diagram"


# def _safe_name(text: str) -> str:
#     return re.sub(r"[^a-zA-Z0-9_-]+", "_", str(text or "")).strip("_").lower() or "diagram"


# def _clean_label(value: Any) -> str:
#     text = html.unescape(str(value or "").strip()).strip('"')
#     text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
#     text = re.sub(r"<[^>]+>", " ", text)
#     text = text.replace("\\\\N", "\n").replace("\\\\n", "\n")
#     text = text.replace("\\N", "\n").replace("\\n", "\n").replace("&nbsp;", " ")
#     text = re.sub(r"\s+\n", "\n", text)
#     text = re.sub(r"\n\s+", "\n", text)
#     text = re.sub(r"[ \t]+", " ", text)
#     return text.strip()


# def _label(raw: Optional[str], fallback: str) -> str:
#     label = _clean_label(raw or "")
#     if not label:
#         label = str(fallback or "").replace("_", " ").replace("-", " ").title()
#     return label[:180]


# def _safe_color(value: Optional[str], default: str = DEFAULT_EDGE_COLOR) -> str:
#     raw = str(value or "").strip().strip('"').lstrip("#")
#     if not raw:
#         return default
#     if raw.lower() in COLOR_NAME_MAP:
#         return COLOR_NAME_MAP[raw.lower()]
#     raw = raw.upper()
#     return raw if len(raw) == 6 and all(ch in "0123456789ABCDEF" for ch in raw) else default


# def _rgb(hex_color: str, default: str = "000000") -> RGBColor:
#     return RGBColor.from_string(_safe_color(hex_color, default))


# def _is_pseudo_node(node_id: str, label: Optional[str] = None) -> bool:
#     text = _clean_label(label if label is not None else node_id).lower().strip()
#     raw = str(node_id or "").lower().strip()
#     if text in {"", "n/a", "na", "none", "null"}:
#         return True
#     if raw in {
#         "node", "edge", "graph", "digraph", "subgraph", "rank",
#         "label", "style", "color", "fillcolor", "fontcolor", "rankdir"
#     }:
#         return True
#     if raw.startswith((
#         "digraph", "graph", "subgraph", "cluster_", "cluster ",
#         "label=", "style=", "color=", "fillcolor=", "fontcolor=", "rankdir="
#     )):
#         return True
#     if text.startswith((
#         "label=", "style=", "color=", "fillcolor=", "fontcolor=",
#         "rankdir=", "fontsize=", "fontname=", "margin=", "pad="
#     )):
#         return True
#     if text in {"filled", "lightgrey", "lightgray", "blue", "green", "purple", "orange", "red", "gray", "grey"}:
#         return True
#     return False


# def _is_graphviz_control_statement(statement: str) -> bool:
#     s = html.unescape(str(statement or "")).strip().lower()
#     if not s or s in {"{", "}"}:
#         return True
#     return s.startswith((
#         "digraph ", "graph ", "subgraph ", "node ", "edge ", "rank ",
#         "rankdir", "label", "style", "color", "fillcolor", "fontcolor",
#         "fontsize", "fontname", "margin", "pad", "splines", "bgcolor"
#     ))


# def _try_existing_icon_resolver(icon_hint: str) -> Optional[Path]:
#     try:
#         from . import icon_resolver as ir
#     except Exception:
#         return None

#     for fn_name in ["resolve_icon_path", "resolve_icon", "get_icon_path", "find_icon_path"]:
#         fn = getattr(ir, fn_name, None)
#         if not callable(fn):
#             continue
#         for args in [(icon_hint,), (icon_hint, True), (icon_hint, None)]:
#             try:
#                 result = fn(*args)
#                 if result:
#                     path = Path(str(result))
#                     if path.exists():
#                         return path
#             except TypeError:
#                 continue
#             except Exception:
#                 logger.debug("Icon resolver failed: %s", fn_name, exc_info=True)
#     return None


# def _search_icon_dir(icon_hint: str) -> Optional[Path]:
#     if not icon_hint:
#         return None

#     hint = _safe_name(icon_hint)
#     bases = {
#         hint,
#         icon_hint,
#         icon_hint.lower(),
#         icon_hint.replace(" ", "_").lower(),
#         icon_hint.replace("-", "_").lower(),
#         icon_hint.replace("/", "_").lower(),
#     }

#     for base in bases:
#         if not base:
#             continue
#         for ext in [".png", ".jpg", ".jpeg", ".bmp", ".gif", ".svg"]:
#             candidate = ICON_DIR / f"{base}{ext}"
#             if candidate.exists():
#                 return candidate

#     try:
#         for path in ICON_DIR.rglob("*"):
#             if path.is_file() and hint in path.stem.lower():
#                 return path
#     except Exception:
#         logger.debug("Icon directory search failed", exc_info=True)

#     return None


# def _keyword_icon_hint(label_or_id: str) -> Optional[str]:
#     text = _clean_label(label_or_id).lower()
#     for keyword, icon_hint in ICON_KEYWORD_MAP:
#         if keyword in text:
#             return icon_hint
#     return None


# def _preferred_picture_icon_path(icon_hint: str) -> Optional[Path]:
#     candidates: List[str] = []
#     keyword = _keyword_icon_hint(icon_hint)
#     if keyword:
#         candidates.append(keyword)
#     if icon_hint:
#         candidates.append(icon_hint)

#     for candidate in candidates:
#         path = _try_existing_icon_resolver(candidate) or _search_icon_dir(candidate)
#         if not path:
#             continue
#         if path.suffix.lower() in SUPPORTED_PICTURE_EXTS:
#             return path
#         if path.suffix.lower() == ".svg":
#             for ext in PREFERRED_ICON_EXTS:
#                 sibling = path.with_suffix(ext)
#                 if sibling.exists():
#                     return sibling
#     return None


# def _guess_icon_hint(node_id: str, node_label: str, attrs: Dict[str, str]) -> str:
#     image_attr = attrs.get("image") or attrs.get("icon") or attrs.get("imagepath")
#     if image_attr:
#         stem = Path(str(image_attr)).stem
#         if stem:
#             return stem

#     keyword = _keyword_icon_hint(node_label) or _keyword_icon_hint(node_id)
#     if keyword:
#         return keyword

#     for candidate in [node_id, node_label.split("\n", 1)[0], node_label.replace("\n", " ")]:
#         cleaned = _safe_name(candidate)
#         if cleaned:
#             return cleaned

#     return _safe_name(node_id or node_label or "node")


# def _strip_dot_comments(dot_text: str) -> str:
#     text = BLOCK_COMMENT_RE.sub("", dot_text or "")
#     return COMMENT_LINE_RE.sub("", text)


# def _outer_graph_body(dot: str) -> str:
#     text = dot or ""
#     start = text.find("{")
#     end = text.rfind("}")
#     if start != -1 and end != -1 and end > start:
#         return text[start + 1:end]
#     return text


# def _split_statements(dot_text: str) -> List[str]:
#     statements: List[str] = []
#     buf: List[str] = []
#     bracket_depth = 0
#     brace_depth = 0
#     angle_depth = 0

#     for ch in dot_text or "":
#         if ch == "[":
#             bracket_depth += 1
#         elif ch == "]":
#             bracket_depth = max(0, bracket_depth - 1)
#         elif ch == "{":
#             brace_depth += 1
#         elif ch == "}":
#             brace_depth = max(0, brace_depth - 1)
#         elif ch == "<":
#             angle_depth += 1
#         elif ch == ">":
#             angle_depth = max(0, angle_depth - 1)

#         if ch == ";" and bracket_depth == 0 and brace_depth == 0 and angle_depth == 0:
#             stmt = "".join(buf).strip()
#             if stmt:
#                 statements.append(stmt)
#             buf = []
#         else:
#             buf.append(ch)

#     tail = "".join(buf).strip()
#     if tail:
#         statements.append(tail)

#     return statements


# def _extract_attr_pairs(attr_text: str) -> Dict[str, str]:
#     attrs: Dict[str, str] = {}
#     if not attr_text:
#         return attrs

#     text = html.unescape(str(attr_text))
#     for match in re.finditer(r'([A-Za-z0-9_:-]+)\s*=\s*("[^"]*"|<[^>]*>|[^,\n\r;]+)', text):
#         key = match.group(1).lower().strip()
#         value = match.group(2).strip().strip('"')
#         if value.startswith("<") and value.endswith(">"):
#             value = value[1:-1]
#         attrs[key] = value.strip()
#     return attrs


# def _extract_bracket_attr(statement: str) -> Tuple[str, Dict[str, str]]:
#     statement = html.unescape((statement or "").strip().rstrip(";").strip())
#     if "[" not in statement or "]" not in statement:
#         return statement, {}

#     start = statement.find("[")
#     end = statement.rfind("]")
#     if end <= start:
#         return statement, {}

#     return statement[:start].strip(), _extract_attr_pairs(statement[start + 1:end])


# def _unquote_id(value: str) -> str:
#     return str(value or "").strip().strip('"').strip()


# def _find_subgraph_blocks(dot_text: str) -> List[Tuple[str, str]]:
#     text = dot_text or ""
#     results: List[Tuple[str, str]] = []
#     idx = 0

#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             break

#         start = idx + match.start()
#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break

#         header = text[start:brace_start].strip()
#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1

#         if end >= len(text):
#             break

#         results.append((header, text[brace_start + 1:end]))
#         idx = end + 1

#     return results


# def _remove_subgraph_blocks(dot_text: str) -> str:
#     text = dot_text or ""
#     pieces: List[str] = []
#     idx = 0

#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             pieces.append(text[idx:])
#             break

#         start = idx + match.start()
#         pieces.append(text[idx:start])

#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break

#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1

#         if end >= len(text):
#             break

#         idx = end + 1

#     return "".join(pieces)


# def _extract_cluster_label(body: str, fallback_name: str) -> str:
#     for stmt in _split_statements(body):
#         s = html.unescape(stmt or "").strip()
#         if s.lower().startswith("label") and "=" in s:
#             try:
#                 value = s.split("=", 1)[1].strip().strip('"').strip()
#                 if value:
#                     return _clean_label(value)
#             except Exception:
#                 pass
#     return fallback_name


# def _extract_cluster_nodes(body: str) -> Set[str]:
#     node_ids: Set[str] = set()

#     for stmt in _split_statements(body):
#         s = html.unescape(stmt or "").strip()
#         if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s):
#             continue

#         prefix, attrs = _extract_bracket_attr(s)
#         node_id = _unquote_id(prefix)
#         label = attrs.get("label", node_id)

#         if not node_id or _is_pseudo_node(node_id, label):
#             continue

#         if re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
#             node_ids.add(node_id)

#     for match in EDGE_RE.finditer(body or ""):
#         for idx in [1, 3]:
#             node_id = _unquote_id(match.group(idx))
#             if node_id and not _is_pseudo_node(node_id):
#                 node_ids.add(node_id)

#     return node_ids


# def _parse_clusters_from_dot(dot: str) -> List[Dict[str, Any]]:
#     clusters: List[Dict[str, Any]] = []
#     for header, body in _find_subgraph_blocks(dot):
#         fallback_name = _unquote_id(header.replace("subgraph", "", 1).strip()) or "Cluster"
#         cluster_label = _extract_cluster_label(body, fallback_name)
#         cluster_nodes = sorted(_extract_cluster_nodes(body))
#         if cluster_nodes and not _is_pseudo_node(fallback_name, cluster_label):
#             clusters.append({
#                 "id": _safe_name(fallback_name),
#                 "label": cluster_label,
#                 "nodes": cluster_nodes,
#             })
#     return clusters


# def _parse_node_attrs_from_dot(dot: str) -> Tuple[Dict[str, Dict[str, str]], List[str]]:
#     attrs_by_node: Dict[str, Dict[str, str]] = {}
#     node_order: List[str] = []

#     def parse_node_statements(text: str):
#         for stmt in _split_statements(text):
#             s = html.unescape(stmt or "").strip()
#             if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s):
#                 continue

#             prefix, attrs = _extract_bracket_attr(s)
#             node_id = _unquote_id(prefix)
#             if not node_id:
#                 continue

#             label = attrs.get("label", node_id)
#             if _is_pseudo_node(node_id, label):
#                 continue

#             if not re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
#                 continue

#             if node_id not in attrs_by_node:
#                 node_order.append(node_id)
#             attrs_by_node[node_id] = attrs

#     body = _outer_graph_body(dot)
#     parse_node_statements(_remove_subgraph_blocks(body))
#     for _header, subgraph_body in _find_subgraph_blocks(dot):
#         parse_node_statements(subgraph_body)

#     return attrs_by_node, node_order


# def _parse_edge_attrs_from_dot(dot: str) -> Dict[Tuple[str, str], Dict[str, str]]:
#     attrs_by_edge: Dict[Tuple[str, str], Dict[str, str]] = {}

#     def parse_edge_statements(text: str):
#         for stmt in _split_statements(text):
#             s = html.unescape(stmt or "").strip()
#             if "->" not in s and "--" not in s:
#                 continue

#             prefix, attrs = _extract_bracket_attr(s)
#             tokens = re.findall(r'"[^"]+"|[A-Za-z0-9_.:-]+', prefix)
#             if len(tokens) < 2:
#                 continue

#             for src_raw, dst_raw in zip(tokens[:-1], tokens[1:]):
#                 src = _unquote_id(src_raw)
#                 dst = _unquote_id(dst_raw)
#                 if not src or not dst or _is_pseudo_node(src) or _is_pseudo_node(dst):
#                     continue
#                 attrs_by_edge[(src, dst)] = attrs

#     body = _outer_graph_body(dot)
#     parse_edge_statements(_remove_subgraph_blocks(body))
#     for _header, subgraph_body in _find_subgraph_blocks(dot):
#         parse_edge_statements(subgraph_body)

#     for match in EDGE_RE.finditer(dot or ""):
#         src = _unquote_id(match.group(1))
#         dst = _unquote_id(match.group(3))
#         if src and dst and not _is_pseudo_node(src) and not _is_pseudo_node(dst):
#             attrs_by_edge.setdefault((src, dst), {})

#     return attrs_by_edge


# # -----------------------------------------------------------------------------
# # Model
# # -----------------------------------------------------------------------------
# def _build_editable_model(diagram_text: str) -> Dict[str, Any]:
#     normalized = normalize_to_graphviz_dot(diagram_text)
#     if not normalized:
#         raise ValueError("Diagram text could not be normalized to DOT.")

#     dot = _strip_dot_comments(html.unescape(normalized))
#     clusters = _parse_clusters_from_dot(dot)
#     node_attrs, node_order = _parse_node_attrs_from_dot(dot)
#     edge_attrs = _parse_edge_attrs_from_dot(dot)

#     nodes: Dict[str, Dict[str, Any]] = {}
#     for node_id in node_order:
#         attrs = node_attrs.get(node_id, {})
#         label = _label(attrs.get("label"), node_id)
#         if _is_pseudo_node(node_id, label):
#             continue

#         nodes[node_id] = {
#             "id": node_id,
#             "label": label,
#             "attrs": attrs,
#             "icon_hint": _guess_icon_hint(node_id, label, attrs),
#             "zone": _infer_zone(node_id, label),
#         }

#     edges: List[Dict[str, Any]] = []
#     for (src, dst), attrs in edge_attrs.items():
#         for node_id in [src, dst]:
#             if node_id not in nodes:
#                 label = _label(None, node_id)
#                 if not _is_pseudo_node(node_id, label):
#                     nodes[node_id] = {
#                         "id": node_id,
#                         "label": label,
#                         "attrs": {},
#                         "icon_hint": _guess_icon_hint(node_id, label, {}),
#                         "zone": _infer_zone(node_id, label),
#                     }
#                     node_order.append(node_id)

#         if src in nodes and dst in nodes:
#             edges.append({
#                 "source": src,
#                 "target": dst,
#                 "raw_color": attrs.get("color") or attrs.get("fontcolor"),
#                 "label": _clean_label(attrs.get("label", "")),
#             })

#     if not edges and len(node_order) > 1:
#         for src, dst in zip(node_order[:-1], node_order[1:]):
#             if src in nodes and dst in nodes:
#                 edges.append({
#                     "source": src,
#                     "target": dst,
#                     "raw_color": None,
#                     "label": "",
#                 })

#     for cluster in clusters:
#         for nid in cluster.get("nodes", []):
#             if nid in nodes:
#                 lname = _clean_label(cluster.get("label", "")).lower()
#                 if "prem" in lname or "source" in lname:
#                     nodes[nid]["zone"] = "onprem"
#                 elif "external" in lname or "vault" in lname or "secret" in lname:
#                     nodes[nid]["zone"] = "external"
#                 elif "build" in lname or "deploy" in lname or "ci" in lname or "cd" in lname:
#                     nodes[nid]["zone"] = "cicd"
#                 elif "observ" in lname or "logging" in lname or "monitor" in lname:
#                     nodes[nid]["zone"] = "observability"
#                 elif "gcp" in lname or "cloud" in lname or "target" in lname:
#                     nodes[nid]["zone"] = "gcp"

#     _assign_flow_colors(edges, nodes)
#     legend = _legend_from_edges(edges)
#     ordered_nodes = [nodes[nid] for nid in node_order if nid in nodes]

#     return {
#         "nodes": ordered_nodes,
#         "node_map": nodes,
#         "node_order": [n["id"] for n in ordered_nodes],
#         "edges": edges,
#         "legend": legend,
#     }


# def _infer_zone(node_id: str, label: str) -> str:
#     text = f"{node_id} {label}".lower()
#     if any(k in text for k in ["teradata", "on-prem", "on premise", "mft", "source db", "source system", "interconnect", "vpn"]):
#         return "onprem"
#     if any(k in text for k in ["vault", "hashicorp", "external", "secret manager"]):
#         return "external"
#     if any(k in text for k in ["terraform", "cloud build", "build", "iac", "ci/cd", "cicd"]):
#         return "cicd"
#     if any(k in text for k in ["logging", "monitoring", "observability"]):
#         return "observability"
#     return "gcp"


# def _assign_flow_colors(edges: List[Dict[str, Any]], nodes: Dict[str, Dict[str, Any]]):
#     used_by_label: Dict[str, str] = {}
#     palette_idx = 0

#     for edge in edges:
#         raw_color = _safe_color(edge.get("raw_color"), "") if edge.get("raw_color") else ""
#         flow_label = edge.get("label") or _edge_default_label(edge, nodes)
#         edge["display_label"] = flow_label

#         if raw_color:
#             color = raw_color
#         elif flow_label in used_by_label:
#             color = used_by_label[flow_label]
#         else:
#             color = FLOW_PALETTE[palette_idx % len(FLOW_PALETTE)]
#             used_by_label[flow_label] = color
#             palette_idx += 1

#         edge["color"] = color

#         dash_style, width_pt = _edge_visual_style(edge, nodes)
#         edge["dash_style"] = dash_style
#         edge["width_pt"] = width_pt


# def _edge_default_label(edge: Dict[str, Any], nodes: Dict[str, Dict[str, Any]]) -> str:
#     src_label = _clean_label(nodes.get(edge["source"], {}).get("label", edge["source"])).replace("\n", " ")[:24]
#     dst_label = _clean_label(nodes.get(edge["target"], {}).get("label", edge["target"])).replace("\n", " ")[:24]
#     return f"{src_label} → {dst_label}"


# def _edge_visual_style(edge: Dict[str, Any], nodes: Dict[str, Dict[str, Any]]) -> Tuple[Optional[MSO_LINE_DASH_STYLE], float]:
#     label = (edge.get("display_label") or edge.get("label") or "").lower()
#     src_label = _clean_label(nodes.get(edge["source"], {}).get("label", edge["source"])).lower()
#     dst_label = _clean_label(nodes.get(edge["target"], {}).get("label", edge["target"])).lower()
#     combined = f"{label} {src_label} {dst_label}"

#     if any(k in combined for k in ["vault", "kms", "key", "secret", "cmek", "decrypt key", "encryption key"]):
#         return MSO_LINE_DASH_STYLE.ROUND_DOT, 1.55

#     if any(k in combined for k in ["pub/sub", "pubsub", "notification", "notify", "event", "trigger", "message"]):
#         return MSO_LINE_DASH_STYLE.DASH, 1.55

#     if any(k in combined for k in ["monitor", "logging", "observability", "metric", "alert"]):
#         return MSO_LINE_DASH_STYLE.DASH_DOT, 1.45

#     return None, DEFAULT_EDGE_WIDTH_PT


# def _legend_from_edges(edges: List[Dict[str, Any]]) -> List[Tuple[str, str, Optional[MSO_LINE_DASH_STYLE]]]:
#     legend: List[Tuple[str, str, Optional[MSO_LINE_DASH_STYLE]]] = []
#     seen = set()
#     for edge in edges:
#         label = edge.get("display_label") or edge.get("label")
#         color = edge.get("color") or DEFAULT_EDGE_COLOR
#         dash_style = edge.get("dash_style")
#         key = (color, label, str(dash_style))
#         if label and key not in seen:
#             seen.add(key)
#             legend.append((color, label, dash_style))
#     return legend


# # -----------------------------------------------------------------------------
# # Swimlane layout
# # -----------------------------------------------------------------------------
# def _layout_model(model: Dict[str, Any]) -> Dict[str, Any]:
#     nodes = model["nodes"]
#     edges = model["edges"]
#     order = _topological_order(model.get("node_order", [n["id"] for n in nodes]), edges)
#     by_id = {n["id"]: n for n in nodes}

#     zone_order = _zone_order_for_nodes(nodes)
#     lane_boxes = _compute_lane_boxes(zone_order)
#     zone_nodes: Dict[str, List[str]] = defaultdict(list)

#     for nid in order:
#         if nid in by_id:
#             zone_nodes[by_id[nid].get("zone", "gcp")].append(nid)

#     placements: Dict[str, Dict[str, float]] = {}
#     for zone in zone_order:
#         ids = zone_nodes.get(zone, [])
#         if not ids:
#             continue

#         lane = lane_boxes[zone]
#         inner_left = lane["left"] + LANE_PAD
#         inner_top = lane["top"] + LANE_TITLE_H + LANE_PAD
#         inner_w = max(0.5, lane["width"] - 2 * LANE_PAD)
#         inner_h = max(0.5, lane["height"] - LANE_TITLE_H - 2 * LANE_PAD)

#         if zone in {"onprem", "external", "cicd", "observability"}:
#             cols = 1 if len(ids) <= 3 else 2
#         else:
#             cols = min(4, max(2, math.ceil(math.sqrt(len(ids) * 1.4))))

#         rows = math.ceil(len(ids) / cols)

#         node_w = min(NODE_W, (inner_w - (cols - 1) * H_GAP) / cols)
#         node_h = min(NODE_H, (inner_h - (rows - 1) * V_GAP) / rows)
#         node_w = max(1.15, node_w)
#         node_h = max(0.56, node_h)

#         total_w = cols * node_w + (cols - 1) * H_GAP
#         total_h = rows * node_h + (rows - 1) * V_GAP
#         start_x = inner_left + max(0.0, (inner_w - total_w) / 2)
#         start_y = inner_top + max(0.0, (inner_h - total_h) / 2)

#         for idx, nid in enumerate(ids):
#             row = idx // cols
#             col = idx % cols
#             placements[nid] = {
#                 "left": start_x + col * (node_w + H_GAP),
#                 "top": start_y + row * (node_h + V_GAP),
#                 "width": node_w,
#                 "height": node_h,
#             }

#     return {"placements": placements, "lane_boxes": list(lane_boxes.values())}


# def _zone_order_for_nodes(nodes: List[Dict[str, Any]]) -> List[str]:
#     present = {n.get("zone", "gcp") for n in nodes}
#     order = []
#     for zone in ["onprem", "gcp", "external", "cicd", "observability"]:
#         if zone in present:
#             order.append(zone)
#     return order or ["gcp"]


# def _compute_lane_boxes(zone_order: List[str]) -> Dict[str, Dict[str, Any]]:
#     boxes: Dict[str, Dict[str, Any]] = {}
#     has_onprem = "onprem" in zone_order
#     has_gcp = "gcp" in zone_order
#     right_zones = [z for z in zone_order if z not in {"onprem", "gcp"}]

#     top = CANVAS_TOP_IN
#     height = CANVAS_HEIGHT_IN

#     if has_onprem and has_gcp:
#         onprem_w = 2.45
#         right_w = 2.05 if right_zones else 0
#         gcp_w = CANVAS_WIDTH_IN - onprem_w - right_w - (0.18 if right_w else 0)

#         boxes["onprem"] = _lane_box("onprem", "On-Premise", CANVAS_LEFT_IN, top, onprem_w, height)
#         boxes["gcp"] = _lane_box("gcp", "Google Cloud Platform", CANVAS_LEFT_IN + onprem_w + 0.18, top, gcp_w, height)

#         if right_zones:
#             rz_left = CANVAS_LEFT_IN + onprem_w + 0.18 + gcp_w + 0.18
#             rz_h = height / len(right_zones)
#             for idx, zone in enumerate(right_zones):
#                 boxes[zone] = _lane_box(zone, _zone_label(zone), rz_left, top + idx * rz_h, right_w, rz_h - 0.08)
#     else:
#         n = max(1, len(zone_order))
#         lane_w = (CANVAS_WIDTH_IN - (n - 1) * 0.18) / n
#         for idx, zone in enumerate(zone_order):
#             boxes[zone] = _lane_box(zone, _zone_label(zone), CANVAS_LEFT_IN + idx * (lane_w + 0.18), top, lane_w, height)

#     return boxes


# def _lane_box(zone: str, label: str, left: float, top: float, width: float, height: float) -> Dict[str, Any]:
#     style = {
#         "onprem": ("Source / On-Premise", "EEF4FF", "8AB4F8"),
#         "gcp": ("Google Cloud Platform", "F3F8FF", "A8C7FA"),
#         "external": ("External Secret Management", EXTERNAL_FILL, EXTERNAL_BORDER),
#         "cicd": ("CI/CD & Deployment", CICD_FILL, CICD_BORDER),
#         "observability": ("Observability", OBS_FILL, OBS_BORDER),
#     }.get(zone, (label, LANE_FILL, LANE_BORDER))

#     return {
#         "id": zone,
#         "label": style[0],
#         "left": left,
#         "top": top,
#         "width": width,
#         "height": height,
#         "fill": style[1],
#         "border": style[2],
#     }


# def _zone_label(zone: str) -> str:
#     return {
#         "onprem": "Source / On-Premise",
#         "gcp": "Google Cloud Platform",
#         "external": "External Secret Management",
#         "cicd": "CI/CD & Deployment",
#         "observability": "Observability",
#     }.get(zone, zone.title())


# def _topological_order(node_order: List[str], edges: List[Dict[str, Any]]) -> List[str]:
#     indeg = {nid: 0 for nid in node_order}
#     outgoing: Dict[str, List[str]] = defaultdict(list)

#     for e in edges:
#         s, d = e["source"], e["target"]
#         if s in indeg and d in indeg:
#             outgoing[s].append(d)
#             indeg[d] += 1

#     q = deque([nid for nid in node_order if indeg.get(nid, 0) == 0])
#     result = []

#     while q:
#         cur = q.popleft()
#         result.append(cur)
#         for nxt in outgoing.get(cur, []):
#             indeg[nxt] -= 1
#             if indeg[nxt] == 0:
#                 q.append(nxt)

#     for nid in node_order:
#         if nid not in result:
#             result.append(nid)

#     return result


# # -----------------------------------------------------------------------------
# # Drawing helpers
# # -----------------------------------------------------------------------------
# def _set_text_frame(
#     shape: BaseShape,
#     text: str,
#     font_size_pt: float = 10,
#     bold: bool = False,
#     color_hex: str = DEFAULT_NODE_FONT,
#     align=PP_ALIGN.LEFT,
# ):
#     tf = shape.text_frame
#     tf.clear()
#     tf.word_wrap = True
#     tf.margin_left = Inches(0.03)
#     tf.margin_right = Inches(0.03)
#     tf.margin_top = Inches(0.01)
#     tf.margin_bottom = Inches(0.01)
#     tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE

#     p = tf.paragraphs[0]
#     p.alignment = align
#     run = p.add_run()
#     run.text = _safe_title(text)
#     run.font.name = "Calibri"
#     run.font.size = Pt(font_size_pt)
#     run.font.bold = bold
#     run.font.color.rgb = _rgb(color_hex, DEFAULT_NODE_FONT)


# def _font_for_label(label: str, width: float, height: float) -> float:
#     clean = _safe_title(label)
#     line_count = max(1, clean.count("\n") + 1)
#     length = len(clean.replace("\n", " "))

#     if width < 1.15 or height < 0.58 or length > 58 or line_count >= 3:
#         return 6.8
#     if width < 1.5 or length > 42 or line_count == 2:
#         return 7.8
#     return 8.8


# def _add_title(slide, title_text: str):
#     shape = slide.shapes.add_textbox(
#         Inches(TITLE_LEFT_IN),
#         Inches(TITLE_TOP_IN),
#         Inches(TITLE_WIDTH_IN),
#         Inches(TITLE_HEIGHT_IN),
#     )
#     _set_text_frame(shape, title_text, font_size_pt=18, bold=True, color_hex="202124")


# def _add_lane_box(slide, box: Dict[str, Any]):
#     shape = slide.shapes.add_shape(
#         MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
#         Inches(box["left"]),
#         Inches(box["top"]),
#         Inches(box["width"]),
#         Inches(box["height"]),
#     )
#     shape.fill.solid()
#     shape.fill.fore_color.rgb = _rgb(box.get("fill", LANE_FILL))
#     shape.line.color.rgb = _rgb(box.get("border", LANE_BORDER))
#     shape.line.width = Pt(1.2)

#     try:
#         shape.adjustments[0] = 0.04
#     except Exception:
#         pass

#     label_box = slide.shapes.add_textbox(
#         Inches(box["left"] + 0.08),
#         Inches(box["top"] + 0.03),
#         Inches(max(0.5, box["width"] - 0.16)),
#         Inches(0.20),
#     )
#     _set_text_frame(label_box, box.get("label", ""), font_size_pt=8.7, bold=True, color_hex=LANE_TITLE_COLOR)


# def _add_node_box(slide, box: Dict[str, float], label: str, icon_path: Optional[Path] = None):
#     shape = slide.shapes.add_shape(
#         MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
#         Inches(box["left"]),
#         Inches(box["top"]),
#         Inches(box["width"]),
#         Inches(box["height"]),
#     )
#     shape.fill.solid()
#     shape.fill.fore_color.rgb = _rgb(DEFAULT_NODE_FILL)
#     shape.line.color.rgb = _rgb(DEFAULT_NODE_BORDER)
#     shape.line.width = Pt(1.0)

#     try:
#         shape.adjustments[0] = 0.08
#     except Exception:
#         pass

#     icon_size = min(NODE_ICON_SIZE_IN, max(0.15, box["height"] * 0.36))
#     text_left = box["left"] + NODE_PADDING_IN

#     if icon_path and icon_path.exists() and box["width"] > 0.75:
#         try:
#             slide.shapes.add_picture(
#                 str(icon_path),
#                 Inches(box["left"] + NODE_PADDING_IN),
#                 Inches(box["top"] + (box["height"] - icon_size) / 2.0),
#                 width=Inches(icon_size),
#                 height=Inches(icon_size),
#             )
#             text_left = box["left"] + NODE_TEXT_LEFT_WITH_ICON_IN
#         except Exception:
#             logger.debug("Could not add icon to PPTX node: %s", icon_path, exc_info=True)

#     text_width = max(0.35, box["width"] - (text_left - box["left"]) - 0.05)
#     text_box = slide.shapes.add_textbox(
#         Inches(text_left),
#         Inches(box["top"] + 0.03),
#         Inches(text_width),
#         Inches(max(0.22, box["height"] - 0.06)),
#     )
#     _set_text_frame(
#         text_box,
#         label,
#         font_size_pt=_font_for_label(label, box["width"], box["height"]),
#         bold=True,
#         color_hex=DEFAULT_NODE_FONT,
#     )


# def _edge_points_between_boxes(src: Dict[str, float], dst: Dict[str, float]) -> List[Tuple[float, float]]:
#     sr, sl = src["left"] + src["width"], src["left"]
#     sy = src["top"] + src["height"] / 2

#     dl, dr = dst["left"], dst["left"] + dst["width"]
#     dy = dst["top"] + dst["height"] / 2

#     if dl >= sr:
#         mx = (sr + dl) / 2
#         return [(sr, sy), (mx, sy), (mx, dy), (dl, dy)]

#     if sl >= dr:
#         mx = (dr + sl) / 2
#         return [(sl, sy), (mx, sy), (mx, dy), (dr, dy)]

#     if dy >= sy:
#         return [
#             (src["left"] + src["width"] / 2, src["top"] + src["height"]),
#             (dst["left"] + dst["width"] / 2, dst["top"]),
#         ]

#     return [
#         (src["left"] + src["width"] / 2, src["top"]),
#         (dst["left"] + dst["width"] / 2, dst["top"] + dst["height"]),
#     ]


# def _shorten_last_segment_for_arrowhead(
#     points: List[Tuple[float, float]],
#     gap: float = ARROW_TIP_GAP_IN,
# ) -> List[Tuple[float, float]]:
#     if len(points) < 2:
#         return points

#     adjusted = list(points)
#     x1, y1 = adjusted[-2]
#     x2, y2 = adjusted[-1]

#     dx = x2 - x1
#     dy = y2 - y1
#     length = math.sqrt(dx * dx + dy * dy)

#     if length <= 0.001:
#         return adjusted

#     usable_gap = min(gap, length * 0.35)
#     adjusted[-1] = (
#         x2 - (dx / length) * usable_gap,
#         y2 - (dy / length) * usable_gap,
#     )
#     return adjusted


# def _set_line_end_arrow(connector: BaseShape):
#     """
#     Add a visible PowerPoint arrowhead to connector end using OOXML.
#     """
#     try:
#         ln = connector._element.spPr.get_or_add_ln()

#         for child in list(ln):
#             if child.tag.endswith("}tailEnd"):
#                 ln.remove(child)

#         tail_end = OxmlElement("a:tailEnd")
#         tail_end.set("type", ARROW_HEAD_TYPE)
#         tail_end.set("w", ARROW_HEAD_WIDTH)
#         tail_end.set("len", ARROW_HEAD_LENGTH)
#         ln.append(tail_end)
#     except Exception:
#         logger.debug("Could not apply arrowhead to connector.", exc_info=True)


# def _add_line_segment(
#     slide,
#     p1: Tuple[float, float],
#     p2: Tuple[float, float],
#     color_hex: str,
#     add_arrowhead: bool = False,
#     dash_style: Optional[MSO_LINE_DASH_STYLE] = None,
#     width_pt: float = DEFAULT_EDGE_WIDTH_PT,
# ):
#     connector = slide.shapes.add_connector(
#         MSO_CONNECTOR.STRAIGHT,
#         Inches(p1[0]),
#         Inches(p1[1]),
#         Inches(p2[0]),
#         Inches(p2[1]),
#     )
#     connector.line.color.rgb = _rgb(color_hex, DEFAULT_EDGE_COLOR)
#     connector.line.width = Pt(width_pt)

#     if dash_style is not None:
#         try:
#             connector.line.dash_style = dash_style
#         except Exception:
#             logger.debug("Could not apply line dash style.", exc_info=True)

#     if add_arrowhead:
#         _set_line_end_arrow(connector)

#     return connector


# def _add_polyline_arrow(
#     slide,
#     points: List[Tuple[float, float]],
#     color_hex: str,
#     dash_style: Optional[MSO_LINE_DASH_STYLE] = None,
#     width_pt: float = DEFAULT_EDGE_WIDTH_PT,
# ):
#     if len(points) < 2:
#         return

#     clean = [points[0]]
#     for point in points[1:]:
#         if abs(point[0] - clean[-1][0]) > 0.01 or abs(point[1] - clean[-1][1]) > 0.01:
#             clean.append(point)

#     if len(clean) < 2:
#         return

#     clean = _shorten_last_segment_for_arrowhead(clean)

#     for idx in range(len(clean) - 1):
#         is_last_segment = idx == len(clean) - 2
#         _add_line_segment(
#             slide,
#             clean[idx],
#             clean[idx + 1],
#             color_hex,
#             add_arrowhead=is_last_segment,
#             dash_style=dash_style,
#             width_pt=width_pt,
#         )


# def _add_legend(slide, legend_items: List[Tuple[str, str, Optional[MSO_LINE_DASH_STYLE]]]):
#     if not legend_items:
#         return

#     title_box = slide.shapes.add_textbox(
#         Inches(LEGEND_LEFT_IN),
#         Inches(LEGEND_TOP_IN),
#         Inches(1.15),
#         Inches(0.18),
#     )
#     _set_text_frame(title_box, "Flow Legend", font_size_pt=9.2, bold=True, color_hex="202124")

#     x, y, max_x = LEGEND_LEFT_IN, LEGEND_TOP_IN + 0.24, LEGEND_LEFT_IN + LEGEND_WIDTH_IN

#     for color, label, dash_style in legend_items[:18]:
#         label = _safe_title(label)
#         item_w = max(1.55, min(3.30, 0.62 + len(label) * 0.050))
#         if x + item_w > max_x:
#             x, y = LEGEND_LEFT_IN, y + 0.22
#         if y > LEGEND_TOP_IN + LEGEND_MAX_HEIGHT_IN:
#             break

#         chip = slide.shapes.add_shape(
#             MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
#             Inches(x),
#             Inches(y),
#             Inches(item_w),
#             Inches(0.18),
#         )
#         chip.fill.solid()
#         chip.fill.fore_color.rgb = _rgb(LEGEND_FILL)
#         chip.line.color.rgb = _rgb(LEGEND_BORDER)
#         chip.line.width = Pt(0.6)

#         line_y = y + 0.09
#         _add_line_segment(
#             slide,
#             (x + 0.05, line_y),
#             (x + 0.18, line_y),
#             color,
#             add_arrowhead=True,
#             dash_style=dash_style,
#             width_pt=1.35,
#         )

#         label_box = slide.shapes.add_textbox(
#             Inches(x + 0.22),
#             Inches(y + 0.01),
#             Inches(max(0.3, item_w - 0.24)),
#             Inches(0.15),
#         )
#         _set_text_frame(label_box, label, font_size_pt=7.2, bold=False, color_hex="202124")
#         x += item_w + 0.08


# def _add_fallback_preview_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str):
#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, f"{slide_title} (Preview)")

#     try:
#         assets = render_graphviz_to_assets(diagram_text, output_dir=output_dir, base_name=base_name)
#     except Exception:
#         logger.exception("Fallback preview asset generation failed.")
#         assets = None

#     if assets and assets.get("png_path") and Path(assets["png_path"]).exists():
#         slide.shapes.add_picture(
#             str(assets["png_path"]),
#             Inches(CANVAS_LEFT_IN),
#             Inches(CANVAS_TOP_IN),
#             width=Inches(PREVIEW_MAX_WIDTH_IN),
#         )
#         _add_legend(slide, assets.get("legend", []))


# # -----------------------------------------------------------------------------
# # Diagram extraction
# # -----------------------------------------------------------------------------
# def _is_diagram_text(value: Any) -> bool:
#     return isinstance(value, str) and normalize_to_graphviz_dot(_clean_internal_tokens(value)) is not None


# def _walk_for_diagrams(value: Any, path_parts: List[str], results: List[Dict[str, Any]]):
#     if value is None:
#         return

#     if isinstance(value, str):
#         if _is_diagram_text(value):
#             results.append({
#                 "title": " / ".join(path_parts) if path_parts else "Diagram",
#                 "diagram_text": _clean_internal_tokens(value),
#                 "section_key": _safe_name(path_parts[-1] if path_parts else "diagram"),
#             })
#         return

#     if isinstance(value, list):
#         for idx, item in enumerate(value, start=1):
#             if _is_diagram_text(item):
#                 results.append({
#                     "title": " / ".join(path_parts) if path_parts else f"Diagram {idx}",
#                     "diagram_text": _clean_internal_tokens(item),
#                     "section_key": _safe_name(path_parts[-1] if path_parts else f"diagram_{idx}"),
#                 })
#             elif isinstance(item, (dict, list)):
#                 _walk_for_diagrams(item, path_parts, results)
#         return

#     if isinstance(value, dict):
#         for key, child in value.items():
#             key_text = str(key).replace("_", " ").title()
#             next_path = path_parts if str(key).lower() == "diagrams" else path_parts + [key_text]
#             _walk_for_diagrams(child, next_path, results)


# def extract_diagram_specs(
#     data: Dict[str, Any],
#     hld_model: Any,
#     selected_sections_set: Optional[Set[str]] = None,
# ) -> List[Dict[str, Any]]:
#     results: List[Dict[str, Any]] = []

#     for field_name, field_info in hld_model.model_fields.items():
#         if not should_include_section(field_name, field_info, selected_sections_set or set()):
#             continue
#         section_title = field_info.title or field_name.replace("_", " ").title()
#         _walk_for_diagrams(data.get(field_name), [section_title], results)

#     counts: Dict[str, int] = defaultdict(int)
#     for item in results:
#         counts[item["title"]] += 1

#     seen: Dict[str, int] = defaultdict(int)
#     for item in results:
#         if counts[item["title"]] > 1:
#             seen[item["title"]] += 1
#             item["title"] = f'{item["title"]} ({seen[item["title"]]})'

#     return results


# # -----------------------------------------------------------------------------
# # Slide rendering
# # -----------------------------------------------------------------------------
# def add_diagram_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str):
#     try:
#         model = _build_editable_model(diagram_text)
#         if not model["nodes"]:
#             raise ValueError("No editable nodes parsed")
#         layout = _layout_model(model)
#         placements = layout["placements"]
#         lane_boxes = layout["lane_boxes"]
#     except Exception:
#         logger.exception("Editable PPTX rendering failed for slide '%s'. Falling back to preview image.", slide_title)
#         _add_fallback_preview_slide(prs, slide_title, diagram_text, output_dir, base_name)
#         return

#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, slide_title)

#     for box in lane_boxes:
#         _add_lane_box(slide, box)

#     # Draw edges first, then nodes on top
#     for edge in model["edges"]:
#         src_box = placements.get(edge["source"])
#         dst_box = placements.get(edge["target"])
#         if src_box and dst_box:
#             _add_polyline_arrow(
#                 slide,
#                 _edge_points_between_boxes(src_box, dst_box),
#                 edge.get("color") or DEFAULT_EDGE_COLOR,
#                 dash_style=edge.get("dash_style"),
#                 width_pt=edge.get("width_pt", DEFAULT_EDGE_WIDTH_PT),
#             )

#     for node in model["nodes"]:
#         box = placements.get(node["id"])
#         if not box:
#             continue
#         icon_path = _preferred_picture_icon_path(node.get("icon_hint") or node.get("label") or node["id"])
#         _add_node_box(slide, box, node.get("label") or node["id"], icon_path)

#     _add_legend(slide, model.get("legend", []))


# # -----------------------------------------------------------------------------
# # Public builder
# # -----------------------------------------------------------------------------
# def build_editable_pptx(
#     data: Dict[str, Any],
#     pptx_path: str,
#     hld_model: Any,
#     selected_sections_set: Optional[Set[str]] = None,
# ):
#     output_path = Path(pptx_path)
#     output_path.parent.mkdir(parents=True, exist_ok=True)

#     diagram_assets_dir = output_path.parent / f"{output_path.stem}_diagram_assets"
#     diagram_assets_dir.mkdir(parents=True, exist_ok=True)

#     diagram_specs = extract_diagram_specs(
#         data=data,
#         hld_model=hld_model,
#         selected_sections_set=selected_sections_set,
#     )

#     prs = Presentation()
#     prs.slide_width = Inches(SLIDE_WIDTH_IN)
#     prs.slide_height = Inches(SLIDE_HEIGHT_IN)

#     if not diagram_specs:
#         slide = prs.slides.add_slide(prs.slide_layouts[6])
#         _add_title(slide, "Editable Diagrams")
#         msg = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11.5), Inches(1.0))
#         _set_text_frame(
#             msg,
#             "No diagrams were detected in the selected sections.",
#             font_size_pt=16,
#             bold=True,
#             color_hex="5F6368",
#             align=PP_ALIGN.CENTER,
#         )
#         prs.save(str(output_path))
#         return

#     for counter, spec in enumerate(diagram_specs, start=1):
#         add_diagram_slide(
#             prs,
#             spec["title"],
#             spec["diagram_text"],
#             diagram_assets_dir,
#             f"{_safe_name(spec['section_key'])}_{counter:02d}",
#         )

#     prs.save(str(output_path))
#     logger.info("Editable PPTX generated: %s", output_path)


# 5th Working Code


# from __future__ import annotations

# import html
# import json
# import logging
# import math
# import re
# from collections import defaultdict, deque
# from pathlib import Path
# from typing import Any, Dict, List, Optional, Set, Tuple

# from pptx import Presentation
# from pptx.dml.color import RGBColor
# from pptx.enum.dml import MSO_LINE_DASH_STYLE
# from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
# from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
# from pptx.oxml.xmlchemy import OxmlElement
# from pptx.shapes.base import BaseShape
# from pptx.util import Inches, Pt

# from .config import ICON_DIR
# from .graphviz_renderer import normalize_to_graphviz_dot, render_graphviz_to_assets
# from .schema_utils import should_include_section
# from .text_sanitizer import renderer_safe_plain_text

# logger = logging.getLogger(__name__)

# # =============================================================================
# # Editable PPTX Diagram Renderer - aligned swimlane version
# # =============================================================================
# # Features:
# #   - Editable PowerPoint diagrams
# #   - Generic swimlane layout
# #   - Rank-aligned flow layout, including Process View / numbered sequence views
# #   - Clear same-row horizontal arrows for sequence/process diagrams
# #   - Generic fan-in/fan-out port separation
# #   - Dynamic icon resolution from DOT attrs, optional icon manifest, and icon filenames
# #   - No service-specific routing/styling hardcoding
# # =============================================================================

# SLIDE_WIDTH_IN = 13.333
# SLIDE_HEIGHT_IN = 7.5
# TITLE_LEFT_IN = 0.35
# TITLE_TOP_IN = 0.12
# TITLE_WIDTH_IN = 12.6
# TITLE_HEIGHT_IN = 0.42
# CANVAS_LEFT_IN = 0.35
# CANVAS_TOP_IN = 0.68
# CANVAS_WIDTH_IN = 12.3
# CANVAS_HEIGHT_IN = 5.25
# LEGEND_LEFT_IN = 0.35
# LEGEND_TOP_IN = 6.12
# LEGEND_WIDTH_IN = 12.3
# LEGEND_MAX_HEIGHT_IN = 1.22

# NODE_W = 1.78
# NODE_H = 0.78
# NODE_ICON_SIZE_IN = 0.22
# NODE_PADDING_IN = 0.07
# NODE_TEXT_LEFT_WITH_ICON_IN = 0.33

# LANE_PAD = 0.18
# LANE_TITLE_H = 0.26
# H_GAP = 0.30
# V_GAP = 0.24

# DEFAULT_NODE_FILL = "F8F9FA"
# DEFAULT_NODE_BORDER = "DADCE0"
# DEFAULT_NODE_FONT = "202124"

# LANE_FILL = "F3F8FF"
# LANE_BORDER = "A8C7FA"
# LANE_TITLE_COLOR = "174EA6"
# EXTERNAL_FILL = "FFF8E1"
# EXTERNAL_BORDER = "F9AB00"
# CICD_FILL = "FCE8E6"
# CICD_BORDER = "F28B82"
# OBS_FILL = "F3E8FD"
# OBS_BORDER = "D7AEFB"

# DEFAULT_EDGE_COLOR = "5F6368"
# DEFAULT_EDGE_WIDTH_PT = 1.55
# LEGEND_FILL = "FFFFFF"
# LEGEND_BORDER = "DADCE0"

# PREVIEW_MAX_WIDTH_IN = 12.0
# SUPPORTED_PICTURE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif"}
# PREFERRED_ICON_EXTS = [".png", ".jpg", ".jpeg", ".bmp", ".gif"]

# COMMENT_LINE_RE = re.compile(r"//.*?$", re.MULTILINE)
# BLOCK_COMMENT_RE = re.compile(r"/\*.*?\*/", re.DOTALL)
# EDGE_RE = re.compile(r'("[^"]+"|[A-Za-z0-9_.:-]+)\s*(->|--)\s*("[^"]+"|[A-Za-z0-9_.:-]+)')

# FLOW_PALETTE = [
#     "E60000", "1D70B8", "28A197", "F47738", "6F35A5",
#     "007C89", "D53880", "85994B", "B58840", "5F6368"
# ]

# COLOR_NAME_MAP = {
#     "red": "E60000",
#     "blue": "1D70B8",
#     "green": "28A197",
#     "orange": "F47738",
#     "purple": "6F35A5",
#     "teal": "007C89",
#     "gray": "5F6368",
#     "grey": "5F6368",
#     "black": "202124",
#     "lightgrey": "DADCE0",
#     "lightgray": "DADCE0",
# }

# ARROW_HEAD_TYPE = "triangle"
# ARROW_HEAD_WIDTH = "med"
# ARROW_HEAD_LENGTH = "med"
# ARROW_TIP_GAP_IN = 0.07

# # Generic routing controls. No service-specific routing hardcoding.
# EDGE_PORT_SPREAD_IN = 0.16
# EDGE_CHANNEL_SPACING_IN = 0.20
# EDGE_MARGIN_FROM_NODE_IN = 0.07
# EDGE_PAIR_SEPARATION_IN = 0.13
# EDGE_MIN_CHANNEL_CLEARANCE_IN = 0.24
# SAME_ROW_TOLERANCE_IN = 0.16

# # Layout controls for readable process/sequence diagrams.
# MIN_RANK_NODE_W = 1.05
# RANK_COL_GAP = 0.14
# ISOLATED_NODE_SCALE = 0.88
# ISOLATED_ROW_GAP = 0.12

# # Icon matching controls. Uses existing files/manifest; no Python service-name mapping.
# ICON_MATCH_MIN_SCORE = 1.75
# ICON_MATCH_TOKEN_BONUS = 1.0
# ICON_MATCH_CONSECUTIVE_BONUS = 0.75
# ICON_MATCH_EXACT_STEM_BONUS = 4.0
# ICON_MATCH_PREFIX_BONUS = 1.2
# ICON_MANIFEST_NAMES = ["icon_manifest.json", "icons_manifest.json"]

# # Broad lane classification only; not used for service-specific routing/styling.
# ZONE_RULES = {
#     "onprem": ["on-prem", "on premise", "source", "datacenter", "data center"],
#     "external": ["external", "third party", "3rd party"],
#     "cicd": ["ci/cd", "cicd", "deployment", "release", "pipeline"],
#     "observability": ["observability", "monitor", "logging", "alerting", "metrics"],
#     "gcp": ["gcp", "google cloud", "cloud", "target"],
# }


# # -----------------------------------------------------------------------------
# # Basic helpers
# # -----------------------------------------------------------------------------
# def _clean_internal_tokens(text: Any) -> str:
#     value = str(text or "")
#     for token in [
#         "AIASECTIONBLOCKSTARTTOKEN",
#         "AIASECTIONBLOCKENDTOKEN",
#         "[[AIASECTIONBLOCKSTARTTOKEN]]",
#         "[[AIASECTIONBLOCKENDTOKEN]]",
#         "SECTION_BLOCK_START",
#         "SECTION_BLOCK_END",
#         "[[SECTION_BLOCK_START]]",
#         "[[SECTION_BLOCK_END]]",
#         "[[SECTIONBLOCKSTART]]",
#         "[[SECTIONBLOCKEND]]",
#     ]:
#         value = value.replace(token, "")
#     return html.unescape(value.strip())


# def _safe_title(text: Any) -> str:
#     return renderer_safe_plain_text(_clean_internal_tokens(text)).strip() or "Diagram"


# def _safe_name(text: Any) -> str:
#     return re.sub(r"[^a-zA-Z0-9_-]+", "_", str(text or "")).strip("_").lower() or "diagram"


# def _clean_label(value: Any) -> str:
#     text = html.unescape(str(value or "").strip()).strip('"')
#     text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
#     text = re.sub(r"<[^>]+>", " ", text)
#     text = text.replace("\\\\N", "\n").replace("\\\\n", "\n")
#     text = text.replace("\\N", "\n").replace("\\n", "\n")
#     text = text.replace("&nbsp;", " ")
#     text = re.sub(r"\s+\n", "\n", text)
#     text = re.sub(r"\n\s+", "\n", text)
#     text = re.sub(r"[ \t]+", " ", text)
#     return text.strip()


# def _label(raw: Optional[str], fallback: str) -> str:
#     label = _clean_label(raw or "")
#     if not label:
#         label = str(fallback or "").replace("_", " ").replace("-", " ").title()
#     return label[:180]


# def _safe_color(value: Optional[str], default: str = DEFAULT_EDGE_COLOR) -> str:
#     raw = str(value or "").strip().strip('"').lstrip("#")
#     if not raw:
#         return default
#     if raw.lower() in COLOR_NAME_MAP:
#         return COLOR_NAME_MAP[raw.lower()]
#     raw = raw.upper()
#     return raw if len(raw) == 6 and all(ch in "0123456789ABCDEF" for ch in raw) else default


# def _rgb(hex_color: str, default: str = "000000") -> RGBColor:
#     return RGBColor.from_string(_safe_color(hex_color, default))


# def _safe_float(value: Any, default: float) -> float:
#     try:
#         return float(str(value).strip())
#     except Exception:
#         return default


# def _is_pseudo_node(node_id: str, label: Optional[str] = None) -> bool:
#     text = _clean_label(label if label is not None else node_id).lower().strip()
#     raw = str(node_id or "").lower().strip()
#     if text in {"", "n/a", "na", "none", "null"}:
#         return True
#     if raw in {
#         "node", "edge", "graph", "digraph", "subgraph", "rank",
#         "label", "style", "color", "fillcolor", "fontcolor", "rankdir"
#     }:
#         return True
#     if raw.startswith((
#         "digraph", "graph", "subgraph", "cluster_", "cluster ",
#         "label=", "style=", "color=", "fillcolor=", "fontcolor=", "rankdir="
#     )):
#         return True
#     if text.startswith((
#         "label=", "style=", "color=", "fillcolor=", "fontcolor=",
#         "rankdir=", "fontsize=", "fontname=", "margin=", "pad="
#     )):
#         return True
#     if text in {"filled", "lightgrey", "lightgray", "blue", "green", "purple", "orange", "red", "gray", "grey"}:
#         return True
#     return False


# def _is_graphviz_control_statement(statement: str) -> bool:
#     s = html.unescape(str(statement or "")).strip().lower()
#     if not s or s in {"{", "}"}:
#         return True
#     return s.startswith((
#         "digraph ", "graph ", "subgraph ", "node ", "edge ", "rank ",
#         "rankdir", "label", "style", "color", "fillcolor", "fontcolor",
#         "fontsize", "fontname", "margin", "pad", "splines", "bgcolor"
#     ))


# # -----------------------------------------------------------------------------
# # Icon helpers - dynamic, based on document attrs / manifest / icon filenames
# # -----------------------------------------------------------------------------
# def _icon_tokens(value: Any) -> List[str]:
#     text = _clean_label(value).lower()
#     text = re.sub(r"[^a-z0-9]+", " ", text)
#     raw_tokens = [token.strip() for token in text.split() if token.strip()]
#     tokens: List[str] = []
#     for token in raw_tokens:
#         if len(token) < 2:
#             continue
#         tokens.append(token)
#         if token.endswith("ies") and len(token) > 4:
#             tokens.append(token[:-3] + "y")
#         elif token.endswith("s") and len(token) > 3:
#             tokens.append(token[:-1])
#     seen = set()
#     result = []
#     for token in tokens:
#         if token not in seen:
#             seen.add(token)
#             result.append(token)
#     return result


# def _load_icon_manifest() -> Dict[str, List[str]]:
#     manifest: Dict[str, List[str]] = {}
#     try:
#         for manifest_name in ICON_MANIFEST_NAMES:
#             path = ICON_DIR / manifest_name
#             if not path.exists():
#                 continue
#             data = json.loads(path.read_text(encoding="utf-8"))
#             if not isinstance(data, dict):
#                 continue
#             for file_name, aliases in data.items():
#                 values: List[str] = []
#                 if isinstance(aliases, str):
#                     values.append(aliases)
#                 elif isinstance(aliases, list):
#                     values.extend(str(item) for item in aliases)
#                 values.append(str(file_name))
#                 manifest[str(file_name)] = values
#     except Exception:
#         logger.debug("Could not load icon manifest.", exc_info=True)
#     return manifest


# def _score_icon_candidate(icon_hint: str, candidate_path: Path, manifest: Optional[Dict[str, List[str]]] = None) -> float:
#     hint_tokens = _icon_tokens(icon_hint)
#     candidate_text_parts = [candidate_path.stem]
#     if manifest:
#         aliases = manifest.get(candidate_path.name) or manifest.get(candidate_path.stem) or []
#         candidate_text_parts.extend(aliases)
#     stem_tokens = _icon_tokens(" ".join(str(part) for part in candidate_text_parts))
#     if not hint_tokens or not stem_tokens:
#         return 0.0
#     hint_set = set(hint_tokens)
#     stem_set = set(stem_tokens)
#     score = 0.0
#     if _safe_name(icon_hint) == _safe_name(candidate_path.stem):
#         score += ICON_MATCH_EXACT_STEM_BONUS
#     score += len(hint_set.intersection(stem_set)) * ICON_MATCH_TOKEN_BONUS
#     for stem_token in stem_tokens:
#         for hint_token in hint_tokens:
#             if stem_token == hint_token:
#                 continue
#             if hint_token.startswith(stem_token) or stem_token.startswith(hint_token):
#                 score += ICON_MATCH_PREFIX_BONUS
#                 break
#     for idx in range(len(stem_tokens) - 1):
#         first = stem_tokens[idx]
#         second = stem_tokens[idx + 1]
#         for hidx in range(len(hint_tokens) - 1):
#             if hint_tokens[hidx] == first and hint_tokens[hidx + 1] == second:
#                 score += ICON_MATCH_CONSECUTIVE_BONUS
#     score -= max(0, len(_icon_tokens(candidate_path.stem)) - 3) * 0.05
#     return score


# def _try_existing_icon_resolver(icon_hint: str) -> Optional[Path]:
#     try:
#         from . import icon_resolver as ir
#     except Exception:
#         return None
#     for fn_name in ["resolve_icon_path", "resolve_icon", "get_icon_path", "find_icon_path"]:
#         fn = getattr(ir, fn_name, None)
#         if not callable(fn):
#             continue
#         for args in [(icon_hint,), (icon_hint, True), (icon_hint, None)]:
#             try:
#                 result = fn(*args)
#                 if result:
#                     path = Path(str(result))
#                     if path.exists():
#                         return path
#             except TypeError:
#                 continue
#             except Exception:
#                 logger.debug("Icon resolver failed: %s", fn_name, exc_info=True)
#     return None


# def _search_icon_dir(icon_hint: str) -> Optional[Path]:
#     if not icon_hint:
#         return None
#     raw_hint = str(icon_hint or "").strip()
#     safe_hint = _safe_name(raw_hint)
#     candidate_names = {
#         safe_hint,
#         raw_hint,
#         raw_hint.lower(),
#         raw_hint.replace(" ", "_").lower(),
#         raw_hint.replace("-", "_").lower(),
#         raw_hint.replace("/", "_").lower(),
#     }
#     for name in candidate_names:
#         if not name:
#             continue
#         for ext in [".png", ".jpg", ".jpeg", ".bmp", ".gif", ".svg"]:
#             candidate = ICON_DIR / f"{name}{ext}"
#             if candidate.exists():
#                 return candidate
#     best_path: Optional[Path] = None
#     best_score = 0.0
#     manifest = _load_icon_manifest()
#     try:
#         for path in ICON_DIR.rglob("*"):
#             if not path.is_file():
#                 continue
#             if path.suffix.lower() not in SUPPORTED_PICTURE_EXTS and path.suffix.lower() != ".svg":
#                 continue
#             score = _score_icon_candidate(raw_hint, path, manifest=manifest)
#             if score > best_score:
#                 best_score = score
#                 best_path = path
#     except Exception:
#         logger.debug("Icon directory search failed", exc_info=True)
#     if best_path and best_score >= ICON_MATCH_MIN_SCORE:
#         logger.debug("Matched icon '%s' for hint '%s' with score %.2f", best_path, raw_hint, best_score)
#         return best_path
#     logger.debug("No icon matched for hint '%s'. Best score was %.2f", raw_hint, best_score)
#     return None


# def _preferred_picture_icon_path(icon_hint: str) -> Optional[Path]:
#     path = _try_existing_icon_resolver(icon_hint) or _search_icon_dir(icon_hint)
#     if not path:
#         return None
#     if path.suffix.lower() in SUPPORTED_PICTURE_EXTS:
#         return path
#     if path.suffix.lower() == ".svg":
#         for ext in PREFERRED_ICON_EXTS:
#             sibling = path.with_suffix(ext)
#             if sibling.exists():
#                 return sibling
#     return None


# def _guess_icon_hint(node_id: str, node_label: str, attrs: Dict[str, str]) -> str:
#     for attr_name in ["image", "icon", "imagepath", "iconpath"]:
#         image_attr = attrs.get(attr_name)
#         if image_attr:
#             stem = Path(str(image_attr)).stem
#             if stem:
#                 return stem
#     for candidate in [node_label.replace("\n", " "), node_label.split("\n", 1)[0], node_id]:
#         cleaned = _clean_label(candidate)
#         if cleaned:
#             return cleaned
#     return _safe_name(node_id or node_label or "node")


# # -----------------------------------------------------------------------------
# # DOT parsing helpers
# # -----------------------------------------------------------------------------
# def _strip_dot_comments(dot_text: str) -> str:
#     text = BLOCK_COMMENT_RE.sub("", dot_text or "")
#     return COMMENT_LINE_RE.sub("", text)


# def _outer_graph_body(dot: str) -> str:
#     text = dot or ""
#     start = text.find("{")
#     end = text.rfind("}")
#     if start != -1 and end != -1 and end > start:
#         return text[start + 1:end]
#     return text


# def _split_statements(dot_text: str) -> List[str]:
#     statements: List[str] = []
#     buf: List[str] = []
#     bracket_depth = 0
#     brace_depth = 0
#     angle_depth = 0
#     for ch in dot_text or "":
#         if ch == "[":
#             bracket_depth += 1
#         elif ch == "]":
#             bracket_depth = max(0, bracket_depth - 1)
#         elif ch == "{":
#             brace_depth += 1
#         elif ch == "}":
#             brace_depth = max(0, brace_depth - 1)
#         elif ch == "<":
#             angle_depth += 1
#         elif ch == ">":
#             angle_depth = max(0, angle_depth - 1)
#         if ch == ";" and bracket_depth == 0 and brace_depth == 0 and angle_depth == 0:
#             stmt = "".join(buf).strip()
#             if stmt:
#                 statements.append(stmt)
#             buf = []
#         else:
#             buf.append(ch)
#     tail = "".join(buf).strip()
#     if tail:
#         statements.append(tail)
#     return statements


# def _extract_attr_pairs(attr_text: str) -> Dict[str, str]:
#     attrs: Dict[str, str] = {}
#     if not attr_text:
#         return attrs
#     text = html.unescape(str(attr_text))
#     for match in re.finditer(r'([A-Za-z0-9_:-]+)\s*=\s*("[^"]*"|<[^>]*>|[^,\n\r;]+)', text):
#         key = match.group(1).lower().strip()
#         value = match.group(2).strip().strip('"')
#         if value.startswith("<") and value.endswith(">"):
#             value = value[1:-1]
#         attrs[key] = value.strip()
#     return attrs


# def _extract_bracket_attr(statement: str) -> Tuple[str, Dict[str, str]]:
#     statement = html.unescape((statement or "").strip().rstrip(";").strip())
#     if "[" not in statement or "]" not in statement:
#         return statement, {}
#     start = statement.find("[")
#     end = statement.rfind("]")
#     if end <= start:
#         return statement, {}
#     return statement[:start].strip(), _extract_attr_pairs(statement[start + 1:end])


# def _unquote_id(value: str) -> str:
#     return str(value or "").strip().strip('"').strip()


# def _find_subgraph_blocks(dot_text: str) -> List[Tuple[str, str]]:
#     text = dot_text or ""
#     results: List[Tuple[str, str]] = []
#     idx = 0
#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             break
#         start = idx + match.start()
#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break
#         header = text[start:brace_start].strip()
#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1
#         if end >= len(text):
#             break
#         results.append((header, text[brace_start + 1:end]))
#         idx = end + 1
#     return results


# def _remove_subgraph_blocks(dot_text: str) -> str:
#     text = dot_text or ""
#     pieces: List[str] = []
#     idx = 0
#     while True:
#         match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
#         if not match:
#             pieces.append(text[idx:])
#             break
#         start = idx + match.start()
#         pieces.append(text[idx:start])
#         brace_start = text.find("{", start)
#         if brace_start == -1:
#             break
#         depth = 0
#         end = brace_start
#         while end < len(text):
#             if text[end] == "{":
#                 depth += 1
#             elif text[end] == "}":
#                 depth -= 1
#                 if depth == 0:
#                     break
#             end += 1
#         if end >= len(text):
#             break
#         idx = end + 1
#     return "".join(pieces)


# def _extract_cluster_label(body: str, fallback_name: str) -> str:
#     for stmt in _split_statements(body):
#         s = html.unescape(stmt or "").strip()
#         if s.lower().startswith("label") and "=" in s:
#             try:
#                 value = s.split("=", 1)[1].strip().strip('"').strip()
#                 if value:
#                     return _clean_label(value)
#             except Exception:
#                 pass
#     return fallback_name


# def _extract_cluster_nodes(body: str) -> Set[str]:
#     node_ids: Set[str] = set()
#     for stmt in _split_statements(body):
#         s = html.unescape(stmt or "").strip()
#         if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s):
#             continue
#         prefix, attrs = _extract_bracket_attr(s)
#         node_id = _unquote_id(prefix)
#         label = attrs.get("label", node_id)
#         if not node_id or _is_pseudo_node(node_id, label):
#             continue
#         if re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
#             node_ids.add(node_id)
#     for match in EDGE_RE.finditer(body or ""):
#         for idx in [1, 3]:
#             node_id = _unquote_id(match.group(idx))
#             if node_id and not _is_pseudo_node(node_id):
#                 node_ids.add(node_id)
#     return node_ids


# def _parse_clusters_from_dot(dot: str) -> List[Dict[str, Any]]:
#     clusters: List[Dict[str, Any]] = []
#     for header, body in _find_subgraph_blocks(dot):
#         fallback_name = _unquote_id(header.replace("subgraph", "", 1).strip()) or "Cluster"
#         cluster_label = _extract_cluster_label(body, fallback_name)
#         cluster_nodes = sorted(_extract_cluster_nodes(body))
#         if cluster_nodes and not _is_pseudo_node(fallback_name, cluster_label):
#             clusters.append({"id": _safe_name(fallback_name), "label": cluster_label, "nodes": cluster_nodes})
#     return clusters


# def _parse_node_attrs_from_dot(dot: str) -> Tuple[Dict[str, Dict[str, str]], List[str]]:
#     attrs_by_node: Dict[str, Dict[str, str]] = {}
#     node_order: List[str] = []

#     def parse_node_statements(text: str) -> None:
#         for stmt in _split_statements(text):
#             s = html.unescape(stmt or "").strip()
#             if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s):
#                 continue
#             prefix, attrs = _extract_bracket_attr(s)
#             node_id = _unquote_id(prefix)
#             if not node_id:
#                 continue
#             label = attrs.get("label", node_id)
#             if _is_pseudo_node(node_id, label):
#                 continue
#             if not re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
#                 continue
#             if node_id not in attrs_by_node:
#                 node_order.append(node_id)
#             attrs_by_node[node_id] = attrs

#     body = _outer_graph_body(dot)
#     parse_node_statements(_remove_subgraph_blocks(body))
#     for _header, subgraph_body in _find_subgraph_blocks(dot):
#         parse_node_statements(subgraph_body)
#     return attrs_by_node, node_order


# def _parse_edge_attrs_from_dot(dot: str) -> Dict[Tuple[str, str], Dict[str, str]]:
#     attrs_by_edge: Dict[Tuple[str, str], Dict[str, str]] = {}

#     def parse_edge_statements(text: str) -> None:
#         for stmt in _split_statements(text):
#             s = html.unescape(stmt or "").strip()
#             if "->" not in s and "--" not in s:
#                 continue
#             prefix, attrs = _extract_bracket_attr(s)
#             tokens = re.findall(r'"[^"]+"|[A-Za-z0-9_.:-]+', prefix)
#             if len(tokens) < 2:
#                 continue
#             for src_raw, dst_raw in zip(tokens[:-1], tokens[1:]):
#                 src = _unquote_id(src_raw)
#                 dst = _unquote_id(dst_raw)
#                 if not src or not dst or _is_pseudo_node(src) or _is_pseudo_node(dst):
#                     continue
#                 attrs_by_edge[(src, dst)] = attrs

#     body = _outer_graph_body(dot)
#     parse_edge_statements(_remove_subgraph_blocks(body))
#     for _header, subgraph_body in _find_subgraph_blocks(dot):
#         parse_edge_statements(subgraph_body)
#     for match in EDGE_RE.finditer(dot or ""):
#         src = _unquote_id(match.group(1))
#         dst = _unquote_id(match.group(3))
#         if src and dst and not _is_pseudo_node(src) and not _is_pseudo_node(dst):
#             attrs_by_edge.setdefault((src, dst), {})
#     return attrs_by_edge


# # -----------------------------------------------------------------------------
# # Model
# # -----------------------------------------------------------------------------
# def _infer_zone(node_id: str, label: str) -> str:
#     text = f"{node_id} {label}".lower()
#     for zone, keywords in ZONE_RULES.items():
#         if any(keyword in text for keyword in keywords):
#             return zone
#     return "gcp"


# def _zone_from_cluster_label(cluster_label: str, default: str = "gcp") -> str:
#     text = _clean_label(cluster_label).lower()
#     for zone, keywords in ZONE_RULES.items():
#         if any(keyword in text for keyword in keywords):
#             return zone
#     return default


# def _dash_style_from_raw_style(raw_style: Optional[str]) -> Optional[MSO_LINE_DASH_STYLE]:
#     style = str(raw_style or "").lower()
#     if not style:
#         return None
#     if "dotted" in style:
#         return MSO_LINE_DASH_STYLE.ROUND_DOT
#     if "dashed" in style:
#         return MSO_LINE_DASH_STYLE.DASH
#     if "dashdot" in style or "dash-dot" in style:
#         return MSO_LINE_DASH_STYLE.DASH_DOT
#     return None


# def _edge_default_label(edge: Dict[str, Any], nodes: Dict[str, Dict[str, Any]]) -> str:
#     src_label = _clean_label(nodes.get(edge["source"], {}).get("label", edge["source"])).replace("\n", " ")[:28]
#     dst_label = _clean_label(nodes.get(edge["target"], {}).get("label", edge["target"])).replace("\n", " ")[:28]
#     return f"{src_label} → {dst_label}"


# def _assign_flow_visuals(edges: List[Dict[str, Any]], nodes: Dict[str, Dict[str, Any]]) -> None:
#     used_by_label: Dict[str, str] = {}
#     palette_idx = 0
#     for edge in edges:
#         flow_label = edge.get("label") or _edge_default_label(edge, nodes)
#         edge["display_label"] = flow_label
#         raw_color = _safe_color(edge.get("raw_color"), "") if edge.get("raw_color") else ""
#         if raw_color:
#             color = raw_color
#         elif flow_label in used_by_label:
#             color = used_by_label[flow_label]
#         else:
#             color = FLOW_PALETTE[palette_idx % len(FLOW_PALETTE)]
#             used_by_label[flow_label] = color
#             palette_idx += 1
#         edge["color"] = color
#         edge["dash_style"] = _dash_style_from_raw_style(edge.get("raw_style"))
#         edge["width_pt"] = _safe_float(edge.get("raw_penwidth"), DEFAULT_EDGE_WIDTH_PT)


# def _legend_from_edges(edges: List[Dict[str, Any]]) -> List[Tuple[str, str, Optional[MSO_LINE_DASH_STYLE]]]:
#     legend: List[Tuple[str, str, Optional[MSO_LINE_DASH_STYLE]]] = []
#     seen = set()
#     for edge in edges:
#         label = edge.get("display_label") or edge.get("label")
#         color = edge.get("color") or DEFAULT_EDGE_COLOR
#         dash_style = edge.get("dash_style")
#         key = (color, label, str(dash_style))
#         if label and key not in seen:
#             seen.add(key)
#             legend.append((color, label, dash_style))
#     return legend


# def _build_editable_model(diagram_text: str) -> Dict[str, Any]:
#     normalized = normalize_to_graphviz_dot(diagram_text)
#     if not normalized:
#         raise ValueError("Diagram text could not be normalized to DOT.")
#     dot = _strip_dot_comments(html.unescape(normalized))
#     clusters = _parse_clusters_from_dot(dot)
#     node_attrs, node_order = _parse_node_attrs_from_dot(dot)
#     edge_attrs = _parse_edge_attrs_from_dot(dot)
#     nodes: Dict[str, Dict[str, Any]] = {}
#     for node_id in node_order:
#         attrs = node_attrs.get(node_id, {})
#         label = _label(attrs.get("label"), node_id)
#         if _is_pseudo_node(node_id, label):
#             continue
#         nodes[node_id] = {
#             "id": node_id,
#             "label": label,
#             "attrs": attrs,
#             "icon_hint": _guess_icon_hint(node_id, label, attrs),
#             "zone": _infer_zone(node_id, label),
#         }
#     edges: List[Dict[str, Any]] = []
#     for (src, dst), attrs in edge_attrs.items():
#         for node_id in [src, dst]:
#             if node_id not in nodes:
#                 label = _label(None, node_id)
#                 if not _is_pseudo_node(node_id, label):
#                     nodes[node_id] = {
#                         "id": node_id,
#                         "label": label,
#                         "attrs": {},
#                         "icon_hint": _guess_icon_hint(node_id, label, {}),
#                         "zone": _infer_zone(node_id, label),
#                     }
#                     node_order.append(node_id)
#         if src in nodes and dst in nodes:
#             edges.append({
#                 "source": src,
#                 "target": dst,
#                 "raw_color": attrs.get("color") or attrs.get("fontcolor"),
#                 "raw_style": attrs.get("style"),
#                 "raw_penwidth": attrs.get("penwidth"),
#                 "label": _clean_label(attrs.get("label", "")),
#                 "attrs": dict(attrs),
#             })
#     if not edges and len(node_order) > 1:
#         for src, dst in zip(node_order[:-1], node_order[1:]):
#             if src in nodes and dst in nodes:
#                 edges.append({
#                     "source": src,
#                     "target": dst,
#                     "raw_color": None,
#                     "raw_style": None,
#                     "raw_penwidth": None,
#                     "label": "",
#                     "attrs": {},
#                 })
#     for cluster in clusters:
#         zone = _zone_from_cluster_label(cluster.get("label", ""), default="gcp")
#         for nid in cluster.get("nodes", []):
#             if nid in nodes:
#                 nodes[nid]["zone"] = zone
#     _assign_flow_visuals(edges, nodes)
#     legend = _legend_from_edges(edges)
#     ordered_nodes = [nodes[nid] for nid in node_order if nid in nodes]
#     return {
#         "nodes": ordered_nodes,
#         "node_map": nodes,
#         "node_order": [n["id"] for n in ordered_nodes],
#         "edges": edges,
#         "legend": legend,
#     }


# # -----------------------------------------------------------------------------
# # Swimlane layout - rank aligned, including Process View
# # -----------------------------------------------------------------------------
# def _zone_order_for_nodes(nodes: List[Dict[str, Any]]) -> List[str]:
#     present = {n.get("zone", "gcp") for n in nodes}
#     order = []
#     for zone in ["onprem", "gcp", "external", "cicd", "observability"]:
#         if zone in present:
#             order.append(zone)
#     return order or ["gcp"]


# def _zone_label(zone: str) -> str:
#     return {
#         "onprem": "Source / On-Premise",
#         "gcp": "Google Cloud Platform",
#         "external": "External",
#         "cicd": "CI/CD & Deployment",
#         "observability": "Observability",
#     }.get(zone, zone.title())


# def _lane_box(zone: str, label: str, left: float, top: float, width: float, height: float) -> Dict[str, Any]:
#     style = {
#         "onprem": ("Source / On-Premise", "EEF4FF", "8AB4F8"),
#         "gcp": ("Google Cloud Platform", "F3F8FF", "A8C7FA"),
#         "external": ("External", EXTERNAL_FILL, EXTERNAL_BORDER),
#         "cicd": ("CI/CD & Deployment", CICD_FILL, CICD_BORDER),
#         "observability": ("Observability", OBS_FILL, OBS_BORDER),
#     }.get(zone, (label, LANE_FILL, LANE_BORDER))
#     return {
#         "id": zone,
#         "label": style[0],
#         "left": left,
#         "top": top,
#         "width": width,
#         "height": height,
#         "fill": style[1],
#         "border": style[2],
#     }


# def _compute_lane_boxes(zone_order: List[str]) -> Dict[str, Dict[str, Any]]:
#     boxes: Dict[str, Dict[str, Any]] = {}
#     has_onprem = "onprem" in zone_order
#     has_gcp = "gcp" in zone_order
#     right_zones = [z for z in zone_order if z not in {"onprem", "gcp"}]
#     top = CANVAS_TOP_IN
#     height = CANVAS_HEIGHT_IN
#     if has_onprem and has_gcp:
#         onprem_w = 2.45
#         right_w = 2.05 if right_zones else 0.0
#         gcp_w = CANVAS_WIDTH_IN - onprem_w - right_w - (0.18 if right_w else 0.0)
#         boxes["onprem"] = _lane_box("onprem", "On-Premise", CANVAS_LEFT_IN, top, onprem_w, height)
#         boxes["gcp"] = _lane_box("gcp", "Google Cloud Platform", CANVAS_LEFT_IN + onprem_w + 0.18, top, gcp_w, height)
#         if right_zones:
#             rz_left = CANVAS_LEFT_IN + onprem_w + 0.18 + gcp_w + 0.18
#             rz_h = height / len(right_zones)
#             for idx, zone in enumerate(right_zones):
#                 boxes[zone] = _lane_box(zone, _zone_label(zone), rz_left, top + idx * rz_h, right_w, rz_h - 0.08)
#     else:
#         n = max(1, len(zone_order))
#         lane_w = (CANVAS_WIDTH_IN - (n - 1) * 0.18) / n
#         for idx, zone in enumerate(zone_order):
#             boxes[zone] = _lane_box(zone, _zone_label(zone), CANVAS_LEFT_IN + idx * (lane_w + 0.18), top, lane_w, height)
#     return boxes


# def _topological_order(node_order: List[str], edges: List[Dict[str, Any]]) -> List[str]:
#     indeg = {nid: 0 for nid in node_order}
#     outgoing: Dict[str, List[str]] = defaultdict(list)
#     for edge in edges:
#         src = edge["source"]
#         dst = edge["target"]
#         if src in indeg and dst in indeg:
#             outgoing[src].append(dst)
#             indeg[dst] += 1
#     q = deque([nid for nid in node_order if indeg.get(nid, 0) == 0])
#     result = []
#     while q:
#         cur = q.popleft()
#         result.append(cur)
#         for nxt in outgoing.get(cur, []):
#             indeg[nxt] -= 1
#             if indeg[nxt] == 0:
#                 q.append(nxt)
#     for nid in node_order:
#         if nid not in result:
#             result.append(nid)
#     return result


# def _rank_nodes_by_flow(node_order: List[str], edges: List[Dict[str, Any]]) -> Dict[str, int]:
#     rank = {nid: 0 for nid in node_order}
#     order_index = {nid: idx for idx, nid in enumerate(node_order)}
#     ordered_edges: List[Tuple[str, str]] = []
#     for edge in edges:
#         src = edge.get("source")
#         dst = edge.get("target")
#         if src not in rank or dst not in rank:
#             continue
#         if order_index.get(src, 0) >= order_index.get(dst, 0):
#             continue
#         ordered_edges.append((src, dst))
#     for _ in range(max(1, len(node_order))):
#         changed = False
#         for src, dst in ordered_edges:
#             candidate = rank[src] + 1
#             if candidate > rank[dst]:
#                 rank[dst] = candidate
#                 changed = True
#         if not changed:
#             break
#     return rank


# def _compact_rank_values(ids: List[str], rank_by_node: Dict[str, int], max_cols: int) -> Dict[str, int]:
#     if not ids:
#         return {}
#     unique_ranks = sorted({rank_by_node.get(nid, 0) for nid in ids})
#     if not unique_ranks:
#         return {nid: 0 for nid in ids}
#     rank_to_idx = {rank: idx for idx, rank in enumerate(unique_ranks)}
#     raw_cols = max(1, len(unique_ranks))
#     cols = max(1, min(max_cols, raw_cols))
#     compact: Dict[str, int] = {}
#     for nid in ids:
#         raw_idx = rank_to_idx.get(rank_by_node.get(nid, 0), 0)
#         if raw_cols <= cols:
#             compact[nid] = raw_idx
#         else:
#             compact[nid] = min(cols - 1, round(raw_idx * (cols - 1) / max(1, raw_cols - 1)))
#     return compact


# def _node_degrees(edges: List[Dict[str, Any]]) -> Dict[str, int]:
#     degrees: Dict[str, int] = defaultdict(int)
#     for edge in edges:
#         src = edge.get("source")
#         dst = edge.get("target")
#         if src:
#             degrees[src] += 1
#         if dst:
#             degrees[dst] += 1
#     return degrees


# def _is_mostly_linear_sequence(ids: List[str], edges: List[Dict[str, Any]]) -> bool:
#     if len(ids) < 4:
#         return False
#     id_set = set(ids)
#     in_deg = {nid: 0 for nid in ids}
#     out_deg = {nid: 0 for nid in ids}
#     edge_count = 0
#     for edge in edges:
#         src = edge.get("source")
#         dst = edge.get("target")
#         if src in id_set and dst in id_set:
#             out_deg[src] += 1
#             in_deg[dst] += 1
#             edge_count += 1
#     if edge_count < len(ids) - 1:
#         return False
#     branch_nodes = sum(1 for nid in ids if in_deg[nid] > 1 or out_deg[nid] > 1)
#     return branch_nodes <= max(1, len(ids) // 5)


# def _compute_column_capacity(inner_w: float, preferred_w: float = NODE_W, min_w: float = MIN_RANK_NODE_W) -> int:
#     preferred_cols = int((inner_w + H_GAP) // max(0.75, preferred_w + H_GAP))
#     compact_cols = int((inner_w + RANK_COL_GAP) // max(0.65, min_w + RANK_COL_GAP))
#     return max(1, max(preferred_cols, compact_cols))


# def _connected_weight(nid: str, edges: List[Dict[str, Any]]) -> int:
#     weight = 0
#     for edge in edges:
#         if edge.get("source") == nid or edge.get("target") == nid:
#             weight += 1
#     return weight


# def _layout_model(model: Dict[str, Any]) -> Dict[str, Any]:
#     nodes = model["nodes"]
#     edges = model["edges"]
#     original_order = model.get("node_order", [n["id"] for n in nodes])
#     order = _topological_order(original_order, edges)
#     by_id = {n["id"]: n for n in nodes}
#     rank_by_node = _rank_nodes_by_flow(order, edges)
#     degree_by_node = _node_degrees(edges)
#     zone_order = _zone_order_for_nodes(nodes)
#     lane_boxes = _compute_lane_boxes(zone_order)
#     zone_nodes: Dict[str, List[str]] = defaultdict(list)
#     for nid in order:
#         if nid in by_id:
#             zone_nodes[by_id[nid].get("zone", "gcp")].append(nid)
#     placements: Dict[str, Dict[str, float]] = {}
#     order_index = {nid: idx for idx, nid in enumerate(order)}
#     for zone in zone_order:
#         ids = zone_nodes.get(zone, [])
#         if not ids:
#             continue
#         lane = lane_boxes[zone]
#         inner_left = lane["left"] + LANE_PAD
#         inner_top = lane["top"] + LANE_TITLE_H + LANE_PAD
#         inner_w = max(0.5, lane["width"] - 2 * LANE_PAD)
#         inner_h = max(0.5, lane["height"] - LANE_TITLE_H - 2 * LANE_PAD)
#         connected_ids = [nid for nid in ids if degree_by_node.get(nid, 0) > 0]
#         isolated_ids = [nid for nid in ids if degree_by_node.get(nid, 0) == 0]
#         isolated_row_h = 0.0
#         if isolated_ids and connected_ids:
#             isolated_row_h = min(0.78, max(0.54, NODE_H * ISOLATED_NODE_SCALE)) + ISOLATED_ROW_GAP
#         flow_top = inner_top
#         flow_h = max(0.5, inner_h - isolated_row_h)
#         flow_ids = connected_ids if connected_ids else ids
#         if not flow_ids:
#             continue
#         max_cols_by_width = _compute_column_capacity(inner_w)
#         mostly_linear = _is_mostly_linear_sequence(flow_ids, edges)
#         if mostly_linear and len(flow_ids) <= max_cols_by_width:
#             rank_col_by_node = {nid: idx for idx, nid in enumerate(flow_ids)}
#         else:
#             if zone in {"onprem", "external", "cicd", "observability"} and not mostly_linear:
#                 max_cols = min(2, max_cols_by_width)
#             else:
#                 max_cols = min(max_cols_by_width, max(2, len(flow_ids)))
#             rank_col_by_node = _compact_rank_values(flow_ids, rank_by_node, max_cols)
#         cols = max(1, max(rank_col_by_node.values(), default=0) + 1)
#         col_nodes: Dict[int, List[str]] = defaultdict(list)
#         for nid in flow_ids:
#             col_nodes[rank_col_by_node.get(nid, 0)].append(nid)
#         max_rows = max((len(col_nodes.get(col, [])) for col in range(cols)), default=1)
#         effective_gap = RANK_COL_GAP if mostly_linear and cols >= 6 else H_GAP
#         node_w = min(NODE_W, (inner_w - (cols - 1) * effective_gap) / cols)
#         node_h = min(NODE_H, (flow_h - (max_rows - 1) * V_GAP) / max_rows)
#         node_w = max(MIN_RANK_NODE_W, node_w)
#         node_h = max(0.54, node_h)
#         total_w = cols * node_w + (cols - 1) * effective_gap
#         start_x = inner_left + max(0.0, (inner_w - total_w) / 2)
#         for col in range(cols):
#             col_ids = col_nodes.get(col, [])
#             if not col_ids:
#                 continue
#             col_ids.sort(key=lambda nid: (-_connected_weight(nid, edges), order_index.get(nid, 9999)))
#             rows = len(col_ids)
#             total_h = rows * node_h + (rows - 1) * V_GAP
#             start_y = flow_top + max(0.0, (flow_h - total_h) / 2)
#             for row, nid in enumerate(col_ids):
#                 placements[nid] = {
#                     "left": start_x + col * (node_w + effective_gap),
#                     "top": start_y + row * (node_h + V_GAP),
#                     "width": node_w,
#                     "height": node_h,
#                 }
#         if isolated_ids and connected_ids:
#             iso_h = min(0.64, max(0.48, node_h * ISOLATED_NODE_SCALE))
#             iso_w = min(1.35, max(1.05, node_w * ISOLATED_NODE_SCALE))
#             iso_gap = 0.16
#             iso_cols = max(1, min(len(isolated_ids), int((inner_w + iso_gap) // (iso_w + iso_gap))))
#             iso_total_w = iso_cols * iso_w + (iso_cols - 1) * iso_gap
#             iso_start_x = inner_left + max(0.0, (inner_w - iso_total_w) / 2)
#             iso_y = inner_top + inner_h - iso_h
#             isolated_ids.sort(key=lambda nid: order_index.get(nid, 9999))
#             for idx, nid in enumerate(isolated_ids[:iso_cols]):
#                 placements[nid] = {
#                     "left": iso_start_x + idx * (iso_w + iso_gap),
#                     "top": iso_y,
#                     "width": iso_w,
#                     "height": iso_h,
#                 }
#     return {"placements": placements, "lane_boxes": list(lane_boxes.values())}


# # -----------------------------------------------------------------------------
# # Generic edge routing - no service-specific hardcoding
# # -----------------------------------------------------------------------------
# def _spread_offset(index: int, total: int, step: float) -> float:
#     if total <= 1:
#         return 0.0
#     center = (total - 1) / 2.0
#     return (index - center) * step


# def _alternating_offset(index: int, step: float) -> float:
#     if index <= 0:
#         return 0.0
#     n = (index + 1) // 2
#     return n * step if index % 2 == 1 else -n * step


# def _box_center(box: Dict[str, float]) -> Tuple[float, float]:
#     return (box["left"] + box["width"] / 2.0, box["top"] + box["height"] / 2.0)


# def _choose_edge_sides(src: Dict[str, float], dst: Dict[str, float]) -> Tuple[str, str]:
#     src_x, src_y = _box_center(src)
#     dst_x, dst_y = _box_center(dst)
#     x_gap = dst_x - src_x
#     y_gap = dst_y - src_y
#     if abs(x_gap) >= abs(y_gap) * 1.10:
#         return ("right", "left") if x_gap >= 0 else ("left", "right")
#     return ("bottom", "top") if y_gap >= 0 else ("top", "bottom")


# def _port_point(box: Dict[str, float], side: str, offset: float = 0.0) -> Tuple[float, float]:
#     cx = box["left"] + box["width"] / 2.0
#     cy = box["top"] + box["height"] / 2.0
#     if side == "left":
#         return (box["left"] - EDGE_MARGIN_FROM_NODE_IN, cy + offset)
#     if side == "right":
#         return (box["left"] + box["width"] + EDGE_MARGIN_FROM_NODE_IN, cy + offset)
#     if side == "top":
#         return (cx + offset, box["top"] - EDGE_MARGIN_FROM_NODE_IN)
#     return (cx + offset, box["top"] + box["height"] + EDGE_MARGIN_FROM_NODE_IN)


# def _route_edge_points(
#     src_box: Dict[str, float],
#     dst_box: Dict[str, float],
#     src_side: str,
#     dst_side: str,
#     src_offset: float,
#     dst_offset: float,
#     channel_rank: int = 0,
#     pair_rank: int = 0,
# ) -> List[Tuple[float, float]]:
#     start = _port_point(src_box, src_side, src_offset)
#     end = _port_point(dst_box, dst_side, dst_offset)
#     sx, sy = start
#     ex, ey = end
#     total_offset = _alternating_offset(channel_rank, EDGE_CHANNEL_SPACING_IN) + _alternating_offset(pair_rank, EDGE_PAIR_SEPARATION_IN)
#     if src_side in {"left", "right"} and dst_side in {"left", "right"} and abs(sy - ey) <= SAME_ROW_TOLERANCE_IN:
#         if abs(total_offset) <= 0.001:
#             return [start, end]
#         route_y = sy + total_offset
#         return [
#             start,
#             (sx + (EDGE_MIN_CHANNEL_CLEARANCE_IN if src_side == "right" else -EDGE_MIN_CHANNEL_CLEARANCE_IN), route_y),
#             (ex + (-EDGE_MIN_CHANNEL_CLEARANCE_IN if dst_side == "left" else EDGE_MIN_CHANNEL_CLEARANCE_IN), route_y),
#             end,
#         ]
#     if src_side in {"left", "right"} and dst_side in {"left", "right"}:
#         base_x = (sx + ex) / 2.0
#         channel_x = base_x + total_offset
#         if src_side == "right":
#             channel_x = max(channel_x, sx + EDGE_MIN_CHANNEL_CLEARANCE_IN)
#         else:
#             channel_x = min(channel_x, sx - EDGE_MIN_CHANNEL_CLEARANCE_IN)
#         if dst_side == "left":
#             channel_x = min(channel_x, ex - EDGE_MIN_CHANNEL_CLEARANCE_IN)
#         else:
#             channel_x = max(channel_x, ex + EDGE_MIN_CHANNEL_CLEARANCE_IN)
#         return [start, (channel_x, sy), (channel_x, ey), end]
#     if src_side in {"top", "bottom"} and dst_side in {"top", "bottom"}:
#         base_y = (sy + ey) / 2.0
#         channel_y = base_y + total_offset
#         if src_side == "bottom":
#             channel_y = max(channel_y, sy + EDGE_MIN_CHANNEL_CLEARANCE_IN)
#         else:
#             channel_y = min(channel_y, sy - EDGE_MIN_CHANNEL_CLEARANCE_IN)
#         if dst_side == "top":
#             channel_y = min(channel_y, ey - EDGE_MIN_CHANNEL_CLEARANCE_IN)
#         else:
#             channel_y = max(channel_y, ey + EDGE_MIN_CHANNEL_CLEARANCE_IN)
#         return [start, (sx, channel_y), (ex, channel_y), end]
#     if src_side in {"left", "right"}:
#         direction = 1 if src_side == "right" else -1
#         mid_x = sx + direction * (EDGE_MIN_CHANNEL_CLEARANCE_IN + abs(total_offset))
#         return [start, (mid_x, sy), (mid_x, ey), end]
#     direction = 1 if src_side == "bottom" else -1
#     mid_y = sy + direction * (EDGE_MIN_CHANNEL_CLEARANCE_IN + abs(total_offset))
#     return [start, (sx, mid_y), (ex, mid_y), end]


# def _build_edge_routes(model: Dict[str, Any], placements: Dict[str, Dict[str, float]]) -> Dict[int, List[Tuple[float, float]]]:
#     edges = model.get("edges", [])
#     edge_meta: Dict[int, Dict[str, Any]] = {}
#     outgoing_groups: Dict[Tuple[str, str], List[int]] = defaultdict(list)
#     incoming_groups: Dict[Tuple[str, str], List[int]] = defaultdict(list)
#     pair_groups: Dict[Tuple[str, str], List[int]] = defaultdict(list)
#     for idx, edge in enumerate(edges):
#         src_box = placements.get(edge["source"])
#         dst_box = placements.get(edge["target"])
#         if not src_box or not dst_box:
#             continue
#         src_side, dst_side = _choose_edge_sides(src_box, dst_box)
#         edge_meta[idx] = {"src_side": src_side, "dst_side": dst_side, "source": edge["source"], "target": edge["target"]}
#         outgoing_groups[(edge["source"], src_side)].append(idx)
#         incoming_groups[(edge["target"], dst_side)].append(idx)
#         pair_groups[(edge["source"], edge["target"])].append(idx)
#     routes: Dict[int, List[Tuple[float, float]]] = {}
#     for idx, edge in enumerate(edges):
#         meta = edge_meta.get(idx)
#         if not meta:
#             continue
#         src_box = placements.get(edge["source"])
#         dst_box = placements.get(edge["target"])
#         if not src_box or not dst_box:
#             continue
#         src_side = meta["src_side"]
#         dst_side = meta["dst_side"]
#         out_group = outgoing_groups.get((edge["source"], src_side), [])
#         in_group = incoming_groups.get((edge["target"], dst_side), [])
#         pair_group = pair_groups.get((edge["source"], edge["target"]), [])
#         src_rank = out_group.index(idx) if idx in out_group else 0
#         dst_rank = in_group.index(idx) if idx in in_group else 0
#         pair_rank = pair_group.index(idx) if idx in pair_group else 0
#         src_offset = _spread_offset(src_rank, len(out_group), EDGE_PORT_SPREAD_IN)
#         dst_offset = _spread_offset(dst_rank, len(in_group), EDGE_PORT_SPREAD_IN)
#         routes[idx] = _route_edge_points(
#             src_box=src_box,
#             dst_box=dst_box,
#             src_side=src_side,
#             dst_side=dst_side,
#             src_offset=src_offset,
#             dst_offset=dst_offset,
#             channel_rank=src_rank,
#             pair_rank=pair_rank,
#         )
#     return routes


# # -----------------------------------------------------------------------------
# # Drawing helpers
# # -----------------------------------------------------------------------------
# def _set_text_frame(
#     shape: BaseShape,
#     text: str,
#     font_size_pt: float = 10,
#     bold: bool = False,
#     color_hex: str = DEFAULT_NODE_FONT,
#     align=PP_ALIGN.LEFT,
# ) -> None:
#     tf = shape.text_frame
#     tf.clear()
#     tf.word_wrap = True
#     tf.margin_left = Inches(0.03)
#     tf.margin_right = Inches(0.03)
#     tf.margin_top = Inches(0.01)
#     tf.margin_bottom = Inches(0.01)
#     tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
#     p = tf.paragraphs[0]
#     p.alignment = align
#     run = p.add_run()
#     run.text = _safe_title(text)
#     run.font.name = "Calibri"
#     run.font.size = Pt(font_size_pt)
#     run.font.bold = bold
#     run.font.color.rgb = _rgb(color_hex, DEFAULT_NODE_FONT)


# def _font_for_label(label: str, width: float, height: float) -> float:
#     clean = _safe_title(label)
#     line_count = max(1, clean.count("\n") + 1)
#     length = len(clean.replace("\n", " "))
#     if width < 1.15 or height < 0.58 or length > 58 or line_count >= 3:
#         return 6.8
#     if width < 1.5 or length > 42 or line_count == 2:
#         return 7.8
#     return 8.8


# def _add_title(slide, title_text: str) -> None:
#     shape = slide.shapes.add_textbox(Inches(TITLE_LEFT_IN), Inches(TITLE_TOP_IN), Inches(TITLE_WIDTH_IN), Inches(TITLE_HEIGHT_IN))
#     _set_text_frame(shape, title_text, font_size_pt=18, bold=True, color_hex="202124")


# def _add_lane_box(slide, box: Dict[str, Any]) -> None:
#     shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
#     shape.fill.solid()
#     shape.fill.fore_color.rgb = _rgb(box.get("fill", LANE_FILL))
#     shape.line.color.rgb = _rgb(box.get("border", LANE_BORDER))
#     shape.line.width = Pt(1.2)
#     try:
#         shape.adjustments[0] = 0.04
#     except Exception:
#         pass
#     label_box = slide.shapes.add_textbox(Inches(box["left"] + 0.08), Inches(box["top"] + 0.03), Inches(max(0.5, box["width"] - 0.16)), Inches(0.20))
#     _set_text_frame(label_box, box.get("label", ""), font_size_pt=8.7, bold=True, color_hex=LANE_TITLE_COLOR)


# def _add_node_box(slide, box: Dict[str, float], label: str, icon_path: Optional[Path] = None) -> None:
#     shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
#     shape.fill.solid()
#     shape.fill.fore_color.rgb = _rgb(DEFAULT_NODE_FILL)
#     shape.line.color.rgb = _rgb(DEFAULT_NODE_BORDER)
#     shape.line.width = Pt(1.0)
#     try:
#         shape.adjustments[0] = 0.08
#     except Exception:
#         pass
#     icon_size = min(NODE_ICON_SIZE_IN, max(0.15, box["height"] * 0.36))
#     text_left = box["left"] + NODE_PADDING_IN
#     if icon_path and icon_path.exists() and box["width"] > 0.75:
#         try:
#             slide.shapes.add_picture(str(icon_path), Inches(box["left"] + NODE_PADDING_IN), Inches(box["top"] + (box["height"] - icon_size) / 2.0), width=Inches(icon_size), height=Inches(icon_size))
#             text_left = box["left"] + NODE_TEXT_LEFT_WITH_ICON_IN
#         except Exception:
#             logger.debug("Could not add icon to PPTX node: %s", icon_path, exc_info=True)
#     text_width = max(0.35, box["width"] - (text_left - box["left"]) - 0.05)
#     text_box = slide.shapes.add_textbox(Inches(text_left), Inches(box["top"] + 0.03), Inches(text_width), Inches(max(0.22, box["height"] - 0.06)))
#     _set_text_frame(text_box, label, font_size_pt=_font_for_label(label, box["width"], box["height"]), bold=True, color_hex=DEFAULT_NODE_FONT)


# def _shorten_last_segment_for_arrowhead(points: List[Tuple[float, float]], gap: float = ARROW_TIP_GAP_IN) -> List[Tuple[float, float]]:
#     if len(points) < 2:
#         return points
#     adjusted = list(points)
#     x1, y1 = adjusted[-2]
#     x2, y2 = adjusted[-1]
#     dx = x2 - x1
#     dy = y2 - y1
#     length = math.sqrt(dx * dx + dy * dy)
#     if length <= 0.001:
#         return adjusted
#     usable_gap = min(gap, length * 0.35)
#     adjusted[-1] = (x2 - (dx / length) * usable_gap, y2 - (dy / length) * usable_gap)
#     return adjusted


# def _set_line_end_arrow(connector: BaseShape) -> None:
#     try:
#         ln = connector._element.spPr.get_or_add_ln()
#         for child in list(ln):
#             if child.tag.endswith("}tailEnd"):
#                 ln.remove(child)
#         tail_end = OxmlElement("a:tailEnd")
#         tail_end.set("type", ARROW_HEAD_TYPE)
#         tail_end.set("w", ARROW_HEAD_WIDTH)
#         tail_end.set("len", ARROW_HEAD_LENGTH)
#         ln.append(tail_end)
#     except Exception:
#         logger.debug("Could not apply arrowhead to connector.", exc_info=True)


# def _add_line_segment(
#     slide,
#     p1: Tuple[float, float],
#     p2: Tuple[float, float],
#     color_hex: str,
#     add_arrowhead: bool = False,
#     dash_style: Optional[MSO_LINE_DASH_STYLE] = None,
#     width_pt: float = DEFAULT_EDGE_WIDTH_PT,
# ):
#     connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p1[0]), Inches(p1[1]), Inches(p2[0]), Inches(p2[1]))
#     connector.line.color.rgb = _rgb(color_hex, DEFAULT_EDGE_COLOR)
#     connector.line.width = Pt(width_pt)
#     if dash_style is not None:
#         try:
#             connector.line.dash_style = dash_style
#         except Exception:
#             logger.debug("Could not apply line dash style.", exc_info=True)
#     if add_arrowhead:
#         _set_line_end_arrow(connector)
#     return connector


# def _add_polyline_arrow(
#     slide,
#     points: List[Tuple[float, float]],
#     color_hex: str,
#     dash_style: Optional[MSO_LINE_DASH_STYLE] = None,
#     width_pt: float = DEFAULT_EDGE_WIDTH_PT,
# ) -> None:
#     if len(points) < 2:
#         return
#     clean = [points[0]]
#     for point in points[1:]:
#         if abs(point[0] - clean[-1][0]) > 0.01 or abs(point[1] - clean[-1][1]) > 0.01:
#             clean.append(point)
#     if len(clean) < 2:
#         return
#     clean = _shorten_last_segment_for_arrowhead(clean)
#     simplified = [clean[0]]
#     for i in range(1, len(clean) - 1):
#         prev_pt = simplified[-1]
#         cur_pt = clean[i]
#         next_pt = clean[i + 1]
#         same_x = abs(prev_pt[0] - cur_pt[0]) < 0.01 and abs(cur_pt[0] - next_pt[0]) < 0.01
#         same_y = abs(prev_pt[1] - cur_pt[1]) < 0.01 and abs(cur_pt[1] - next_pt[1]) < 0.01
#         if same_x or same_y:
#             continue
#         simplified.append(cur_pt)
#     simplified.append(clean[-1])
#     for idx in range(len(simplified) - 1):
#         is_last_segment = idx == len(simplified) - 2
#         _add_line_segment(slide, simplified[idx], simplified[idx + 1], color_hex, add_arrowhead=is_last_segment, dash_style=dash_style, width_pt=width_pt)


# def _add_legend(slide, legend_items: List[Tuple[str, str, Optional[MSO_LINE_DASH_STYLE]]]) -> None:
#     if not legend_items:
#         return
#     title_box = slide.shapes.add_textbox(Inches(LEGEND_LEFT_IN), Inches(LEGEND_TOP_IN), Inches(1.15), Inches(0.18))
#     _set_text_frame(title_box, "Flow Legend", font_size_pt=9.2, bold=True, color_hex="202124")
#     x = LEGEND_LEFT_IN
#     y = LEGEND_TOP_IN + 0.24
#     max_x = LEGEND_LEFT_IN + LEGEND_WIDTH_IN
#     for color, label, dash_style in legend_items[:18]:
#         label = _safe_title(label)
#         item_w = max(1.55, min(3.30, 0.62 + len(label) * 0.050))
#         if x + item_w > max_x:
#             x = LEGEND_LEFT_IN
#             y += 0.22
#         if y > LEGEND_TOP_IN + LEGEND_MAX_HEIGHT_IN:
#             break
#         chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(item_w), Inches(0.18))
#         chip.fill.solid()
#         chip.fill.fore_color.rgb = _rgb(LEGEND_FILL)
#         chip.line.color.rgb = _rgb(LEGEND_BORDER)
#         chip.line.width = Pt(0.6)
#         line_y = y + 0.09
#         _add_line_segment(slide, (x + 0.05, line_y), (x + 0.18, line_y), color, add_arrowhead=True, dash_style=dash_style, width_pt=1.35)
#         label_box = slide.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.01), Inches(max(0.3, item_w - 0.24)), Inches(0.15))
#         _set_text_frame(label_box, label, font_size_pt=7.2, bold=False, color_hex="202124")
#         x += item_w + 0.08


# def _add_fallback_preview_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str) -> None:
#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, f"{slide_title} (Preview)")
#     try:
#         assets = render_graphviz_to_assets(diagram_text, output_dir=output_dir, base_name=base_name)
#     except Exception:
#         logger.exception("Fallback preview asset generation failed.")
#         assets = None
#     if assets and assets.get("png_path") and Path(assets["png_path"]).exists():
#         slide.shapes.add_picture(str(assets["png_path"]), Inches(CANVAS_LEFT_IN), Inches(CANVAS_TOP_IN), width=Inches(PREVIEW_MAX_WIDTH_IN))
#         _add_legend(slide, assets.get("legend", []))


# # -----------------------------------------------------------------------------
# # Diagram extraction
# # -----------------------------------------------------------------------------
# def _is_diagram_text(value: Any) -> bool:
#     return isinstance(value, str) and normalize_to_graphviz_dot(_clean_internal_tokens(value)) is not None


# def _walk_for_diagrams(value: Any, path_parts: List[str], results: List[Dict[str, Any]]) -> None:
#     if value is None:
#         return
#     if isinstance(value, str):
#         if _is_diagram_text(value):
#             results.append({
#                 "title": " / ".join(path_parts) if path_parts else "Diagram",
#                 "diagram_text": _clean_internal_tokens(value),
#                 "section_key": _safe_name(path_parts[-1] if path_parts else "diagram"),
#             })
#         return
#     if isinstance(value, list):
#         for idx, item in enumerate(value, start=1):
#             if _is_diagram_text(item):
#                 results.append({
#                     "title": " / ".join(path_parts) if path_parts else f"Diagram {idx}",
#                     "diagram_text": _clean_internal_tokens(item),
#                     "section_key": _safe_name(path_parts[-1] if path_parts else f"diagram_{idx}"),
#                 })
#             elif isinstance(item, (dict, list)):
#                 _walk_for_diagrams(item, path_parts, results)
#         return
#     if isinstance(value, dict):
#         for key, child in value.items():
#             key_text = str(key).replace("_", " ").title()
#             next_path = path_parts if str(key).lower() == "diagrams" else path_parts + [key_text]
#             _walk_for_diagrams(child, next_path, results)


# def extract_diagram_specs(data: Dict[str, Any], hld_model: Any, selected_sections_set: Optional[Set[str]] = None) -> List[Dict[str, Any]]:
#     results: List[Dict[str, Any]] = []
#     for field_name, field_info in hld_model.model_fields.items():
#         if not should_include_section(field_name, field_info, selected_sections_set or set()):
#             continue
#         section_title = field_info.title or field_name.replace("_", " ").title()
#         _walk_for_diagrams(data.get(field_name), [section_title], results)
#     counts: Dict[str, int] = defaultdict(int)
#     for item in results:
#         counts[item["title"]] += 1
#     seen: Dict[str, int] = defaultdict(int)
#     for item in results:
#         if counts[item["title"]] > 1:
#             seen[item["title"]] += 1
#             item["title"] = f'{item["title"]} ({seen[item["title"]]})'
#     return results


# # -----------------------------------------------------------------------------
# # Slide rendering
# # -----------------------------------------------------------------------------
# def _edge_priority(edge: Dict[str, Any]) -> int:
#     return 0 if edge.get("dash_style") is None else 1


# def add_diagram_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str) -> None:
#     try:
#         model = _build_editable_model(diagram_text)
#         if not model["nodes"]:
#             raise ValueError("No editable nodes parsed")
#         layout = _layout_model(model)
#         placements = layout["placements"]
#         lane_boxes = layout["lane_boxes"]
#     except Exception:
#         logger.exception("Editable PPTX rendering failed for slide '%s'. Falling back to preview image.", slide_title)
#         _add_fallback_preview_slide(prs, slide_title, diagram_text, output_dir, base_name)
#         return
#     slide = prs.slides.add_slide(prs.slide_layouts[6])
#     _add_title(slide, slide_title)
#     for box in lane_boxes:
#         _add_lane_box(slide, box)
#     edge_routes = _build_edge_routes(model, placements)
#     indexed_edges = list(enumerate(model["edges"]))
#     indexed_edges.sort(key=lambda item: _edge_priority(item[1]))
#     for idx, edge in indexed_edges:
#         points = edge_routes.get(idx)
#         if not points:
#             continue
#         _add_polyline_arrow(
#             slide,
#             points,
#             edge.get("color") or DEFAULT_EDGE_COLOR,
#             dash_style=edge.get("dash_style"),
#             width_pt=edge.get("width_pt", DEFAULT_EDGE_WIDTH_PT),
#         )
#     for node in model["nodes"]:
#         box = placements.get(node["id"])
#         if not box:
#             continue
#         icon_path = _preferred_picture_icon_path(node.get("icon_hint") or node.get("label") or node["id"])
#         _add_node_box(slide, box, node.get("label") or node["id"], icon_path)
#     _add_legend(slide, model.get("legend", []))


# # -----------------------------------------------------------------------------
# # Public builder
# # -----------------------------------------------------------------------------
# def build_editable_pptx(data: Dict[str, Any], pptx_path: str, hld_model: Any, selected_sections_set: Optional[Set[str]] = None) -> None:
#     output_path = Path(pptx_path)
#     output_path.parent.mkdir(parents=True, exist_ok=True)
#     diagram_assets_dir = output_path.parent / f"{output_path.stem}_diagram_assets"
#     diagram_assets_dir.mkdir(parents=True, exist_ok=True)
#     diagram_specs = extract_diagram_specs(data=data, hld_model=hld_model, selected_sections_set=selected_sections_set)
#     prs = Presentation()
#     prs.slide_width = Inches(SLIDE_WIDTH_IN)
#     prs.slide_height = Inches(SLIDE_HEIGHT_IN)
#     if not diagram_specs:
#         slide = prs.slides.add_slide(prs.slide_layouts[6])
#         _add_title(slide, "Editable Diagrams")
#         msg = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11.5), Inches(1.0))
#         _set_text_frame(msg, "No diagrams were detected in the selected sections.", font_size_pt=16, bold=True, color_hex="5F6368", align=PP_ALIGN.CENTER)
#         prs.save(str(output_path))
#         return
#     for counter, spec in enumerate(diagram_specs, start=1):
#         add_diagram_slide(prs, spec["title"], spec["diagram_text"], diagram_assets_dir, f"{_safe_name(spec['section_key'])}_{counter:02d}")
#     prs.save(str(output_path))
#     logger.info("Editable PPTX generated: %s", output_path)

#6th working code





from __future__ import annotations

import html
import json
import logging
import math
import re
import shutil
import subprocess
import tempfile
import urllib.request
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.shapes.base import BaseShape
from pptx.util import Inches, Pt

from .graphviz_renderer import normalize_to_graphviz_dot, render_graphviz_to_assets
from .schema_utils import should_include_section
from .text_sanitizer import renderer_safe_plain_text

logger = logging.getLogger(__name__)

@dataclass
class DiagramConfig:
    """Dynamic configuration for layout and styling without hardcoding constants."""
    slide_width_in: float = 13.333
    slide_height_in: float = 7.5
    title_left_in: float = 0.35
    title_top_in: float = 0.12
    title_width_in: float = 12.6
    title_height_in: float = 0.42
    
    canvas_left_in: float = 0.35
    canvas_top_in: float = 0.68
    canvas_width_in: float = 12.3
    canvas_height_in: float = 5.28
    
    legend_left_in: float = 0.35
    legend_top_in: float = 6.12
    legend_width_in: float = 12.3
    legend_max_height_in: float = 1.22

    node_icon_size_in: float = 0.20
    node_padding_in: float = 0.06
    node_text_left_with_icon_in: float = 0.30
    
    default_node_fill: str = "F8F9FA"
    default_node_border: str = "DADCE0"
    default_node_font: str = "202124"
    lane_fill: str = "F3F8FF"
    lane_border: str = "A8C7FA"
    lane_title_color: str = "174EA6"
    default_edge_color: str = "5F6368"
    default_edge_width_pt: float = 1.70
    legend_fill: str = "FFFFFF"
    legend_border: str = "DADCE0"
    preview_max_width_in: float = 12.0
    
    flow_palette: List[str] = field(default_factory=lambda: [
        "E60000", "1D70B8", "28A197", "F47738", "6F35A5", 
        "007C89", "D53880", "85994B", "B58840", "5F6368"
    ])
    color_name_map: Dict[str, str] = field(default_factory=lambda: {
        "red":"E60000", "blue":"1D70B8", "green":"28A197", "orange":"F47738", 
        "purple":"6F35A5", "teal":"007C89", "gray":"5F6368", "grey":"5F6368", 
        "black":"202124", "lightgrey":"DADCE0", "lightgray":"DADCE0", "white":"FFFFFF"
    })
    
    arrow_head_type: str = "triangle"
    arrow_head_width: str = "med"
    arrow_head_length: str = "med"
    arrow_tip_gap_in: float = 0.055

COMMENT_LINE_RE = re.compile(r"//.*?$", re.MULTILINE)
BLOCK_COMMENT_RE = re.compile(r"/\*.*?\*/", re.DOTALL)
EDGE_RE = re.compile(r'("[^"]+"|[A-Za-z0-9_.:-]+)\s*(->|--)\s*("[^"]+"|[A-Za-z0-9_.:-]+)')

def _clean_internal_tokens(text: Any) -> str:
    value = str(text or "")
    for token in ["AIASECTIONBLOCKSTARTTOKEN", "AIASECTIONBLOCKENDTOKEN", "[[AIASECTIONBLOCKSTARTTOKEN]]", "[[AIASECTIONBLOCKENDTOKEN]]", "SECTION_BLOCK_START", "SECTION_BLOCK_END", "[[SECTION_BLOCK_START]]", "[[SECTION_BLOCK_END]]", "[[SECTIONBLOCKSTART]]", "[[SECTIONBLOCKEND]]"]:
        value = value.replace(token, "")
    return html.unescape(value.strip())

def _safe_title(text: Any) -> str:
    return renderer_safe_plain_text(_clean_internal_tokens(text)).strip() or "Diagram"

def _safe_name(text: Any) -> str:
    return re.sub(r"[^a-zA-Z0-9_-]+", "_", str(text or "")).strip("_").lower() or "diagram"

def _clean_label(value: Any) -> str:
    text = html.unescape(str(value or "").strip()).strip('"')
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("\\\\N", "\n").replace("\\\\n", "\n")
    text = text.replace("\\N", "\n").replace("\\n", "\n").replace("&nbsp;", " ")
    text = re.sub(r"\s+\n", "\n", text)
    text = re.sub(r"\n\s+", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()

def _safe_color(value: Optional[str], config: DiagramConfig, default: str) -> str:
    raw = str(value or "").strip().strip('"').lstrip("#")
    if not raw:
        return default
    if raw.lower() in config.color_name_map:
        return config.color_name_map[raw.lower()]
    raw = raw.upper()
    return raw if len(raw) == 6 and all(ch in "0123456789ABCDEF" for ch in raw) else default

def _rgb(hex_color: str, default: str = "000000") -> RGBColor:
    return RGBColor.from_string(hex_color) if hex_color else RGBColor.from_string(default)

def _safe_float(value: Any, default: float) -> float:
    try:
        return float(str(value).strip())
    except Exception:
        return default

def _shell_split_graphviz(line: str) -> List[str]:
    import shlex
    return shlex.split(line, posix=True)

def _strip_dot_comments(dot_text: str) -> str:
    return COMMENT_LINE_RE.sub("", BLOCK_COMMENT_RE.sub("", dot_text or ""))

def _outer_graph_body(dot: str) -> str:
    text = dot or ""
    start = text.find("{")
    end = text.rfind("}")
    return text[start + 1:end] if start != -1 and end != -1 and end > start else text

def _split_statements(dot_text: str) -> List[str]:
    statements, buf = [], []
    bracket_depth = brace_depth = angle_depth = 0
    for ch in dot_text or "":
        if ch == "[": bracket_depth += 1
        elif ch == "]": bracket_depth = max(0, bracket_depth - 1)
        elif ch == "{": brace_depth += 1
        elif ch == "}": brace_depth = max(0, brace_depth - 1)
        elif ch == "<": angle_depth += 1
        elif ch == ">": angle_depth = max(0, angle_depth - 1)
        if ch == ";" and bracket_depth == 0 and brace_depth == 0 and angle_depth == 0:
            stmt = "".join(buf).strip()
            if stmt: statements.append(stmt)
            buf = []
        else:
            buf.append(ch)
    tail = "".join(buf).strip()
    if tail: statements.append(tail)
    return statements

def _extract_attr_pairs(attr_text: str) -> Dict[str, str]:
    attrs: Dict[str, str] = {}
    text = html.unescape(str(attr_text or ""))
    for match in re.finditer(r'([A-Za-z0-9_:-]+)\s*=\s*("[^"]*"|<[^>]*>|[^,\n\r;]+)', text):
        key = match.group(1).lower().strip()
        value = match.group(2).strip().strip('"')
        if value.startswith("<") and value.endswith(">"):
            value = value[1:-1]
        attrs[key] = value.strip()
    return attrs

def _extract_bracket_attr(statement: str) -> Tuple[str, Dict[str, str]]:
    statement = html.unescape((statement or "").strip().rstrip(";").strip())
    if "[" not in statement or "]" not in statement:
        return statement, {}
    start = statement.find("[")
    end = statement.rfind("]")
    if end <= start:
        return statement, {}
    return statement[:start].strip(), _extract_attr_pairs(statement[start + 1:end])

def _unquote_id(value: str) -> str:
    return str(value or "").strip().strip('"').strip()

def _is_graphviz_control_statement(statement: str) -> bool:
    s = html.unescape(str(statement or "")).strip().lower()
    if not s or s in {"{", "}"}:
        return True
    return s.startswith(("digraph ", "graph ", "subgraph ", "node ", "edge ", "rank ", "rankdir", "label", "style", "color", "fillcolor", "fontcolor", "fontsize", "fontname", "margin", "pad", "splines", "bgcolor"))

def _is_pseudo_node(node_id: str, label: Optional[str] = None) -> bool:
    text = _clean_label(label if label is not None else node_id).lower().strip()
    raw = str(node_id or "").lower().strip()
    if text in {"", "n/a", "na", "none", "null"}:
        return True
    if raw in {"node", "edge", "graph", "digraph", "subgraph", "rank", "label", "style", "color", "fillcolor", "fontcolor", "rankdir"}:
        return True
    if raw.startswith(("digraph", "graph", "subgraph", "cluster_", "cluster ", "label=", "style=", "color=", "fillcolor=", "fontcolor=", "rankdir=")):
        return True
    if text.startswith(("label=", "style=", "color=", "fillcolor=", "fontcolor=", "rankdir=", "fontsize=", "fontname=", "margin=", "pad=")):
        return True
    return False

def _find_subgraph_blocks(dot_text: str) -> List[Tuple[str, str]]:
    text = dot_text or ""
    results: List[Tuple[str, str]] = []
    idx = 0
    while True:
        match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
        if not match: break
        start = idx + match.start()
        brace_start = text.find("{", start)
        if brace_start == -1: break
        header = text[start:brace_start].strip()
        depth = 0
        end = brace_start
        while end < len(text):
            if text[end] == "{": depth += 1
            elif text[end] == "}":
                depth -= 1
                if depth == 0: break
            end += 1
        if end >= len(text): break
        results.append((header, text[brace_start + 1:end]))
        idx = end + 1
    return results

def _remove_subgraph_blocks(dot_text: str) -> str:
    text = dot_text or ""
    pieces: List[str] = []
    idx = 0
    while True:
        match = re.search(r"\bsubgraph\b", text[idx:], flags=re.IGNORECASE)
        if not match:
            pieces.append(text[idx:]); break
        start = idx + match.start()
        pieces.append(text[idx:start])
        brace_start = text.find("{", start)
        if brace_start == -1: break
        depth = 0
        end = brace_start
        while end < len(text):
            if text[end] == "{": depth += 1
            elif text[end] == "}":
                depth -= 1
                if depth == 0: break
            end += 1
        if end >= len(text): break
        idx = end + 1
    return "".join(pieces)

def _extract_cluster_label(body: str, fallback_name: str) -> str:
    for stmt in _split_statements(body):
        s = html.unescape(stmt or "").strip()
        if s.lower().startswith("label") and "=" in s:
            value = s.split("=", 1)[1].strip().strip('"').strip()
            if value: return _clean_label(value)
    return fallback_name

def _extract_cluster_nodes(body: str) -> Set[str]:
    node_ids: Set[str] = set()
    for stmt in _split_statements(body):
        s = html.unescape(stmt or "").strip()
        if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s): continue
        prefix, attrs = _extract_bracket_attr(s)
        node_id = _unquote_id(prefix)
        label = attrs.get("label", node_id)
        if node_id and not _is_pseudo_node(node_id, label) and re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
            node_ids.add(node_id)
    for match in EDGE_RE.finditer(body or ""):
        for group_idx in [1, 3]:
            node_id = _unquote_id(match.group(group_idx))
            if node_id and not _is_pseudo_node(node_id): node_ids.add(node_id)
    return node_ids

def _parse_clusters_from_dot(dot: str) -> List[Dict[str, Any]]:
    clusters: List[Dict[str, Any]] = []
    for header, body in _find_subgraph_blocks(dot):
        fallback_name = _unquote_id(header.replace("subgraph", "", 1).strip()) or "Cluster"
        label = _extract_cluster_label(body, fallback_name)
        nodes = sorted(_extract_cluster_nodes(body))
        if nodes and not _is_pseudo_node(fallback_name, label):
            clusters.append({"id": _safe_name(fallback_name), "label": label, "nodes": nodes})
    return clusters

def _parse_node_attrs_from_dot(dot: str) -> Tuple[Dict[str, Dict[str, str]], List[str]]:
    attrs_by_node: Dict[str, Dict[str, str]] = {}
    node_order: List[str] = []
    def parse_node_statements(text: str) -> None:
        for stmt in _split_statements(text):
            s = html.unescape(stmt or "").strip()
            if not s or "->" in s or "--" in s or _is_graphviz_control_statement(s): continue
            prefix, attrs = _extract_bracket_attr(s)
            node_id = _unquote_id(prefix)
            label = attrs.get("label", node_id)
            if node_id and not _is_pseudo_node(node_id, label) and re.match(r'^[A-Za-z0-9_".:-]+$', prefix.strip()):
                if node_id not in attrs_by_node: node_order.append(node_id)
                attrs_by_node[node_id] = attrs
    body = _outer_graph_body(dot)
    parse_node_statements(_remove_subgraph_blocks(body))
    for _header, subgraph_body in _find_subgraph_blocks(dot): parse_node_statements(subgraph_body)
    return attrs_by_node, node_order

def _parse_edge_attrs_from_dot(dot: str) -> Dict[Tuple[str, str], Dict[str, str]]:
    attrs_by_edge: Dict[Tuple[str, str], Dict[str, str]] = {}
    def parse_edge_statements(text: str) -> None:
        for stmt in _split_statements(text):
            s = html.unescape(stmt or "").strip()
            if "->" not in s and "--" not in s: continue
            prefix, attrs = _extract_bracket_attr(s)
            tokens = re.findall(r'"[^"]+"|[A-Za-z0-9_.:-]+', prefix)
            for src_raw, dst_raw in zip(tokens[:-1], tokens[1:]):
                src, dst = _unquote_id(src_raw), _unquote_id(dst_raw)
                if src and dst and not _is_pseudo_node(src) and not _is_pseudo_node(dst):
                    attrs_by_edge[(src, dst)] = attrs
    body = _outer_graph_body(dot)
    parse_edge_statements(_remove_subgraph_blocks(body))
    for _header, subgraph_body in _find_subgraph_blocks(dot): parse_edge_statements(subgraph_body)
    for match in EDGE_RE.finditer(dot or ""):
        src, dst = _unquote_id(match.group(1)), _unquote_id(match.group(3))
        if src and dst and not _is_pseudo_node(src) and not _is_pseudo_node(dst): attrs_by_edge.setdefault((src, dst), {})
    return attrs_by_edge

def _run_graphviz_plain(dot_text: str) -> str:
    dot_bin = shutil.which("dot")
    if not dot_bin:
        raise RuntimeError("Graphviz 'dot' command not found. Install Graphviz or ensure it is on PATH.")
    with tempfile.TemporaryDirectory() as tmpdir:
        dot_path = Path(tmpdir) / "diagram.dot"
        dot_path.write_text(dot_text, encoding="utf-8")
        result = subprocess.run([dot_bin, "-Tplain", str(dot_path)], capture_output=True, text=True, check=False)
        if result.returncode != 0:
            raise RuntimeError(result.stderr or "Graphviz plain layout failed")
        return result.stdout

def _parse_graphviz_plain(plain_text: str) -> Dict[str, Any]:
    graph = {"scale": 1.0, "width": 1.0, "height": 1.0, "nodes": {}, "edges": []}
    for raw_line in (plain_text or "").splitlines():
        line = raw_line.strip()
        if not line: continue
        parts = _shell_split_graphviz(line)
        if not parts: continue
        kind = parts[0]
        if kind == "graph" and len(parts) >= 4:
            graph["scale"] = float(parts[1])
            graph["width"] = max(0.1, float(parts[2]))
            graph["height"] = max(0.1, float(parts[3]))
        elif kind == "node" and len(parts) >= 6:
            name = parts[1]
            graph["nodes"][name] = {
                "id": name,
                "x": float(parts[2]),
                "y": float(parts[3]),
                "width": float(parts[4]),
                "height": float(parts[5]),
                "label": _clean_label(parts[6]) if len(parts) >= 7 else name,
                "raw": parts,
            }
        elif kind == "edge" and len(parts) >= 5:
            tail, head, n = parts[1], parts[2], int(parts[3])
            coords = []
            idx = 4
            for _ in range(n):
                if idx + 1 >= len(parts): break
                coords.append((float(parts[idx]), float(parts[idx + 1])))
                idx += 2
            label = ""
            if idx < len(parts):
                possible = parts[idx]
                if not re.fullmatch(r"-?\d+(\.\d+)?", possible):
                    label = _clean_label(possible)
            graph["edges"].append({"source": tail, "target": head, "points": coords, "label": label, "raw": parts})
    return graph

def _resolve_image_path(attrs: Dict[str, str], download_dir: Path) -> Optional[Path]:
    """Dynamically pull the exact image path specified in the DOT language without relying on local manifests."""
    image_uri = attrs.get("image") or attrs.get("icon")
    if not image_uri:
        return None
    
    if image_uri.startswith("http://") or image_uri.startswith("https://"):
        safe_name = _safe_name(image_uri.split("/")[-1])
        local_path = download_dir / safe_name
        if not local_path.exists():
            try:
                urllib.request.urlretrieve(image_uri, local_path)
            except Exception as e:
                logger.debug("Failed to download image %s: %s", image_uri, e)
                return None
        return local_path
    
    local_path = Path(image_uri)
    return local_path if local_path.exists() else None

def _build_model_from_dot(diagram_text: str, config: DiagramConfig, assets_dir: Path) -> Dict[str, Any]:
    normalized = normalize_to_graphviz_dot(diagram_text)
    if not normalized:
        raise ValueError("Diagram text could not be normalized to DOT.")
    dot = _strip_dot_comments(html.unescape(normalized))
    node_attrs, node_order = _parse_node_attrs_from_dot(dot)
    edge_attrs = _parse_edge_attrs_from_dot(dot)
    clusters = _parse_clusters_from_dot(dot)
    plain = _run_graphviz_plain(dot)
    gv = _parse_graphviz_plain(plain)

    nodes: Dict[str, Dict[str, Any]] = {}
    for node_id, gv_node in gv["nodes"].items():
        attrs = node_attrs.get(node_id, {})
        label = _clean_label(attrs.get("label") or gv_node.get("label") or node_id)
        if _is_pseudo_node(node_id, label):
            continue
            
        nodes[node_id] = {
            **gv_node,
            "label": label,
            "attrs": attrs,
            "icon_path": _resolve_image_path(attrs, assets_dir),
        }
        if node_id not in node_order:
            node_order.append(node_id)

    edges: List[Dict[str, Any]] = []
    for gv_edge in gv["edges"]:
        src, dst = gv_edge["source"], gv_edge["target"]
        if src not in nodes or dst not in nodes:
            continue
        attrs = edge_attrs.get((src, dst), {})
        edges.append({
            **gv_edge,
            "raw_color": attrs.get("color") or attrs.get("fontcolor"),
            "raw_style": attrs.get("style"),
            "raw_penwidth": attrs.get("penwidth"),
            "label": _clean_label(attrs.get("label") or gv_edge.get("label") or ""),
            "attrs": attrs,
        })

    _assign_flow_visuals(edges, nodes, config)
    return {"dot": dot, "graph": gv, "nodes": [nodes[nid] for nid in node_order if nid in nodes], "node_map": nodes, "edges": edges, "clusters": clusters, "legend": _legend_from_edges(edges, config)}

def _assign_flow_visuals(edges: List[Dict[str, Any]], nodes: Dict[str, Dict[str, Any]], config: DiagramConfig) -> None:
    used_by_label: Dict[str, str] = {}
    palette_idx = 0
    for edge in edges:
        src_label = _clean_label(nodes.get(edge["source"], {}).get("label", edge["source"])).replace("\n", " ")[:28]
        dst_label = _clean_label(nodes.get(edge["target"], {}).get("label", edge["target"])).replace("\n", " ")[:28]
        flow_label = edge.get("label") or f"{src_label} → {dst_label}"
        edge["display_label"] = flow_label
        raw_color = _safe_color(edge.get("raw_color"), config, "") if edge.get("raw_color") else ""
        if raw_color:
            color = raw_color
        elif flow_label in used_by_label:
            color = used_by_label[flow_label]
        else:
            color = config.flow_palette[palette_idx % len(config.flow_palette)]
            used_by_label[flow_label] = color
            palette_idx += 1
        edge["color"] = color
        edge["dash_style"] = _dash_style_from_raw_style(edge.get("raw_style"))
        edge["width_pt"] = _safe_float(edge.get("raw_penwidth"), config.default_edge_width_pt)

def _dash_style_from_raw_style(raw_style: Optional[str]) -> Optional[MSO_LINE_DASH_STYLE]:
    style = str(raw_style or "").lower()
    if "dotted" in style: return MSO_LINE_DASH_STYLE.ROUND_DOT
    if "dashed" in style: return MSO_LINE_DASH_STYLE.DASH
    if "dashdot" in style or "dash-dot" in style: return MSO_LINE_DASH_STYLE.DASH_DOT
    return None

def _legend_from_edges(edges: List[Dict[str, Any]], config: DiagramConfig) -> List[Tuple[str, str, Optional[MSO_LINE_DASH_STYLE]]]:
    legend, seen = [], set()
    for edge in edges:
        label = edge.get("display_label") or edge.get("label")
        color = edge.get("color") or config.default_edge_color
        dash = edge.get("dash_style")
        key = (color, label, str(dash))
        if label and key not in seen:
            seen.add(key)
            legend.append((color, label, dash))
    return legend

def _compute_transform(model: Dict[str, Any], config: DiagramConfig) -> Dict[str, float]:
    graph = model["graph"]
    g_w = max(0.1, float(graph.get("width", 1.0)))
    g_h = max(0.1, float(graph.get("height", 1.0)))
    scale = min(config.canvas_width_in / g_w, config.canvas_height_in / g_h)
    draw_w = g_w * scale
    draw_h = g_h * scale
    left = config.canvas_left_in + (config.canvas_width_in - draw_w) / 2.0
    top = config.canvas_top_in + (config.canvas_height_in - draw_h) / 2.0
    return {"scale": scale, "left": left, "top": top, "graph_h": g_h}

def _gv_to_ppt(x: float, y: float, t: Dict[str, float]) -> Tuple[float, float]:
    return t["left"] + x * t["scale"], t["top"] + (t["graph_h"] - y) * t["scale"]

def _node_box_from_gv(node: Dict[str, Any], t: Dict[str, float]) -> Dict[str, float]:
    cx, cy = _gv_to_ppt(node["x"], node["y"], t)
    w = max(0.42, node["width"] * t["scale"])
    h = max(0.30, node["height"] * t["scale"])
    return {"left": cx - w / 2.0, "top": cy - h / 2.0, "width": w, "height": h}

def _cluster_boxes(model: Dict[str, Any], placements: Dict[str, Dict[str, float]], config: DiagramConfig) -> List[Dict[str, Any]]:
    boxes = []
    for cluster in model.get("clusters", []):
        node_boxes = [placements[nid] for nid in cluster.get("nodes", []) if nid in placements]
        if not node_boxes:
            continue
        left = min(b["left"] for b in node_boxes) - 0.30
        top = min(b["top"] for b in node_boxes) - 0.36
        right = max(b["left"] + b["width"] for b in node_boxes) + 0.30
        bottom = max(b["top"] + b["height"] for b in node_boxes) + 0.28
        boxes.append({"left": left, "top": top, "width": right-left, "height": bottom-top, "label": cluster.get("label", ""), "fill": config.lane_fill, "border": config.lane_border})
    return boxes

def _set_text_frame(shape: BaseShape, text: str, font_size_pt: float = 10, bold: bool = False, color_hex: str = "000000", align=PP_ALIGN.LEFT) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = _safe_title(text)
    run.font.name = "Calibri"
    run.font.size = Pt(font_size_pt)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color_hex)

def _font_for_label(label: str, width: float, height: float) -> float:
    clean = _safe_title(label)
    line_count = max(1, clean.count("\n") + 1)
    length = len(clean.replace("\n", " "))
    if width < 0.9 or height < 0.45 or length > 58 or line_count >= 3: return 6.6
    if width < 1.25 or length > 42 or line_count == 2: return 7.4
    return 8.4

def _add_title(slide, title_text: str, config: DiagramConfig) -> None:
    shape = slide.shapes.add_textbox(Inches(config.title_left_in), Inches(config.title_top_in), Inches(config.title_width_in), Inches(config.title_height_in))
    _set_text_frame(shape, title_text, font_size_pt=18, bold=True, color_hex="202124")

def _add_cluster_box(slide, box: Dict[str, Any], config: DiagramConfig) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(box.get("fill", config.lane_fill))
    shape.line.color.rgb = _rgb(box.get("border", config.lane_border))
    shape.line.width = Pt(1.0)
    try:
        shape.adjustments[0] = 0.03
    except Exception:
        pass
    label = box.get("label") or ""
    if label:
        label_box = slide.shapes.add_textbox(Inches(box["left"] + 0.07), Inches(box["top"] + 0.03), Inches(max(0.5, box["width"] - 0.14)), Inches(0.18))
        _set_text_frame(label_box, label, font_size_pt=8.2, bold=True, color_hex=config.lane_title_color)

def _add_node_box(slide, box: Dict[str, float], label: str, icon_path: Optional[Path], config: DiagramConfig) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(box["left"]), Inches(box["top"]), Inches(box["width"]), Inches(box["height"]))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(config.default_node_fill)
    shape.line.color.rgb = _rgb(config.default_node_border)
    shape.line.width = Pt(0.9)
    try:
        shape.adjustments[0] = 0.07
    except Exception:
        pass
    icon_size = min(config.node_icon_size_in, max(0.13, box["height"] * 0.34))
    text_left = box["left"] + config.node_padding_in
    if icon_path and icon_path.exists() and box["width"] > 0.68:
        try:
            slide.shapes.add_picture(str(icon_path), Inches(box["left"] + config.node_padding_in), Inches(box["top"] + (box["height"] - icon_size) / 2.0), width=Inches(icon_size), height=Inches(icon_size))
            text_left = box["left"] + config.node_text_left_with_icon_in
        except Exception:
            logger.debug("Could not add icon %s", icon_path, exc_info=True)
    text_width = max(0.32, box["width"] - (text_left - box["left"]) - 0.04)
    text_box = slide.shapes.add_textbox(Inches(text_left), Inches(box["top"] + 0.02), Inches(text_width), Inches(max(0.20, box["height"] - 0.04)))
    _set_text_frame(text_box, label, font_size_pt=_font_for_label(label, box["width"], box["height"]), bold=True, color_hex=config.default_node_font)

def _shorten_last_segment_for_arrowhead(points: List[Tuple[float, float]], gap: float) -> List[Tuple[float, float]]:
    if len(points) < 2: return points
    adjusted = list(points)
    x1, y1 = adjusted[-2]
    x2, y2 = adjusted[-1]
    dx, dy = x2 - x1, y2 - y1
    length = math.sqrt(dx * dx + dy * dy)
    if length <= 0.001: return adjusted
    usable_gap = min(gap, length * 0.35)
    adjusted[-1] = (x2 - (dx / length) * usable_gap, y2 - (dy / length) * usable_gap)
    return adjusted

def _set_line_end_arrow(connector: BaseShape, config: DiagramConfig) -> None:
    try:
        ln = connector._element.spPr.get_or_add_ln()
        for child in list(ln):
            if child.tag.endswith("}tailEnd") or child.tag.endswith("}headEnd"):
                ln.remove(child)
        tail_end = OxmlElement("a:tailEnd")
        tail_end.set("type", config.arrow_head_type)
        tail_end.set("w", config.arrow_head_width)
        tail_end.set("len", config.arrow_head_length)
        ln.append(tail_end)
    except Exception:
        logger.debug("Could not apply arrowhead", exc_info=True)

def _add_line_segment(slide, p1: Tuple[float, float], p2: Tuple[float, float], color_hex: str, config: DiagramConfig, add_arrowhead: bool = False, dash_style: Optional[MSO_LINE_DASH_STYLE] = None, width_pt: float = 1.70):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(p1[0]), Inches(p1[1]), Inches(p2[0]), Inches(p2[1]))
    connector.line.color.rgb = _rgb(color_hex, config.default_edge_color)
    connector.line.width = Pt(width_pt)
    if dash_style is not None:
        try: connector.line.dash_style = dash_style
        except Exception: pass
    if add_arrowhead:
        _set_line_end_arrow(connector, config)
    return connector

def _add_polyline_arrow(slide, points: List[Tuple[float, float]], color_hex: str, config: DiagramConfig, dash_style: Optional[MSO_LINE_DASH_STYLE] = None, width_pt: float = 1.70) -> None:
    if len(points) < 2: return
    clean = [points[0]]
    for point in points[1:]:
        if abs(point[0] - clean[-1][0]) > 0.01 or abs(point[1] - clean[-1][1]) > 0.01:
            clean.append(point)
    if len(clean) < 2: return
    clean = _shorten_last_segment_for_arrowhead(clean, config.arrow_tip_gap_in)
    simplified = [clean[0]]
    for i in range(1, len(clean) - 1):
        prev, cur, nxt = simplified[-1], clean[i], clean[i + 1]
        same_x = abs(prev[0] - cur[0]) < 0.01 and abs(cur[0] - nxt[0]) < 0.01
        same_y = abs(prev[1] - cur[1]) < 0.01 and abs(cur[1] - nxt[1]) < 0.01
        if same_x or same_y: continue
        simplified.append(cur)
    simplified.append(clean[-1])
    for idx in range(len(simplified) - 1):
        _add_line_segment(slide, simplified[idx], simplified[idx + 1], color_hex, config, add_arrowhead=(idx == len(simplified) - 2), dash_style=dash_style, width_pt=width_pt)

def _add_legend(slide, legend_items: List[Tuple[str, str, Optional[MSO_LINE_DASH_STYLE]]], config: DiagramConfig) -> None:
    if not legend_items: return
    title_box = slide.shapes.add_textbox(Inches(config.legend_left_in), Inches(config.legend_top_in), Inches(1.15), Inches(0.18))
    _set_text_frame(title_box, "Flow Legend", font_size_pt=9.2, bold=True, color_hex="202124")
    x, y, max_x = config.legend_left_in, config.legend_top_in + 0.24, config.legend_left_in + config.legend_width_in
    for color, label, dash_style in legend_items[:18]:
        label = _safe_title(label)
        item_w = max(1.55, min(3.30, 0.62 + len(label) * 0.050))
        if x + item_w > max_x:
            x, y = config.legend_left_in, y + 0.22
        if y > config.legend_top_in + config.legend_max_height_in: break
        chip = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(item_w), Inches(0.18))
        chip.fill.solid(); chip.fill.fore_color.rgb = _rgb(config.legend_fill); chip.line.color.rgb = _rgb(config.legend_border); chip.line.width = Pt(0.6)
        line_y = y + 0.09
        _add_line_segment(slide, (x + 0.05, line_y), (x + 0.18, line_y), color, config, add_arrowhead=True, dash_style=dash_style, width_pt=1.35)
        label_box = slide.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.01), Inches(max(0.3, item_w - 0.24)), Inches(0.15))
        _set_text_frame(label_box, label, font_size_pt=7.2, bold=False, color_hex="202124")
        x += item_w + 0.08

def _add_fallback_preview_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str, config: DiagramConfig) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title(slide, f"{slide_title} (Preview)", config)
    try:
        assets = render_graphviz_to_assets(diagram_text, output_dir=output_dir, base_name=base_name)
    except Exception:
        logger.exception("Fallback preview asset generation failed")
        assets = None
    if assets and assets.get("png_path") and Path(assets["png_path"]).exists():
        slide.shapes.add_picture(str(assets["png_path"]), Inches(config.canvas_left_in), Inches(config.canvas_top_in), width=Inches(config.preview_max_width_in))
        _add_legend(slide, assets.get("legend", []), config)

def extract_diagram_specs(data: Dict[str, Any], hld_model: Any, selected_sections_set: Optional[Set[str]] = None) -> List[Dict[str, Any]]:
    results: List[Dict[str, Any]] = []
    
    def _walk_for_diagrams(value: Any, path_parts: List[str], results: List[Dict[str, Any]]) -> None:
        if value is None: return
        if isinstance(value, str):
            if normalize_to_graphviz_dot(_clean_internal_tokens(value)) is not None:
                results.append({"title": " / ".join(path_parts) if path_parts else "Diagram", "diagram_text": _clean_internal_tokens(value), "section_key": _safe_name(path_parts[-1] if path_parts else "diagram")})
            return
        if isinstance(value, list):
            for idx, item in enumerate(value, start=1):
                if isinstance(item, str) and normalize_to_graphviz_dot(_clean_internal_tokens(item)) is not None:
                    results.append({"title": " / ".join(path_parts) if path_parts else f"Diagram {idx}", "diagram_text": _clean_internal_tokens(item), "section_key": _safe_name(path_parts[-1] if path_parts else f"diagram_{idx}")})
                elif isinstance(item, (dict, list)):
                    _walk_for_diagrams(item, path_parts, results)
            return
        if isinstance(value, dict):
            for key, child in value.items():
                key_text = str(key).replace("_", " ").title()
                next_path = path_parts if str(key).lower() == "diagrams" else path_parts + [key_text]
                _walk_for_diagrams(child, next_path, results)

    for field_name, field_info in hld_model.model_fields.items():
        if not should_include_section(field_name, field_info, selected_sections_set or set()): continue
        section_title = field_info.title or field_name.replace("_", " ").title()
        _walk_for_diagrams(data.get(field_name), [section_title], results)
    counts: Dict[str, int] = defaultdict(int)
    for item in results: counts[item["title"]] += 1
    seen: Dict[str, int] = defaultdict(int)
    for item in results:
        if counts[item["title"]] > 1:
            seen[item["title"]] += 1
            item["title"] = f'{item["title"]} ({seen[item["title"]]})'
    return results

def add_diagram_slide(prs: Presentation, slide_title: str, diagram_text: str, output_dir: Path, base_name: str, config: DiagramConfig) -> None:
    try:
        model = _build_model_from_dot(diagram_text, config, output_dir)
        t = _compute_transform(model, config)
        placements = {node["id"]: _node_box_from_gv(node, t) for node in model["nodes"]}
        cluster_boxes = _cluster_boxes(model, placements, config)
    except Exception:
        logger.exception("Graphviz-coordinate PPTX rendering failed for '%s'. Falling back to preview image.", slide_title)
        _add_fallback_preview_slide(prs, slide_title, diagram_text, output_dir, base_name, config)
        return

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_title(slide, slide_title, config)

    for box in cluster_boxes:
        _add_cluster_box(slide, box, config)

    for edge in model["edges"]:
        pts = [_gv_to_ppt(x, y, t) for x, y in edge.get("points", [])]
        _add_polyline_arrow(slide, pts, edge.get("color") or config.default_edge_color, config, dash_style=edge.get("dash_style"), width_pt=edge.get("width_pt", config.default_edge_width_pt))

    for node in model["nodes"]:
        box = placements.get(node["id"])
        if not box: continue
        _add_node_box(slide, box, node.get("label") or node["id"], node.get("icon_path"), config)

    _add_legend(slide, model.get("legend", []), config)

def build_editable_pptx(data: Dict[str, Any], pptx_path: str, hld_model: Any, selected_sections_set: Optional[Set[str]] = None) -> None:
    output_path = Path(pptx_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    diagram_assets_dir = output_path.parent / f"{output_path.stem}_diagram_assets"
    diagram_assets_dir.mkdir(parents=True, exist_ok=True)
    
    # Initialize our generic config payload to replace hardcoded values
    config = DiagramConfig()
    
    diagram_specs = extract_diagram_specs(data=data, hld_model=hld_model, selected_sections_set=selected_sections_set)
    prs = Presentation()
    prs.slide_width = Inches(config.slide_width_in)
    prs.slide_height = Inches(config.slide_height_in)
    
    if not diagram_specs:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_title(slide, "Editable Diagrams", config)
        msg = slide.shapes.add_textbox(Inches(0.7), Inches(1.8), Inches(11.5), Inches(1.0))
        _set_text_frame(msg, "No diagrams were detected in the selected sections.", font_size_pt=16, bold=True, color_hex="5F6368", align=PP_ALIGN.CENTER)
        prs.save(str(output_path))
        return
        
    for counter, spec in enumerate(diagram_specs, start=1):
        add_diagram_slide(prs, spec["title"], spec["diagram_text"], diagram_assets_dir, f"{_safe_name(spec['section_key'])}_{counter:02d}", config)
        
    prs.save(str(output_path))
    logger.info("editable_pptx_generated: %s", output_path)